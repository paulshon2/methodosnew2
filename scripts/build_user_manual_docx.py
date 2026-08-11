# -*- coding: utf-8 -*-
"""
MethodOS / AI-Research-OS · 출판수준 사용자 매뉴얼 (DOCX)
대상 제품: methodos--s23-all-23 (연구방법 OS) + 듀얼 플랜(연구준비자 RDOS / 연구자)
"""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\MethodOS_AI-Research-OS_사용자매뉴얼_출판본.docx")


def set_run_font(run, name="맑은 고딕", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_para(doc, text, *, size=11, bold=False, space_after=8, space_before=0, align=None, color=None):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.35
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, color=color)
    return p


def add_h(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size={1: 18, 2: 14, 3: 12}.get(level, 11), bold=True, color=RGBColor(0x1A, 0x1A, 0x2E))
    return h


def add_bullets(doc, items, size=10.5):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        p.clear()
        run = p.add_run(it)
        set_run_font(run, size=size)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25


def add_steps(doc, items, size=10.5):
    for i, it in enumerate(items, 1):
        p = doc.add_paragraph(style="List Number")
        p.clear()
        run = p.add_run(it)
        set_run_font(run, size=size)
        p.paragraph_format.space_after = Pt(3)


def add_table(doc, headers, rows, col_widths=None):
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        # shade
        shading = hdr[i]._tePr if False else None
        from docx.oxml import OxmlElement

        tc = hdr[i]._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "1A1A2E")
        shd.set(qn("w:val"), "clear")
        tcPr.append(shd)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.rows[ri + 1].cells[ci]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(val))
            set_run_font(run, size=9)
    doc.add_paragraph()
    return t


def callout(doc, title, body):
    add_para(doc, f"【안내】{title}", size=10.5, bold=True, color=RGBColor(0x6B, 0x2D, 0x5B), space_after=2)
    add_para(doc, body, size=10, space_after=10)


def build():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.2)
        section.bottom_margin = Cm(2.2)
        section.left_margin = Cm(2.4)
        section.right_margin = Cm(2.4)

    # —— 표지 ——
    add_para(doc, "", space_after=40)
    add_para(
        doc,
        "AI-Research-OS / MethodOS",
        size=14,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=RGBColor(0x6B, 0x2D, 0x5B),
    )
    add_para(
        doc,
        "사용자 매뉴얼",
        size=28,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=RGBColor(0x1A, 0x1A, 0x2E),
        space_after=6,
    )
    add_para(
        doc,
        "듀얼 플랜(연구준비자 · 연구자) × 등급(Free · Basic · Scholar)\n메뉴·기능·절차적 프로세스 단계별 사용 지침",
        size=12,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_after=18,
    )
    add_para(
        doc,
        f"기준 제품 빌드: methodos--s23-all-23\n참조 플랫폼: AI-Research-OS / RDOS 듀얼 트랙\n문서 버전: 1.0  ·  발행일: {date.today().isoformat()}\n문서 성격: 출판·교육용 사용자 지침서",
        size=10,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        color=RGBColor(0x55, 0x55, 0x55),
        space_after=24,
    )
    add_para(
        doc,
        "본 매뉴얼은 화면 메뉴를 ‘연구의 절차’ 순서에 맞춰 설명합니다. "
        "처음 쓰는 분은 제1부(듀얼 플랜) → 제2부(등급) → 제3부(연구자 MethodOS 메뉴 절차) 순으로 읽으십시오.",
        size=10.5,
        align=WD_ALIGN_PARAGRAPH.CENTER,
    )
    doc.add_page_break()

    # —— 목차 안내 ——
    add_h(doc, "목차", 1)
    toc = [
        "제1부  플랫폼 구조와 듀얼 플랜모드",
        "제2부  회원 등급(Free · Basic · Scholar)과 권한",
        "제3부  연구자 플랜 · MethodOS 메뉴 절차 매뉴얼 (전체)",
        "제4부  연구준비자 플랜 · RDOS 메뉴 절차 매뉴얼",
        "제5부  Free 등급 전용 사용 시나리오",
        "제6부  Basic 등급 전용 사용 시나리오",
        "제7부  Scholar 등급 전용 사용 시나리오(전체 기능)",
        "제8부  AI 기능·설정·내보내기·문제 해결",
        "부록  용어·단축 경로·체크리스트",
    ]
    add_bullets(doc, toc)
    doc.add_page_break()

    # ===================== 제1부 =====================
    add_h(doc, "제1부  플랫폼 구조와 듀얼 플랜모드", 1)
    add_para(
        doc,
        "AI-Research-OS는 하나의 브랜드 아래에서 두 가지 플랜모드(트랙)로 나뉩니다. "
        "메뉴와 목적이 다르므로, 가입·승인 후 ‘어느 트랙으로 연구할 것인가’를 먼저 고르는 것이 올바른 시작입니다.",
    )

    add_h(doc, "1.1 두 플랜모드 한눈에 보기", 2)
    add_table(
        doc,
        ["구분", "연구준비자 플랜 (RDOS)", "연구자 플랜 (AI-Research-OS / MethodOS)"],
        [
            [
                "대상",
                "대학원 진학 예정자, 석사 입문자, 연구 초심자",
                "논문·연구를 실제로 수행하는 연구자(석사 이상·연구자·교수 등)",
            ],
            [
                "목적",
                "논문·연구 역량 개발(교육·훈련)",
                "연구 수행·방법 선택·분석·산출물 생산(생산성)",
            ],
            [
                "접근",
                "교육 중심 학습 모듈",
                "연구방법 OS·워크벤치·분석 도구",
            ],
            [
                "대표 경로",
                "/rdos (RDOS 셸)",
                "MethodOS 앱(#/… 해시 라우트) · /dashboard",
            ],
            [
                "승급",
                "미션 완료 후 연구자 플랜으로 승급 가능",
                "Free → Basic → Scholar(·University) 등급",
            ],
        ],
    )
    callout(
        doc,
        "methodos--s23-all-23의 위치",
        "본 매뉴얼의 상세 메뉴 절차(제3부)는 methodos--s23-all-23(연구방법 가상 학습·아키타입·기본통계 등)을 "
        "연구자 플랜의 핵심 연구방법 OS로 설명합니다. 연구준비자 플랜 메뉴는 제4부에서 별도로 다룹니다.",
    )

    add_h(doc, "1.2 사용자가 고르는 첫 분기", 2)
    add_steps(
        doc,
        [
            "랜딩(또는 컨트롤 허브)에서 「연구준비자 플랜」과 「연구자 플랜」을 확인합니다.",
            "연구 입문·기초 훈련이 목표면 연구준비자(RDOS)로 가입·승인을 받습니다.",
            "실제 논문 방법·설문·코딩·통계·시각화가 목표면 연구자 플랜으로 들어가 MethodOS를 사용합니다.",
            "연구준비자에서 미션을 마치면 연구자 플랜 승급을 신청할 수 있습니다(관리자 승인 정책에 따름).",
        ],
    )

    add_h(doc, "1.3 연구자 플랜 안에서 MethodOS가 하는 일", 2)
    add_para(
        doc,
        "연구자 플랜의 MethodOS는 ‘어떤 연구방법으로, 어떤 절차로, 어떤 산출물을 만들까’를 "
        "메뉴로 분해해 둔 운영체제입니다. 추천 → 프로세스 이해 → 방법 가상 학습 → 아키타입 예비학습·실습 → "
        "설문·코딩·기본통계·시각화 → 연구방법작성 → 내 프로젝트 저장으로 이어집니다.",
    )
    doc.add_page_break()

    # ===================== 제2부 =====================
    add_h(doc, "제2부  회원 등급(Free · Basic · Scholar)과 권한", 1)
    add_para(
        doc,
        "연구자 플랜(MethodOS)에서는 사이드바 메뉴 대부분은 Free/Basic에서도 보이지만, "
        "실제로 열 수 있는 연구방법 개수와 그래프 종류가 등급에 따라 달라집니다. "
        "Scholar(및 University)는 전체 연구방법·전체 시각화를 사용합니다.",
    )

    add_h(doc, "2.1 등급별 권한 요약표", 2)
    add_table(
        doc,
        ["항목", "Free", "Basic", "Scholar"],
        [
            [
                "사이드바 메뉴",
                "대시보드~설정(관리자 제외) 표시",
                "Free와 동일(관리자 제외)",
                "전체 메뉴(+기관은 University)",
            ],
            [
                "연구방법(가상 학습)",
                "2개: SV-003(설문지 작성), QT-013(사회과학 통계분석2)",
                "5개: Free + QL-044(포커스그룹), QL-041(심층면접), QL-060(질적 코딩 매뉴얼)",
                "전체(레지스트리 약 210개)",
            ],
            [
                "데이터 시각화",
                "기본 그래프 10종",
                "기본 10 + 통계·분포 8 + 과학·관계 8 ≈ 26종",
                "전 카테고리(재무·계층·지도·3D 등 포함)",
            ],
            [
                "아키타입·설문·코딩·기본통계·방법작성",
                "메뉴 진입 가능(방법·차트 권한이 필요한 하위 기능은 제한)",
                "Free보다 질적 방법·차트 확장",
                "제한 없이 전체 기능",
            ],
            ["프로젝트 저장", "가능", "가능", "가능"],
            ["관리자", "불가", "불가", "슈퍼관리자/기관 정책"],
        ],
    )
    callout(
        doc,
        "잠김 메시지",
        "권한이 없는 연구방법·그래프를 열면 “현재 등급에서 사용할 수 없습니다”라는 안내가 나옵니다. "
        "업그레이드 또는 관리자 개별 허용이 필요합니다.",
    )

    add_h(doc, "2.2 Free에서 쓰는 연구방법 ID", 2)
    add_bullets(
        doc,
        [
            "SV-003 — 설문지 작성(Survey Question Writing): 설문 설계·문항 워크벤치",
            "QT-013 — 사회과학 통계분석2: 단계형 통계 워크벤치",
        ],
    )

    add_h(doc, "2.3 Basic에서 추가되는 연구방법 ID", 2)
    add_bullets(
        doc,
        [
            "QL-044 — Focus Groups 계열: 초점집단 절차 학습",
            "QL-041 — 심층면접 절차",
            "QL-060 — 질적연구 코딩 매뉴얼",
        ],
    )

    add_h(doc, "2.4 권한 덮어쓰기(관리자)", 2)
    add_para(
        doc,
        "관리자는 (1) 전 회원 공통 잠금/허용, (2) 회원별 개별 권한으로 menu.* · method.* · viz.* 코드를 조정할 수 있습니다. "
        "개별 권한이 최종 우선입니다. 따라서 같은 Free 회원이라도 기관 정책에 따라 메뉴가 달라질 수 있습니다.",
    )
    doc.add_page_break()

    # ===================== 제3부 MethodOS =====================
    add_h(doc, "제3부  연구자 플랜 · MethodOS 메뉴 절차 매뉴얼", 1)
    add_para(
        doc,
        "이 부는 methodos--s23-all-23의 사이드바를 ‘연구 절차’ 순서로 설명합니다. "
        "각 절은 (A) 이 메뉴가 하는 일 (B) 언제 쓰는지 (C) 화면에서 누르는 순서 (D) 산출물 (E) 등급 주의로 구성합니다.",
    )

    # 3.0 전체 맵
    add_h(doc, "3.0 연구자 MethodOS 절차 지도(권장 순서)", 2)
    add_steps(
        doc,
        [
            "설정: AI 키·테마·모드(규칙/AI/하이브리드) 확인",
            "대시보드: 최근 프로젝트·바로가기 확인",
            "연구방법 추천: 논문·연구 조건을 넣고 후보 방법 받기",
            "연구방법별 프로세스 / 연구방법 체계: 절차·계보 이해",
            "연구방법 가상 학습: 방법 선택 → 프로젝트 시작 → 11단계 워크벤치",
            "방법론적 아키타입: 개요·용어 → 프로젝트 예비학습 → 프로젝트 실습",
            "설문조사 / 코딩 / 기본통계 / 데이터 시각화: 자료·분석·그림",
            "연구방법작성: 산출물을 모아 방법 절 초안",
            "내 프로젝트: 저장·재개",
        ],
    )

    menus_detail = [
        (
            "3.1 대시보드",
            "연구의 ‘상황판’입니다. 방법 수 KPI, 추천·프로세스·방법·시각화 바로가기, 최근 프로젝트가 보입니다.",
            [
                "로그인 후 기본 진입 화면을 확인합니다.",
                "최근 프로젝트를 눌러 이어서 작업하거나, 추천/방법 메뉴로 이동합니다.",
            ],
            "최근 프로젝트 목록, 바로가기",
            "전 등급 사용",
        ),
        (
            "3.2 연구방법 추천",
            "연구 조건(분야·목적·자료·설계 등)을 입력하면 적합한 연구방법 후보를 순위와 함께 제안합니다. AI 보완 해석을 켤 수 있습니다.",
            [
                "체크리스트·입력란에 연구 상황을 채웁니다.",
                "추천 실행 → 후보 방법 카드를 검토합니다.",
                "마음에 드는 방법으로 ‘방법 열기’ 또는 ‘프로젝트 만들기’로 이어갑니다.",
            ],
            "추천 결과 목록, (선택) AI 해석 문단",
            "전 등급 메뉴 사용. 이후 열리는 방법이 등급 허용 ID인지 확인",
        ),
        (
            "3.3 연구방법별 프로세스",
            "방법 공통 절차를 ‘프로세스 요약 / 단계별 상세 / 공통절차’로 학습합니다.",
            [
                "허브에서 세 카드 중 하나를 고릅니다.",
                "요약 → 단계 → 공통절차 순으로 읽으며 자기 연구에 매핑합니다.",
            ],
            "절차 이해 노트(사용자 메모)",
            "전 등급",
        ),
        (
            "3.4 연구방법 체계",
            "방법 가족·계보·분류표를 펼쳐 보는 참고 메뉴입니다.",
            [
                "섹션을 펼쳐 계보·분류를 확인합니다.",
                "가상 학습에서 고를 방법의 ‘가족’을 미리 파악합니다.",
            ],
            "분류 이해",
            "전 등급",
        ),
    ]

    for title, what, steps, out, tier in menus_detail:
        add_h(doc, title, 2)
        add_para(doc, f"(A) 하는 일 — {what}", size=10.5)
        add_para(doc, "(C) 사용 순서", size=10.5, bold=True, space_after=2)
        add_steps(doc, steps)
        add_para(doc, f"(D) 산출물 — {out}", size=10.5)
        add_para(doc, f"(E) 등급 — {tier}", size=10.5, space_after=12)

    # 3.5 연구방법 가상 학습 — 핵심 장
    add_h(doc, "3.5 연구방법 가상 학습 (핵심 · 약 210개 방법)", 2)
    add_para(
        doc,
        "(A) 하는 일 — 방법 카탈로그에서 방법을 고르고, 프로젝트로 들어가 "
        "11단계(영역) → 세부 Task → 수행내용/산출물 분절을 자동화 구분(수동/반자동/자동)에 따라 실행합니다. "
        "각 분절에는 쉬운 설명, 연습 예시, 산출물 서식 샘플, AI 변환기, (자료수집) 표준 표본수 안내와 100행+ 연습 데이터가 아코디언으로 제공됩니다.",
    )
    add_para(doc, "(B) 언제 — 추천 결과를 실제 절차로 연습할 때, 논문 방법 절을 쓰기 전 워밍업할 때.", size=10.5)
    add_para(doc, "(C) 절차적 사용 순서(아주 자세히)", size=10.5, bold=True, space_after=2)
    add_steps(
        doc,
        [
            "사이드바 「연구방법 가상 학습」을 엽니다.",
            "가족(양적·질적·조사·혼합 등) 필터로 목록을 좁힙니다. Free/Basic이면 허용된 방법만 실행 가능합니다.",
            "방법 카드 → 상세(#/method/ID)에서 개요·프로세스를 읽습니다.",
            "「프로젝트 시작」으로 #/project/… 워크벤치에 들어갑니다.",
            "11개 영역 카드 중 1단계부터 순서대로 엽니다(건너뛰면 산출 연동이 끊길 수 있음).",
            "세부 Task → 수행내용/산출물 항목을 엽니다.",
            "자동화 배지(수동/반자동/자동)를 확인합니다. 반자동은 AI 초안 + 연구자 확정입니다.",
            "연습 예시(미니 표)로 서식을 익힌 뒤, 아코디언에서 표준 표본수와 100행+ 데이터를 펼쳐 CSV로 저장해 연습합니다.",
            "산출물 서식 샘플을 보고 TXT/DOCX/PDF로 저장합니다.",
            "AI 변환기에 내 주제·자료를 넣고 ‘내 프로젝트 산출물’을 만듭니다. 자료 연동(이전 단계/합치기/새 자료만)을 선택합니다.",
            "실행 결과를 확인하고 다음 Task로 진행합니다. 특수 방법(SV-003 설문, QT-013 통계)은 전용 워크벤치가 열립니다.",
            "작업을 마치면 내 프로젝트에 상태가 남는지 대시보드/프로젝트 목록에서 확인합니다.",
        ],
    )
    add_para(doc, "(D) 산출물 — Task별 서술문·표·서식 샘플·AI 변환 결과·연습 CSV", size=10.5)
    add_para(
        doc,
        "(E) 등급 — Free 2방법, Basic 5방법, Scholar 전체. 잠긴 방법은 업그레이드 필요.",
        size=10.5,
        space_after=12,
    )

    add_h(doc, "3.5.1 11단계 워크벤치에서 ‘자동화’의 의미", 3)
    add_bullets(
        doc,
        [
            "수동화: 연구자가 직접 기입·판단. 도구는 안내·서식만 제공.",
            "반자동화: AI·규칙이 초안/후보를 만들면 연구자가 검증·확정(가장 흔함).",
            "자동화: 빈도·κ·일부 통계처럼 결정론 계산이 가능한 구간.",
        ],
    )

    add_h(doc, "3.5.2 자료수집 단계의 표본수 읽기", 3)
    add_para(
        doc,
        "미니 표(3~6행)는 서식 익힘용입니다. 실전 양적 조사는 보통 100명 이상(비교·상관 200~400, 모형 400~1,000+), "
        "질적 면담은 포화 기준 15~40명이 안내됩니다. 아코디언의 100행+ 표로 분석 연습을 하십시오.",
    )

    # 3.6 아키타입
    add_h(doc, "3.6 방법론적 아키타입 (31유형)", 2)
    add_para(
        doc,
        "(A) 하는 일 — 절차가 비슷한 방법끼리 묶은 31개 ‘프로세스 유형’을 학습합니다. "
        "허브 카드 4개: 개요 및 단계 / 용어 / 프로젝트 예비학습 / 프로젝트 실습.",
    )
    add_para(doc, "(C) 권장 순서", size=10.5, bold=True, space_after=2)
    add_steps(
        doc,
        [
            "「개요 및 단계」에서 31유형의 수집→전처리→코딩→통계→결과 뼈대를 훑습니다.",
            "「용어」에서 한글(English) 용어를 확인합니다.",
            "「프로젝트 예비학습」에서 유형 1개를 고르고 절차적 단계·태스크를 예행합니다(①왜 ②어떻게 ③프로토콜 ④⑤예시).",
            "「프로젝트 실습」에서 연구방법 → 주제 5개 중 하나(예: 신제품 컨셉 테스트)를 고릅니다.",
            "①자료수집~⑥최종산출 태스크를 순서대로 진행합니다. 「다음 태스크 진행」으로 연동합니다.",
            "자료수집/본조사 태스크에서 표준 표본수 안내와 완료 n·응답률에 맞춘 마이크로데이터(100행+) 아코디언을 펼쳐 CSV를 저장합니다.",
            "하단 AI 변환기 기입란에 내 주제를 넣고 산출물을 만든 뒤 Excel/PDF로 내보냅니다.",
        ],
    )
    add_para(doc, "(D) 산출물 — 태스크별 표·그림·AI 산출·마이크로데이터 CSV", size=10.5)
    add_para(doc, "(E) 등급 — 메뉴는 전 등급. Scholar에서 연계 방법·차트와 함께 쓰는 것이 가장 원활.", size=10.5, space_after=12)

    # 3.7~3.11
    add_h(doc, "3.7 설문조사", 2)
    add_para(
        doc,
        "(A) 설문 설계 스튜디오입니다. 개요·분류·프로토콜·AI 설문 생성·편집기·가져오기·패턴·내보내기 탭으로 진행합니다.",
    )
    add_steps(
        doc,
        [
            "프로토콜·분류를 읽습니다.",
            "AI 설문 생성 또는 수동 편집기로 문항을 만듭니다.",
            "가져오기/패턴으로 다듬고 내보냅니다.",
            "필요 시 SV-003 가상 학습 프로젝트와 병행합니다.",
        ],
    )

    add_h(doc, "3.8 코딩", 2)
    add_para(
        doc,
        "(A) 질적 코딩 스튜디오입니다. 프로토콜 → 코드북 → AI 파이프라인 → AI 변환기(단계별) → 표준 코드북 탭 순을 권장합니다.",
    )
    add_steps(
        doc,
        [
            "코딩 프로토콜과 윤리·신뢰도 기준을 확인합니다.",
            "코드북을 정의하거나 AI로 초안을 만듭니다.",
            "파이프라인으로 빈도·군집·코더 간 일치 등을 점검합니다.",
            "단계 변환기로 원문→코드→테마→보고 문장을 다듬습니다.",
            "Basic 이상이면 QL-060 가상 학습과 교차 연습합니다.",
        ],
    )

    add_h(doc, "3.9 기본통계", 2)
    add_para(
        doc,
        "(A) 기술·추론 분석(약 41)과 구조방정식·잠재변수(약 12)를 스토리형으로 학습·실행합니다. "
        "모듈마다 ①개요·예시 ②분석절차단계 ③결과·시각화가 있으며, 데모 12종·정식/쉬운 해석·검증 여과·그래프 PNG/JPEG/PDF 내보내기가 있습니다.",
    )
    add_steps(
        doc,
        [
            "기본통계 홈에서 「논문 통계 기본 지식」→「분석 고르기」를 먼저 봅니다.",
            "기술·추론 또는 SEM 목록에서 기법(예: 독립표본 t검정 C-02)을 고릅니다.",
            "테스트용 데모 12종 중 연구 주제에 가까운 것을 고르거나 CSV를 불러옵니다.",
            "분석절차단계 1~4를 따라간 뒤 분석 실행합니다.",
            "결과에서 정식 결과·쉬운 해석·검증 설명·그래프를 확인하고 내보냅니다.",
        ],
    )

    add_h(doc, "3.10 데이터 시각화", 2)
    add_para(
        doc,
        "(A) 차트 카탈로그 → 유형 상세 → 스튜디오로 그래프를 만듭니다. 등급에 따라 보이는 차트 ID가 달라집니다.",
    )
    add_steps(
        doc,
        [
            "허용된 카테고리에서 차트를 고릅니다(Free: 기본 10).",
            "유형 설명 확인 후 스튜디오에서 데이터를 넣고 그립니다.",
            "기본통계·코딩 결과와 함께 논문용 그림으로 씁니다.",
        ],
    )

    add_h(doc, "3.11 연구방법작성", 2)
    add_para(
        doc,
        "(A) 다른 메뉴에서 만든 산출물(아티팩트)을 끌어와 논문 「연구방법」절 초안을 조립·내보냅니다.",
    )
    add_steps(
        doc,
        [
            "편집기에 제목·본문 골격을 잡습니다.",
            "연동 가능한 메뉴 산출물(기본통계·방법 프로젝트·아키타입 실습 등)을 삽입합니다.",
            "PDF / Excel / HTML로 내보냅니다.",
        ],
    )

    add_h(doc, "3.12 내 프로젝트 · 설정", 2)
    add_bullets(
        doc,
        [
            "내 프로젝트: 로컬에 저장된 방법 프로젝트를 열어 단계 진행률을 이어갑니다.",
            "설정: 테마, AI 제공자·API 키·모델, 모드 A(규칙)/B(AI)/C(하이브리드). AI 기능 사용 전 키 설정을 권장합니다.",
        ],
    )
    doc.add_page_break()

    # ===================== 제4부 RDOS =====================
    add_h(doc, "제4부  연구준비자 플랜 · RDOS 메뉴 절차 매뉴얼", 1)
    add_para(
        doc,
        "연구준비자 플랜은 ‘연구 수행 OS’가 아니라 ‘연구 역량 개발 OS’입니다. "
        "논문 구조·설계·방법 개요·읽기·APA·글쓰기·AI 튜터·지식 코어·로드맵·인증으로 구성됩니다.",
    )

    add_h(doc, "4.1 연구준비자 메뉴 일람과 사용 목적", 2)
    add_table(
        doc,
        ["메뉴", "제공 내용", "권장 사용"],
        [
            ["Dashboard", "학습 현황", "오늘 할 미션·진도 확인"],
            ["Research Basics", "연구 기초", "연구란 무엇인지 개념 잡기"],
            ["논문 구조 학습", "논문 구성 이해", "IMRaD·장 구성 익히기"],
            ["연구설계 기초", "연구문제 작성", "RQ·가설 초안"],
            ["연구방법론 기초", "양적·질적 개요", "방법 가족 입문(심화는 연구자 MethodOS)"],
            ["논문 읽기 훈련", "논문 분석", "요약·비판적 읽기"],
            ["APA 학습", "인용·참고문헌", "인용 형식 연습"],
            ["학술 글쓰기 훈련", "문단 작성", "문단·논리 연결"],
            ["AI 튜터", "질의응답", "막힌 개념 질문"],
            ["지식 코어", "논문가이드·용어", "용어·가이드 검색"],
            ["성장 로드맵", "연구자 성장 단계", "단계별 목표"],
            ["연구 준비자 인증", "Research-Ready", "미션 완료·승급 준비"],
        ],
    )

    add_h(doc, "4.2 연구준비자 절차(권장 학습 패스)", 2)
    add_steps(
        doc,
        [
            "Dashboard에서 진도와 미션을 확인합니다.",
            "Research Basics → 논문 구조 → 연구설계 기초 순으로 ‘질문’을 만듭니다.",
            "연구방법론 기초로 양적/질적 차이를 이해합니다(실제 11단계 워크벤치는 연구자 MethodOS).",
            "논문 읽기·APA·학술 글쓰기로 문장 근력을 키웁니다.",
            "AI 튜터·지식 코어로 막힌 부분을 메웁니다.",
            "로드맵을 점검하고 인증(Research-Ready) 미션을 완료합니다.",
            "조건이 되면 연구자 플랜으로 승급하여 MethodOS 전체 절차를 사용합니다.",
        ],
    )
    callout(
        doc,
        "연구준비자 vs 연구자 혼동 방지",
        "연구준비자 메뉴에는 ‘210개 방법 11단계 워크벤치·기본통계 53모듈·아키타입 실습’이 없습니다. "
        "그런 실행형 기능은 연구자 플랜의 MethodOS(본 매뉴얼 제3부)에 있습니다.",
    )
    doc.add_page_break()

    # ===================== 제5부 Free =====================
    add_h(doc, "제5부  Free 등급 전용 사용 시나리오", 1)
    add_para(
        doc,
        "Free는 ‘연구방법을 맛보고 프로젝트를 저장하는’ 입문 등급입니다. "
        "메뉴는 넓게 보이지만, 깊게 실행할 수 있는 연구방법은 2개, 그래프는 기본 10종입니다.",
    )

    add_h(doc, "5.1 Free에서 할 수 있는 일(권장 하루 루트)", 2)
    add_steps(
        doc,
        [
            "설정에서 테마·(가능하면) AI 키를 넣습니다.",
            "연구방법 추천으로 자기 주제에 가까운 방법을 탐색합니다(실행은 허용 ID만).",
            "연구방법 가상 학습에서 SV-003(설문지 작성) 프로젝트를 시작합니다.",
            "11단계 중 도입~설계 구간을 수행하고 서식 샘플·AI 초안을 저장합니다.",
            "QT-013(사회과학 통계분석2)으로 기초 통계 워크벤치를 한 사이클 돌립니다.",
            "데이터 시각화에서 막대·선·원형 등 기본 10종으로 연습합니다.",
            "기본통계 메뉴로 들어가 개요·데모를 둘러보되, 고급 차트·비허용 방법 실행은 Scholar를 계획합니다.",
            "내 프로젝트에 저장하고 다음날 이어서 엽니다.",
        ],
    )

    add_h(doc, "5.2 Free에서 막히는 지점과 대처", 2)
    add_bullets(
        doc,
        [
            "SV-003/QT-013 외 방법을 열 수 없음 → Basic/Scholar 업그레이드 또는 관리자 허용.",
            "히트맵·박스플롯 등 통계·과학 차트 없음 → Basic 이상.",
            "질적 포커스그룹·심층면접·코딩 매뉴얼 심화 → Basic의 QL-044/041/060.",
        ],
    )
    doc.add_page_break()

    # ===================== 제6부 Basic =====================
    add_h(doc, "제6부  Basic 등급 전용 사용 시나리오", 1)
    add_para(
        doc,
        "Basic은 Free의 설문·통계에 더해 질적 핵심 3방법과 분포·관계 그래프를 엽니다. "
        "소규모 혼합연구 연습에 적합합니다.",
    )

    add_h(doc, "6.1 Basic 권장 2주 프로세스", 2)
    add_steps(
        doc,
        [
            "1~2일: Free와 동일하게 SV-003으로 설문 초안.",
            "3~5일: QL-041 심층면접 절차 가상 학습(면담 가이드·윤리·전사 연습).",
            "6~8일: QL-044 포커스그룹 절차(모더레이션·기록).",
            "9~11일: QL-060 코딩 매뉴얼 + 「코딩」 메뉴 파이프라인으로 코드북·κ 연습.",
            "12~13일: QT-013·기본통계로 양적 요약, 시각화에서 히스토그램·박스·히트맵 등 Basic 차트 사용.",
            "14일: 연구방법작성에 산출물을 모아 방법 절 초안 PDF 내보내기.",
        ],
    )

    add_h(doc, "6.2 Basic 등급 기능 체크리스트", 2)
    add_table(
        doc,
        ["기능", "Free", "Basic"],
        [
            ["연구방법 실행", "2", "5"],
            ["기본 그래프", "○", "○"],
            ["분포·통계·관계 그래프", "×", "○"],
            ["질적 면접·FGI·코딩매뉴얼 워크벤치", "×", "○"],
            ["전체 210방법·전 차트", "×", "× (Scholar)"],
        ],
    )
    doc.add_page_break()

    # ===================== 제7부 Scholar =====================
    add_h(doc, "제7부  Scholar 등급 전용 사용 시나리오(전체 메뉴·전체 기능)", 1)
    add_para(
        doc,
        "Scholar는 MethodOS의 모든 연구방법(약 210)과 모든 시각화 카테고리, "
        "아키타입 31유형·프로젝트 실습, 기본통계 53모듈, 설문·코딩·방법작성을 제한 없이 쓰는 등급입니다. "
        "아래는 ‘논문 한 편을 MethodOS로 관통하는’ 전체 절차입니다.",
    )

    add_h(doc, "7.1 Scholar 전체 연구 파이프라인(절차)", 2)
    add_steps(
        doc,
        [
            "설정: AI 키·하이브리드 모드 설정.",
            "추천: 논문 주제·설계·자료 유형 입력 → 후보 방법 확정.",
            "프로세스·체계: 선택 방법의 단계·계보 확인.",
            "연구방법 가상 학습: 본방법 프로젝트 생성 → 11단계 전 구간 수행(자료수집 시 표준 n·100행+ 데이터 활용).",
            "방법론적 아키타입: 해당 유형의 예비학습 → 실습 주제 5개 중 유사 주제로 태스크 1~N 완주 → 마이크로데이터·AI 산출 확보.",
            "설문조사: 확정 문항지 생성·내보내기.",
            "코딩: 개방형·질적 자료 코드북·파이프라인·보고 문장.",
            "기본통계: 가설에 맞는 기법(예: t/ANOVA/회귀/SEM) 실행 → 검증·그래프·해석 문장.",
            "데이터 시각화: 논문용 고급 차트까지 제작.",
            "연구방법작성: 위 산출물을 삽입해 방법 절·부록 초안 완성.",
            "내 프로젝트·내보내기 파일로 지도교수/공동연구자와 공유.",
        ],
    )

    add_h(doc, "7.2 Scholar에서만 특히 의미 있는 기능", 2)
    add_bullets(
        doc,
        [
            "전체 방법 가족(양적·질적·혼합·조사·역사·실험 등) 워크벤치.",
            "아키타입 31 × 방법별 주제 5개 정교화 샘플과 완료 n 규모 마이크로데이터.",
            "기본통계 기술·추론·SEM 전 모듈과 검증 여과·데모 12종.",
            "시각화 전 카테고리(재무·계층·지도·3D·ML 포함).",
            "기관(University) 확장 시 팀·라이선스 정책과 병행.",
        ],
    )

    add_h(doc, "7.3 Scholar 메뉴×절차 대응표", 2)
    add_table(
        doc,
        ["연구 절차", "MethodOS 메뉴", "핵심 조작"],
        [
            ["문제 설정", "추천 · 프로세스", "체크리스트·후보 확정"],
            ["방법 학습", "가상 학습 · 체계", "11단계 Task 수행"],
            ["유형 연습", "아키타입 예비학습·실습", "태스크 연동·AI 변환"],
            ["도구 제작", "설문 · 코딩", "문항·코드북"],
            ["분석", "기본통계 · (QT-013)", "데모/CSV·해석"],
            ["시각화", "데이터 시각화", "차트 스튜디오"],
            ["원고", "연구방법작성", "아티팩트 삽입·PDF"],
            ["관리", "내 프로젝트 · 설정", "저장·AI키"],
        ],
    )
    doc.add_page_break()

    # ===================== 제8부 =====================
    add_h(doc, "제8부  AI 기능·설정·내보내기·문제 해결", 1)

    add_h(doc, "8.1 AI가 붙는 지점", 2)
    add_bullets(
        doc,
        [
            "연구방법 추천: 구간/최종 해석",
            "11단계 워크벤치: AI 예시 다듬기 · AI 변환기",
            "아키타입 실습: 태스크별 AI 산출물",
            "설문: AI 설문 생성",
            "코딩: 파이프라인·단계 변환기·표준 코드북",
            "기본통계: AI 해석 문장",
        ],
    )
    add_para(
        doc,
        "AI는 초안입니다. 반자동화 항목은 반드시 연구자가 출처·논리·윤리를 확정하십시오. "
        "API 키는 설정 메뉴의 제공자(모델)에 저장합니다.",
    )

    add_h(doc, "8.2 내보내기 형식", 2)
    add_bullets(
        doc,
        [
            "서식 샘플: TXT / DOCX(호환 HTML) / PDF 인쇄",
            "연습·마이크로데이터: CSV",
            "기본통계·시각화: PNG / JPEG / PDF / Excel",
            "연구방법작성: PDF / Excel / HTML",
        ],
    )

    add_h(doc, "8.3 자주 하는 실수", 2)
    add_bullets(
        doc,
        [
            "연구준비자 메뉴에서 통계·210방법 워크벤치를 찾음 → 연구자 MethodOS로 이동.",
            "Free로 고급 방법을 실행하려 함 → 등급 확인.",
            "AI 키 없이 AI 버튼만 반복 → 설정에서 키 등록 또는 규칙 모드 사용.",
            "아키타입 태스크를 건너뛰어 연동 끊김 → 「다음 태스크 진행」 순서 준수.",
            "미니 표(5행)만으로 실전 결론을 냄 → 표준 n·100행+ 데이터로 연습 후 본조사 설계.",
        ],
    )
    doc.add_page_break()

    # ===================== 부록 =====================
    add_h(doc, "부록 A  주요 화면 경로(해시)", 1)
    add_bullets(
        doc,
        [
            "#/dashboard — 대시보드",
            "#/recommend — 연구방법 추천",
            "#/process · /summary · /stages · /common — 프로세스",
            "#/methods · #/method/{ID} · #/project/{pid} — 가상 학습·워크벤치",
            "#/archetype · /overview · /glossary · /projects · /practice/... — 아키타입",
            "#/survey · #/coding · #/basic-stats · #/viz · #/method-writing · #/projects · #/settings",
        ],
    )

    add_h(doc, "부록 B  빠른 체크리스트(연구자 Scholar)", 1)
    add_bullets(
        doc,
        [
            "□ AI 키·모드 설정",
            "□ 추천으로 방법 후보 확정",
            "□ 가상 학습 11단계 핵심 Task 완료",
            "□ 아키타입 실습 주제 1개 완주 + CSV",
            "□ 설문/코딩/기본통계/시각화 산출",
            "□ 연구방법작성 PDF",
            "□ 내 프로젝트 저장 확인",
        ],
    )

    add_h(doc, "부록 C  문서 정보", 1)
    add_para(
        doc,
        "본 문서는 methodos--s23-all-23 코드(메뉴·권한·워크벤치)와 "
        "AI-Research-OS/RDOS 듀얼 트랙 사양(연구준비자·연구자)을 교차 정리한 출판용 사용자 매뉴얼입니다. "
        "화면 문구는 빌드에 따라 일부 달라질 수 있으며, 등급 명칭 Free/Basic/Scholar는 제품 권한 체계를 따릅니다.",
        size=10,
    )
    add_para(
        doc,
        "— 본편 끝 —",
        size=11,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=12,
        space_after=18,
    )

    doc.add_page_break()
    add_h(doc, "보충편 1  연구방법 가상 학습 · 화면 요소 상세", 1)
    add_para(
        doc,
        "사용자가 한 Task 분절을 열었을 때 위에서 아래로 만나는 요소를 모두 설명합니다.",
    )
    add_bullets(
        doc,
        [
            "쉬운 예시 문장: 오늘 할 일을 생활 언어로 요약한 한 줄.",
            "수행/산출 패널: 할 일 설명, 자동화 배지(수동·반자동·자동), AI→연구자 확정 안내.",
            "왜 이 과정이 필요한가: 초보자가 다음 단계로 넘어가기 전 고정해야 할 기준.",
            "구분 근거: 왜 반자동/자동인지(학술 확정은 연구자 몫 등).",
            "따라 할 단계: 체크리스트형 절차(초안→대조→재작성→확정).",
            "전문가용 원문: 분석표·영문 약어가 필요할 때 펼침.",
            "산출물 서식 샘플: 연습용으로 채워진 표·서명란. TXT/DOCX/PDF 저장.",
            "AI 예시: 미니표 + 표준 표본수 아코디언 + 100행+ CSV.",
            "AI 변환기: 내 프로젝트 문구 생성. 자료 연동 모드 선택 필수.",
            "실행 결과 영역: 요약·표·차트·다형식 내보내기.",
        ],
    )

    add_h(doc, "보충편 2  아키타입 실습 · 자료수집 태스크에서 볼 것", 1)
    add_steps(
        doc,
        [
            "모집단·표본틀·표집방법·목표 n이 표에 있는지 확인합니다.",
            "측정도구 태스크에서 문항코드(C01…)가 데이터 열 이름과 같게 유지되는지 봅니다.",
            "본조사 태스크에서 완료 n·응답률·평균 소요를 읽고, 아코디언 마이크로데이터(100행+)를 펼칩니다.",
            "CSV를 저장해 기본통계 메뉴에 불러와 같은 주제 분석을 연습합니다.",
            "AI 변환기 기입란에 ‘내 주제’로 바꿔 쓴 뒤 산출물을 Excel/PDF로 남깁니다.",
        ],
    )

    add_h(doc, "보충편 3  Free / Basic / Scholar 하루 일과 비교", 1)
    add_table(
        doc,
        ["시간대", "Free", "Basic", "Scholar"],
        [
            ["오전", "SV-003 서식·AI 초안", "QL-041 면담 가이드", "추천→본방법 11단계"],
            ["오후", "QT-013 기초 통계", "QL-060+코딩 파이프라인", "아키타입 실습+기본통계"],
            ["저녁", "기본 그래프 10종", "히스토그램·박스·히트맵", "방법작성 PDF+공유"],
        ],
    )

    add_h(doc, "보충편 4  연구준비자→연구자 승급 후 첫 주", 1)
    add_para(
        doc,
        "RDOS에서 Research-Ready 인증을 마친 뒤 연구자 플랜으로 오면, 첫 주에는 "
        "MethodOS 「연구방법 추천」→「가상 학습」허용(또는 Scholar 전체) 방법 1개→「기본통계」데모 1개만 "
        "완주하는 것을 권합니다. 한꺼번에 31 아키타입을 열지 마십시오.",
    )

    add_h(doc, "보충편 5  FAQ", 1)
    faqs = [
        (
            "Q. methodos--s23-all-23에 ‘연구준비자 메뉴’가 안 보입니다.",
            "A. 연구준비자 메뉴는 RDOS 트랙(/rdos)에 있습니다. MethodOS 빌드는 연구자용 연구방법 OS입니다. 듀얼 플랜 랜딩에서 트랙을 고르십시오.",
        ),
        (
            "Q. Free인데 사이드바에 코딩·아키타입이 보입니다. 다 쓸 수 있나요?",
            "A. 메뉴 진입은 가능하나, 연구방법 실행·차트는 등급 허용 범위로 제한됩니다. 잠김 안내가 뜨면 업그레이드가 필요합니다.",
        ),
        (
            "Q. 연습 데이터가 5행밖에 없는데 논문에 써도 되나요?",
            "A. 안 됩니다. 미니표는 서식용입니다. 표준 n 안내와 100행+ 마이크로데이터로 연습하고, 본조사는 설계된 표본수로 수집하십시오.",
        ),
        (
            "Q. Scholar와 University 차이는?",
            "A. Scholar는 개인 연구자 전체 기능, University는 기관·팀·라이선스 정책이 더해지는 상위/기관 플랜입니다. 방법·차트 기본 해제는 동일하게 ‘전체’입니다.",
        ),
        (
            "Q. AI 결과가 틀리면?",
            "A. 반자동화 원칙상 연구자 확정이 필수입니다. 원자료·수식·검증 여과·문헌과 대조하고, 설정에서 모델/온도를 조정하십시오.",
        ),
    ]
    for q, a in faqs:
        add_para(doc, q, size=10.5, bold=True, space_after=2)
        add_para(doc, a, size=10, space_after=8)

    add_para(
        doc,
        "— 끝 —",
        size=11,
        bold=True,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=24,
    )

    doc.save(OUT)
    print("Wrote", OUT, "size", OUT.stat().st_size)


if __name__ == "__main__":
    build()
