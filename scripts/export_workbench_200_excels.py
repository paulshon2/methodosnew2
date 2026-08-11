from __future__ import annotations

import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook


ROOT = Path(__file__).resolve().parents[1]
METHODS_DIR = ROOT / "public" / "methods"
REGISTRY_PATH = ROOT / "public" / "registry.json"
OUT_DIR = ROOT / "exports" / "workbench_200_excels"


def safe_filename(value: str) -> str:
    banned = '<>:"/\\|?*'
    out = "".join("_" if c in banned else c for c in value)
    return out.strip().strip(".")


def load_registry_map() -> dict[str, dict[str, Any]]:
    registry = json.loads(REGISTRY_PATH.read_text(encoding="utf-8"))
    methods = registry.get("methods", [])
    return {m.get("id", ""): m for m in methods if m.get("id")}


def get_entries(task: dict[str, Any]) -> list[dict[str, Any]]:
    entries: list[dict[str, Any]] = []
    for item in task.get("items") or []:
        for entry in item.get("entries") or []:
            entries.append(entry)
    return entries


def make_workbook(method_id: str, stages: list[dict[str, Any]], meta: dict[str, Any]) -> Workbook:
    wb = Workbook()
    ws_overview = wb.active
    ws_overview.title = "개요"
    ws_stages = wb.create_sheet("11단계")
    ws_tasks = wb.create_sheet("세부태스크")
    ws_entries = wb.create_sheet("세부항목")

    stage_count = len(stages)
    task_count = sum(len(s.get("tasks") or []) for s in stages)
    entry_count = 0
    for stage in stages:
        for task in stage.get("tasks") or []:
            entry_count += len(get_entries(task))

    ws_overview.append(["항목", "값"])
    ws_overview.append(["방법 ID", method_id])
    ws_overview.append(["방법명(EN)", meta.get("name", "")])
    ws_overview.append(["방법명(KO)", meta.get("name_ko", "")])
    ws_overview.append(["계열(Family)", meta.get("family", "")])
    ws_overview.append(["목적(Purpose)", meta.get("purpose", "")])
    ws_overview.append(["철학/패러다임", (meta.get("philosophy") or {}).get("paradigm", "")])
    ws_overview.append(["총 단계 수", stage_count])
    ws_overview.append(["총 세부태스크 수", task_count])
    ws_overview.append(["총 세부항목(수행내용/산출물) 수", entry_count])

    ws_stages.append(
        [
            "방법 ID",
            "단계 번호",
            "단계 제목",
            "단계 전체명",
            "단계 라벨",
            "easyIntro",
            "whyStage",
            "inEasy",
            "outEasy",
        ]
    )
    for stage in stages:
        ws_stages.append(
            [
                method_id,
                stage.get("stageNo", ""),
                stage.get("title", ""),
                stage.get("fullName", ""),
                stage.get("stageLabel", ""),
                stage.get("easyIntro", ""),
                stage.get("whyStage", ""),
                stage.get("inEasy", ""),
                stage.get("outEasy", ""),
            ]
        )

    ws_tasks.append(
        [
            "방법 ID",
            "단계 번호",
            "태스크 번호",
            "태스크명",
            "태스크명(KO)",
            "example",
            "whyTask",
            "verify",
        ]
    )
    for stage in stages:
        stage_no = stage.get("stageNo", "")
        for task in stage.get("tasks") or []:
            ws_tasks.append(
                [
                    method_id,
                    stage_no,
                    task.get("taskNo", ""),
                    task.get("name", ""),
                    task.get("nameKo", ""),
                    task.get("example", ""),
                    task.get("whyTask", ""),
                    task.get("verify", ""),
                ]
            )

    ws_entries.append(
        [
            "방법 ID",
            "단계 번호",
            "태스크 번호",
            "엔트리 ID",
            "종류(kind)",
            "세그먼트(seg)",
            "표시명(segLabel)",
            "내용(content)",
            "자동화 수준(auto)",
            "근거(reason)",
            "입출력(io)",
            "동작(behavior)",
            "프로시저",
        ]
    )
    for stage in stages:
        stage_no = stage.get("stageNo", "")
        for task in stage.get("tasks") or []:
            task_no = task.get("taskNo", "")
            for entry in get_entries(task):
                ws_entries.append(
                    [
                        method_id,
                        stage_no,
                        task_no,
                        entry.get("id", ""),
                        entry.get("kind", ""),
                        entry.get("seg", ""),
                        entry.get("segLabel", ""),
                        entry.get("content", ""),
                        entry.get("auto", ""),
                        entry.get("reason", ""),
                        entry.get("io", ""),
                        entry.get("behavior", ""),
                        " | ".join(entry.get("procedure") or []),
                    ]
                )

    return wb


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    reg_map = load_registry_map()
    perf_paths = sorted(METHODS_DIR.glob("*/perf-registry.json"))

    made = 0
    for perf_path in perf_paths:
        method_id = perf_path.parent.name
        stages = json.loads(perf_path.read_text(encoding="utf-8"))
        if not isinstance(stages, list):
            continue
        meta = reg_map.get(method_id, {})
        name_part = meta.get("name_ko") or meta.get("name") or method_id
        filename = f"{method_id}_{safe_filename(name_part)}_11단계워크벤치.xlsx"
        out_path = OUT_DIR / filename
        wb = make_workbook(method_id, stages, meta)
        wb.save(out_path)
        made += 1

    print(f"exported_files={made}")
    print(f"output_dir={OUT_DIR}")


if __name__ == "__main__":
    main()
