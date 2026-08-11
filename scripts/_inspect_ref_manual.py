# -*- coding: utf-8 -*-
from docx import Document
from pathlib import Path

p = Path(r"c:\Users\saran\OneDrive\Desktop\AI-Research-OS_RDOS_통합_사용자매뉴얼.docx")
print("exists", p.exists(), "size", p.stat().st_size if p.exists() else None)
d = Document(str(p))
print("paragraphs", len(d.paragraphs), "tables", len(d.tables))
heads = []
for para in d.paragraphs:
    style = para.style.name if para.style else ""
    t = para.text.strip()
    if not t:
        continue
    if style.startswith("Heading"):
        heads.append((style, t[:120]))
print("headings", len(heads))
for s, t in heads[:100]:
    print(f"[{s}] {t}")
print("--- body sample ---")
n = 0
for para in d.paragraphs:
    if para.text.strip():
        n += 1
        if 20 <= n <= 45:
            print(para.text[:220].replace("\n", " / "))
print("nonempty", sum(1 for p in d.paragraphs if p.text.strip()))
if d.tables:
    t = d.tables[0]
    print("table0 rows", len(t.rows), "cols", len(t.columns))
    for r in t.rows[:3]:
        print([c.text[:50].replace("\n", " ") for c in r.cells])
