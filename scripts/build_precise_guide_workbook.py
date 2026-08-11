# -*- coding: utf-8 -*-
"""31묶음 정밀가이드 시트 생성 — 로드맵·1번 시트·번들 JSON 병합 후 엑셀·public 복제."""
from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import openpyxl

SCRIPT_DIR = Path(__file__).resolve().parent
SAAS = SCRIPT_DIR.parent
PUBLIC = SAAS / "public"
BUNDLES = PUBLIC / "archetypes" / "bundles"
sys.path.insert(0, str(SCRIPT_DIR))

from excel_precise_guide import (  # noqa: E402
    DESK,
    GUIDE_NAME,
    clean_task,
    ensure_workbook_with_31_sheets,
    load_roadmap,
    norm_task_key,
    parse_bundle_sheet,
    resolve_guide_path,
)
from phd_task_examples import ARCHETYPE_DOMAIN, generate_phd_gallery  # noqa: E402

# gen helpers
from gen_archetype_data import (  # noqa: E402
    clean_process_label,
    gen_how,
    gen_protocol,
    gen_why,
    process_key_from_label,
)


def detect_kind_simple(task: str) -> str:
    from enrich_archetype_tasks import detect_kind

    return detect_kind(task)


def detail_from_bundle_task(bundle_id: int, task: dict, archetype_name: str, process: str, stage: str) -> dict[str, str]:
    raw = task.get("taskRaw") or clean_task(task.get("task", ""))
    pk = task.get("processKey") or process_key_from_label(process)
    why = (task.get("why") or "").strip() or gen_why(raw, process, archetype_name)
    how = (task.get("how") or "").strip() or gen_how(raw, stage)
    protocol = (task.get("protocol") or "").strip() or gen_protocol(raw, pk)
    sample = (task.get("excelSample") or "").strip()
    if not sample:
        for s in task.get("samples") or []:
            if s.get("example"):
                sample = str(s["example"]).strip()
                break
    if not sample:
        gal = generate_phd_gallery(bundle_id, pk, detect_kind_simple(raw), raw, "")
        for g in gal:
            if g.get("source") != "excel" and len(g.get("body", "")) > 80:
                sample = g["body"][:1200]
                break
    explain = (task.get("sampleExplain") or "").strip()
    if not explain and sample:
        explain = (
            f"{ARCHETYPE_DOMAIN.get(bundle_id, ARCHETYPE_DOMAIN[31])[0]} 맥락에서 "
            f"'{clean_process_label(raw)}' 태스크는 다음 단계로 넘기기 전 품질·재현성·감사기록을 확보하기 위한 예시입니다."
        )
    return {
        "why": why,
        "how": how,
        "protocol": protocol,
        "sample": sample,
        "explain": explain,
    }


def load_existing_bundle1(path: Path) -> dict[str, dict[str, str]]:
    if not path.is_file():
        return {}
    wb = openpyxl.load_workbook(path, data_only=True)
    if "1. 1번묶음 정밀가이드" not in wb.sheetnames:
        wb.close()
        return {}
    rows = parse_bundle_sheet(wb["1. 1번묶음 정밀가이드"])
    wb.close()
    out = {}
    for row in rows:
        k = norm_task_key(row["task"])
        out[k] = {
            "why": row["why"],
            "how": row["how"],
            "protocol": row["protocol"],
            "sample": row["sample"],
            "explain": row["explain"],
        }
    return out


def collect_bundle_tasks_from_json() -> dict[int, list[dict]]:
    by_bid: dict[int, list[dict]] = {}
    for path in sorted(BUNDLES.glob("bundle-*.json")):
        data = json.loads(path.read_text(encoding="utf-8"))
        bid = int(data.get("id") or re.search(r"(\d+)", path.stem).group(1))
        archetype = data.get("name_ko") or data.get("name") or f"묶음{bid}"
        items: list[dict] = []
        for stage in data.get("stages", []):
            proc_label = stage.get("name_ko") or stage.get("processLabel") or ""
            stage_chain = stage.get("name_ko") or stage.get("name") or ""
            for task in stage.get("tasks", []):
                proc = proc_label
                if "자료수집" in str(stage.get("processKey", "")):
                    proc = "① 자료수집"
                elif stage.get("processKey") == "preprocess":
                    proc = "② 전처리"
                elif stage.get("processKey") == "coding":
                    proc = "③ 코딩"
                elif stage.get("processKey") == "stats":
                    proc = "④ 통계기법"
                elif stage.get("processKey") == "result":
                    proc = "⑤ 결과도출"
                elif stage.get("processKey") == "deliverable":
                    proc = "⑥ 최종 산출물"
                detail = detail_from_bundle_task(bid, task, archetype, proc, stage_chain)
                items.append(
                    {
                        "process": proc,
                        "stage": stage_chain,
                        "task": clean_task(task.get("taskRaw") or task.get("task", "")),
                        **detail,
                    }
                )
        by_bid[bid] = items
    return by_bid


def merge_roadmap_order(roadmap: dict[int, list[dict]], by_json: dict[int, list[dict]], b1_excel: dict[str, dict]) -> dict[int, list[dict]]:
    """Prefer roadmap task order; fill detail from json/excel."""
    json_keyed: dict[int, dict[str, dict]] = {}
    for bid, rows in by_json.items():
        json_keyed[bid] = {norm_task_key(r["task"]): r for r in rows}

    out: dict[int, list[dict]] = {}
    for bid in range(1, 32):
        ordered: list[dict] = []
        roadmap_rows = roadmap.get(bid, [])
        jmap = json_keyed.get(bid, {})
        if roadmap_rows:
            for rr in roadmap_rows:
                task = clean_task(rr["task"])
                key = norm_task_key(task)
                base = jmap.get(key, {})
                det = {
                    "why": base.get("why", ""),
                    "how": base.get("how", ""),
                    "protocol": base.get("protocol", ""),
                    "sample": base.get("sample", ""),
                    "explain": base.get("explain", ""),
                }
                if bid == 1 and key in b1_excel:
                    det = {**det, **{k: v for k, v in b1_excel[key].items() if v}}
                if not det["why"]:
                    det = detail_from_bundle_task(
                        bid,
                        {"task": task, **base},
                        ARCHETYPE_DOMAIN.get(bid, ARCHETYPE_DOMAIN[31])[0],
                        rr.get("process", ""),
                        rr.get("stage", ""),
                    )
                ordered.append(
                    {
                        "process": rr.get("process") or base.get("process", ""),
                        "stage": rr.get("stage") or base.get("stage", ""),
                        "task": task,
                        **det,
                    }
                )
        else:
            ordered = list(jmap.values())
        out[bid] = ordered
    return out


def main():
    src = DESK / GUIDE_NAME
    public_copy = PUBLIC / "archetypes" / "precise-guide-31bundles.xlsx"
    desk_out = DESK / "연구방법_과정별_태스크_정밀실행가이드_31묶음.xlsx"

    roadmap: dict[int, list] = {}
    if src.is_file():
        wb = openpyxl.load_workbook(src, data_only=True)
        roadmap = load_roadmap(wb)
        wb.close()
    print("roadmap bundles", len(roadmap), "tasks", sum(len(v) for v in roadmap.values()))

    b1_excel = load_existing_bundle1(src)
    print("bundle1 excel rows", len(b1_excel))
    by_json = collect_bundle_tasks_from_json()
    merged = merge_roadmap_order(roadmap, by_json, b1_excel)

    n = ensure_workbook_with_31_sheets(src if src.is_file() else public_copy, public_copy, merged)
    print("wrote public", public_copy, "rows", n)
    if src.is_file():
        n2 = ensure_workbook_with_31_sheets(src, desk_out, merged)
        print("wrote desktop", desk_out, "rows", n2)
    print("done")


if __name__ == "__main__":
    main()
