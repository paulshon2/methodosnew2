# -*- coding: utf-8 -*-
"""정밀 실행 가이드 엑셀 — 31묶음 시트 읽기/쓰기·키 정규화."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl import Workbook

DESK = Path(r"c:\Users\saran\OneDrive\Desktop")
GUIDE_NAME = "연구방법_과정별_태스크_정밀실행가이드.xlsx"
HEADER = (
    "과정",
    "절차적단계(원문)",
    "태스크(단계)",
    "① 왜 하는지",
    "② 어떻게 하는지",
    "③ 표준 프로토콜(준비)",
    "④ 예시 샘플 (실제 테스트용)",
    "⑤ 예시 샘플 해설 (왜 이렇게)",
)


def clean_task(s: str) -> str:
    return re.sub(r"^\d+\)\s*", "", (s or "").strip())


def norm_task_key(task: str) -> str:
    t = clean_task(task)
    t = re.sub(r"\(Procedural Task\)", "", t, flags=re.I)
    return re.sub(r"[^\w가-힣]+", "", t.lower())


def guide_paths(public_dir: Path) -> list[Path]:
    return [
        public_dir / "archetypes" / "precise-guide-31bundles.xlsx",
        DESK / "연구방법_과정별_태스크_정밀실행가이드_31묶음.xlsx",
        DESK / GUIDE_NAME,
        public_dir / "archetypes" / GUIDE_NAME,
    ]


def resolve_guide_path(public_dir: Path) -> Path:
    for p in guide_paths(public_dir):
        if p.is_file():
            return p
    return DESK / GUIDE_NAME


def sheet_name_for_bundle(bundle_id: int) -> str:
    if bundle_id == 1:
        return "1. 1번묶음 정밀가이드"
    return f"{bundle_id + 2}. {bundle_id}번묶음 정밀가이드"


def is_bundle_guide_sheet(name: str) -> bool:
    return bool(re.search(r"\d+번묶음\s*정밀가이드", name or ""))


def bundle_id_from_sheet(name: str) -> int | None:
    m = re.search(r"(\d+)번묶음\s*정밀가이드", name or "")
    return int(m.group(1)) if m else None


def parse_bundle_sheet(ws) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    cur_process = ""
    cur_stage = ""
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if r[0]:
            cur_process = str(r[0]).strip()
        if len(r) > 1 and r[1]:
            cur_stage = str(r[1]).strip()
        if len(r) < 3 or not r[2]:
            continue
        task = str(r[2]).strip()
        rows.append(
            {
                "process": cur_process,
                "stage": cur_stage,
                "task": task,
                "why": str(r[3] or "").strip() if len(r) > 3 else "",
                "how": str(r[4] or "").strip() if len(r) > 4 else "",
                "protocol": str(r[5] or "").strip() if len(r) > 5 else "",
                "sample": str(r[6] or "").strip() if len(r) > 6 else "",
                "explain": str(r[7] or "").strip() if len(r) > 7 else "",
            }
        )
    return rows


def load_roadmap(wb) -> dict[int, list[dict[str, Any]]]:
    if "2. 전체 태스크 로드맵" not in wb.sheetnames:
        return {}
    ws = wb["2. 전체 태스크 로드맵"]
    out: dict[int, list[dict[str, Any]]] = {}
    cur_b = None
    cur_proc = ""
    cur_stage = ""
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if r[0] is not None and str(r[0]).strip() != "":
            try:
                cur_b = int(r[0])
            except (TypeError, ValueError):
                pass
        if len(r) > 1 and r[1]:
            cur_proc = str(r[1]).strip()
        if len(r) > 2 and r[2]:
            cur_stage = str(r[2]).strip()
        if cur_b is None or len(r) < 4 or not r[3]:
            continue
        out.setdefault(cur_b, []).append(
            {
                "process": cur_proc,
                "stage": cur_stage,
                "task": str(r[3]).strip(),
                "status": str(r[4] or "").strip() if len(r) > 4 else "",
            }
        )
    return out


def load_all_bundle_details(public_dir: Path) -> dict[tuple[int, str], dict[str, str]]:
    """Load ①–⑤ for every task keyed by (bundle_id, norm_task_key)."""
    path = resolve_guide_path(public_dir)
    if not path.is_file():
        return {}

    wb = openpyxl.load_workbook(path, data_only=True)
    merged: dict[tuple[int, str], dict[str, str]] = {}

    legacy = "1. 1번묶음 정밀가이드"
    for name in wb.sheetnames:
        bid = bundle_id_from_sheet(name)
        if bid is None:
            continue
        rows = parse_bundle_sheet(wb[name])
        for row in rows:
            key = norm_task_key(row["task"])
            if not key:
                continue
            merged[(bid, key)] = {
                "why": row["why"],
                "how": row["how"],
                "protocol": row["protocol"],
                "sample": row["sample"],
                "explain": row["explain"],
            }

    wb.close()
    return merged


def write_row(ws, row_idx: int, process: str, stage: str, task: str, detail: dict[str, str]):
    ws.cell(row=row_idx, column=1, value=process or None)
    ws.cell(row=row_idx, column=2, value=stage or None)
    ws.cell(row=row_idx, column=3, value=task)
    ws.cell(row=row_idx, column=4, value=detail.get("why", ""))
    ws.cell(row=row_idx, column=5, value=detail.get("how", ""))
    ws.cell(row=row_idx, column=6, value=detail.get("protocol", ""))
    ws.cell(row=row_idx, column=7, value=detail.get("sample", ""))
    ws.cell(row=row_idx, column=8, value=detail.get("explain", ""))


def ensure_workbook_with_31_sheets(
    source_path: Path,
    dest_path: Path,
    bundle_task_details: dict[int, list[dict[str, Any]]],
) -> int:
    """
    Create/update workbook: keep 사용안내 + 로드맵, add sheets 1..31 with full ①–⑤.
    bundle_task_details[bid] = list of {process, stage, task, why, how, protocol, sample, explain}
  Returns row count written.
    """
    if source_path.is_file():
        wb = openpyxl.load_workbook(source_path)
    else:
        wb = Workbook()
        if "Sheet" in wb.sheetnames:
            del wb["Sheet"]

    if "0. 사용안내" not in wb.sheetnames:
        wb.create_sheet("0. 사용안내", 0)

    written = 0
    for bid in range(1, 32):
        sname = sheet_name_for_bundle(bid)
        if sname in wb.sheetnames:
            del wb[sname]
        ws = wb.create_sheet(sname)
        for col, h in enumerate(HEADER, 1):
            ws.cell(row=1, column=col, value=h)
        tasks = bundle_task_details.get(bid, [])
        row_i = 2
        last_proc = ""
        last_stage = ""
        for t in tasks:
            proc = t.get("process") or last_proc
            stage = t.get("stage") or last_stage
            if t.get("process"):
                last_proc = proc
            if t.get("stage"):
                last_stage = stage
            write_row(ws, row_i, proc, stage, t["task"], t)
            row_i += 1
            written += 1

    dest_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dest_path)
    wb.close()
    return written
