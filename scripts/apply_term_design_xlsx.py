# -*- coding: utf-8 -*-
"""
methodos 용어치환 설계서(xlsx) 적용 · 무결성 검증

사용:
  python scripts/apply_term_design_xlsx.py --dry-run
  python scripts/apply_term_design_xlsx.py --apply
  python scripts/apply_term_design_xlsx.py --verify-only
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
from pathlib import Path
from typing import Any

try:
    import openpyxl
except ImportError:
    print("openpyxl required: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parents[1]
PUBLIC = ROOT / "public"
METHODS = PUBLIC / "methods"
DEFAULT_XLSX = Path(
    r"c:\Users\saran\OneDrive\Desktop\methodos--s23-all-11_용어치환_설계서-수정안.xlsx"
)

SKIP_DIRS = {".git", "node_modules", ".next", "dist", "build"}
TEXT_EXTS = {
    ".ts", ".tsx", ".js", ".jsx", ".json", ".md", ".txt", ".py", ".css", ".html",
    ".webmanifest", ".mjs", ".csv",
}


def load_slug_rows(xlsx: Path) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    ws = wb["01_슬러그_치환표_200종"]
    rows: list[dict[str, Any]] = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        rows.append({
            "new_slug": str(r[0]).strip(),
            "old_perf": str(r[1]).strip(),
            "old_method_id": str(r[2]).strip(),
            "new_method_id": str(r[3]).strip(),
            "new_name_ko": str(r[5] or "").strip(),
            "new_name_hint": str(r[7] or "").strip(),
            "legacy_perf": (str(r[11]).strip() if r[11] else None) or None,
        })
    wb.close()
    return rows


def load_safe_replacements(xlsx: Path) -> list[tuple[str, str]]:
    """시트 06(출판사) + 04 안전 + 07 저서명 — 본문 치환용."""
    out: list[tuple[str, str]] = []
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)

    def add_sheet(name: str, old_i: int, new_i: int, grade_i: int | None = None):
        ws = wb[name]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not r or not r[old_i]:
                continue
            if grade_i is not None and r[grade_i] != "안전":
                continue
            old, new = str(r[old_i]).strip(), str(r[new_i]).strip()
            if not old or not new or new.startswith("(삭제)"):
                new = ""
            if old:
                out.append((old, new))

    add_sheet("06_출판사_치환표", 0, 1, 4)
    add_sheet("04_데이터셋_측정도구_치환표", 0, 1, 4)
    # 07: 한글 표기 -> 치환 후 표기
    ws7 = wb["07_콘텐츠내_저서명_치환표"]
    for r in ws7.iter_rows(min_row=2, values_only=True):
        if not r or not r[0]:
            continue
        if len(r) > 5 and r[5] != "안전":
            continue
        old, new = str(r[0]).strip(), str(r[3]).strip()
        if old and new:
            out.append((old, new))

    wb.close()
    # longer strings first
    out.sort(key=lambda x: -len(x[0]))
    return out


def build_maps(rows: list[dict[str, Any]]):
    perf_to_new: dict[str, str] = {}
    mid_to_new: dict[str, str] = {}
    meta_by_old_perf: dict[str, dict] = {}
    for row in rows:
        perf_to_new[row["old_perf"]] = row["new_slug"]
        mid_to_new[row["old_method_id"]] = row["new_method_id"]
        meta_by_old_perf[row["old_perf"]] = row
    return perf_to_new, mid_to_new, meta_by_old_perf


def dry_run(rows: list[dict[str, Any]]) -> int:
    errors: list[str] = []
    for row in rows:
        d = METHODS / row["old_perf"]
        if not d.is_dir():
            errors.append(f"MISSING_DIR {row['old_perf']}")
        nd = METHODS / row["new_slug"]
        if nd.exists():
            errors.append(f"COLLISION {row['new_slug']}")
    cat = json.loads((PUBLIC / "method-catalog.json").read_text(encoding="utf-8"))
    methods = cat["methods"]
    if len(methods) != 200:
        errors.append(f"CATALOG_COUNT {len(methods)}")
    perf_in_cat = {m["perfSlug"] for m in methods}
    for row in rows:
        if row["old_perf"] not in perf_in_cat:
            errors.append(f"CATALOG_MISSING {row['old_perf']}")
    print(f"[dry-run] slug rules: {len(rows)}")
    print(f"[dry-run] method dirs: {len(list(METHODS.iterdir()))}")
    if errors:
        print(f"[dry-run] ERRORS: {len(errors)}")
        for e in errors[:30]:
            print(" ", e)
        return 1
    print("[dry-run] OK — 구조 전제 충족")
    return 0


def rename_method_dirs(rows: list[dict[str, Any]], apply: bool) -> None:
    for row in rows:
        src = METHODS / row["old_perf"]
        dst = METHODS / row["new_slug"]
        if not src.is_dir():
            raise FileNotFoundError(src)
        if dst.exists():
            raise FileExistsError(dst)
        if apply:
            src.rename(dst)
        print(f"  mv {row['old_perf']} -> {row['new_slug']}")


def patch_catalog(rows: list[dict[str, Any]], apply: bool) -> None:
    path = PUBLIC / "method-catalog.json"
    cat = json.loads(path.read_text(encoding="utf-8"))
    by_perf = {r["old_perf"]: r for r in rows}
    for m in cat["methods"]:
        r = by_perf.get(m.get("perfSlug"))
        if not r:
            continue
        m["methodId"] = r["new_method_id"]
        m["perfSlug"] = r["new_slug"]
        if r["legacy_perf"]:
            m["legacyPerf"] = r["legacy_perf"]
        if r["new_name_ko"]:
            m["name_ko"] = r["new_name_ko"]
        if r["new_name_hint"]:
            m["name_hint"] = r["new_name_hint"]
        if m.get("xlsxFile"):
            old_x = m["xlsxFile"]
            if r["old_perf"] in old_x:
                m["xlsxFile"] = old_x.replace(r["old_perf"], r["new_slug"])
        for key in ("purpose",):
            if isinstance(m.get(key), str) and r["old_perf"] in m[key]:
                m[key] = m[key].replace(r["old_perf"], r["new_slug"])
    if apply:
        path.write_text(json.dumps(cat, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"  patched method-catalog.json ({len(cat['methods'])} methods)")


def patch_perf_index(rows: list[dict[str, Any]], apply: bool) -> None:
    path = PUBLIC / "perf-index.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    new_data: dict[str, Any] = {}
    mid_map = {r["old_method_id"]: r for r in rows}
    for old_key, val in data.items():
        r = mid_map.get(old_key)
        if r:
            new_key = r["new_method_id"]
            if isinstance(val, dict):
                val = dict(val)
                val["perfSlug"] = r["new_slug"]
                if r["legacy_perf"]:
                    val["legacyPerf"] = r["legacy_perf"]
            new_data[new_key] = val
        else:
            new_data[old_key] = val
    if apply:
        path.write_text(json.dumps(new_data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("  patched perf-index.json")


def replace_in_obj(obj: Any, perf_map: dict[str, str], mid_map: dict[str, str]) -> Any:
    if isinstance(obj, str):
        s = obj
        for old, new in perf_map.items():
            if old in s:
                s = s.replace(old, new)
        for old, new in mid_map.items():
            if old in s:
                s = s.replace(old, new)
        return s
    if isinstance(obj, list):
        return [replace_in_obj(x, perf_map, mid_map) for x in obj]
    if isinstance(obj, dict):
        out = {}
        for k, v in obj.items():
            nk = k
            if k in mid_map:
                nk = mid_map[k]
            elif k in perf_map:
                nk = perf_map[k]
            out[nk] = replace_in_obj(v, perf_map, mid_map)
        return out
    return obj


def patch_method_process_guide(perf_map: dict[str, str], mid_map: dict[str, str], apply: bool) -> None:
    path = PUBLIC / "method-process-guide.json"
    data = json.loads(path.read_text(encoding="utf-8"))
    data = replace_in_obj(data, perf_map, mid_map)
    if apply:
        path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("  patched method-process-guide.json")


def patch_registry(mid_map: dict[str, str], apply: bool) -> None:
    path = PUBLIC / "registry.json"
    text = path.read_text(encoding="utf-8")
    original = text
    for old, new in sorted(mid_map.items(), key=lambda x: -len(x[0])):
        text = text.replace(f'"{old}.', f'"{new}.')
        text = text.replace(f'"id": "{old}"', f'"id": "{new}"')
    if apply and text != original:
        path.write_text(text, encoding="utf-8")
    print("  patched registry.json (method ids)")


def patch_each_method_folder(perf_map: dict[str, str], mid_map: dict[str, str], apply: bool) -> None:
    count = 0
    for d in METHODS.iterdir():
        if not d.is_dir():
            continue
        for fp in d.rglob("*"):
            if fp.suffix.lower() not in {".json", ".md", ".txt", ".csv"}:
                continue
            try:
                raw = fp.read_text(encoding="utf-8")
            except Exception:
                continue
            new_raw = raw
            for old, new in perf_map.items():
                new_raw = new_raw.replace(old, new)
            for old, new in mid_map.items():
                new_raw = new_raw.replace(old, new)
            if new_raw != raw and apply:
                fp.write_text(new_raw, encoding="utf-8")
                count += 1
    print(f"  patched files under public/methods/: {count}")


def patch_repo_text_files(
    perf_map: dict[str, str],
    mid_map: dict[str, str],
    safe_pairs: list[tuple[str, str]],
    apply: bool,
) -> None:
    changed = 0
    for base in (ROOT, ROOT.parent):
        if not base.exists():
            continue
        for fp in base.rglob("*"):
            if any(p in SKIP_DIRS for p in fp.parts):
                continue
            if not fp.is_file() or fp.suffix.lower() not in TEXT_EXTS:
                continue
            if "apply_term_design_xlsx.py" in fp.name:
                continue
            try:
                raw = fp.read_text(encoding="utf-8")
            except Exception:
                continue
            new_raw = raw
            for old, new in perf_map.items():
                new_raw = new_raw.replace(old, new)
            for old, new in mid_map.items():
                new_raw = new_raw.replace(old, new)
            for old, new in safe_pairs:
                if old in new_raw:
                    new_raw = new_raw.replace(old, new if new else "")
            if new_raw != raw:
                changed += 1
                if apply:
                    fp.write_text(new_raw, encoding="utf-8")
    print(f"  patched text files (repo): {changed}")


def fix_spec_import_paths(apply: bool) -> None:
    """methodId 치환으로 깨진 *-spec.ts 모듈 경로 복원."""
    fixes = (
        ("@/lib/TX-002-spec", "@/lib/content-analysis-spec"),
        ("@/lib/QL-048-spec", "@/lib/q-methodology-spec"),
    )
    changed = 0
    for fp in ROOT.rglob("*"):
        if fp.suffix not in {".ts", ".tsx"} or any(p in SKIP_DIRS for p in fp.parts):
            continue
        try:
            raw = fp.read_text(encoding="utf-8")
        except Exception:
            continue
        new_raw = raw
        for a, b in fixes:
            new_raw = new_raw.replace(a, b)
        if new_raw != raw:
            changed += 1
            if apply:
                fp.write_text(new_raw, encoding="utf-8")
    print(f"  fixed spec import paths: {changed}")


def patch_gt_spec(mid_map: dict[str, str], apply: bool) -> None:
    """근거이론 GT 전용 methodId."""
    path = ROOT / "lib" / "grounded-theory-spec.ts"
    if not path.exists():
        return
    text = path.read_text(encoding="utf-8")
    old = "straussian-grounded-theory"
    new = mid_map.get(old, "QL-008")
    text2 = text.replace(f'methodId === "{old}"', f'methodId === "{new}" || methodId === "{old}"')
    text2 = text2.replace(old, new)
    if apply and text2 != text:
        path.write_text(text2, encoding="utf-8")
    print("  patched grounded-theory-spec.ts")


def patch_public_meta(apply: bool) -> None:
    fixes = [
        (PUBLIC / "manifest.webmanifest", "원전 PDF", "방법론 참고 자료"),
        (PUBLIC / "manifest.webmanifest", "원전 기반", "표준 프로토콜"),
    ]
    layout = ROOT / "app" / "layout.tsx"
    landing = ROOT / "components" / "marketing" / "landing-page.tsx"
    for path, a, b in fixes:
        if path.exists():
            t = path.read_text(encoding="utf-8")
            t2 = t.replace(a, b)
            if apply and t2 != t:
                path.write_text(t2, encoding="utf-8")
    for path in (layout, landing):
        if not path.exists():
            continue
        t = path.read_text(encoding="utf-8")
        t2 = t.replace("원전 PDF", "참고문헌").replace("원전 기반", "표준 프로토콜").replace("원전 안내", "참고문헌 안내")
        if apply and t2 != t:
            path.write_text(t2, encoding="utf-8")
    print("  patched public meta / layout / landing (if present)")


def verify_integrity(rows: list[dict[str, Any]]) -> int:
    errors: list[str] = []
    new_slugs = {r["new_slug"] for r in rows}
    dirs = {d.name for d in METHODS.iterdir() if d.is_dir()}
    for s in new_slugs:
        if s not in dirs:
            errors.append(f"MISSING_NEW_DIR {s}")
    old_left = [r["old_perf"] for r in rows if r["old_perf"] in dirs]
    if old_left:
        errors.append(f"OLD_DIRS_REMAIN {len(old_left)}")
    cat = json.loads((PUBLIC / "method-catalog.json").read_text(encoding="utf-8"))
    for m in cat["methods"]:
        if m["perfSlug"] not in new_slugs:
            errors.append(f"CATALOG_BAD_SLUG {m.get('methodId')} -> {m.get('perfSlug')}")
            break
        req = METHODS / m["perfSlug"] / "perf-registry.json"
        if not req.exists():
            errors.append(f"MISSING perf-registry {m['perfSlug']}")
            break
    if errors:
        print(f"[verify] FAILED ({len(errors)})")
        for e in errors[:25]:
            print(" ", e)
        return 1
    print("[verify] OK — 200 dirs · catalog · perf-registry")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--apply", action="store_true")
    ap.add_argument("--verify-only", action="store_true")
    args = ap.parse_args()

    if not args.xlsx.is_file():
        print("xlsx not found:", args.xlsx, file=sys.stderr)
        return 1

    rows = load_slug_rows(args.xlsx)
    perf_map, mid_map, _ = build_maps(rows)
    safe_pairs = load_safe_replacements(args.xlsx)

    if args.verify_only:
        return verify_integrity(rows)

    if args.dry_run or not args.apply:
        code = dry_run(rows)
        if code:
            return code
        print("[dry-run] Would apply:")
        print(f"  - rename {len(rows)} method directories")
        print(f"  - patch catalog, perf-index, process-guide, registry")
        print(f"  - safe text pairs: {len(safe_pairs)}")
        return 0

    print("[apply] starting...")
    rename_method_dirs(rows, apply=True)
    patch_catalog(rows, apply=True)
    patch_perf_index(rows, apply=True)
    patch_method_process_guide(perf_map, mid_map, apply=True)
    patch_registry(mid_map, apply=True)
    patch_each_method_folder(perf_map, mid_map, apply=True)
    patch_repo_text_files(perf_map, mid_map, safe_pairs, apply=True)
    fix_spec_import_paths(apply=True)
    patch_gt_spec(mid_map, apply=True)
    patch_public_meta(apply=True)
    print("[apply] done. running verify...")
    return verify_integrity(rows)


if __name__ == "__main__":
    raise SystemExit(main())
