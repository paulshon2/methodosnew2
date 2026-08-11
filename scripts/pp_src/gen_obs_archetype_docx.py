# -*- coding: utf-8 -*-
"""Generate Observation Technique DOCX: 1 × 5 × 12 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from obs_coding_examples import (
    coded_rows,
    coding_proc_apply,
    coding_proc_scheme,
    freq_rows,
    scheme_rows,
)
from obs_docx_profiles import METHOD_PROFILES
from obs_figures import figures_for_analysis
from obs_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-26.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\관찰기법 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "관찰기법 아키타입(Observation Technique Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①관찰대상·장면 선정 →②구조적/비구조적 관찰 실시 →③사건표집/시간표집 적용",
        [
            "1) 관찰대상·장면 선정하기",
            "2) 구조적(체크리스트) 또는 비구조적 관찰 실시하기",
            "3) 사건표집/시간표집 방식 적용하기",
        ],
    ),
    (
        "② 전처리",
        "①관찰기록지 정리 →②이중 관찰자료 정리(일치도 점검용)",
        [
            "1) 관찰기록지 정리하기",
            "2) 관찰자간 일치도 점검을 위한 이중 관찰자료 정리하기",
        ],
    ),
    (
        "③ 코딩",
        "①행동범주 코딩체계 개발 →②사전 정의 범주에 코딩",
        [
            "1) 행동범주 코딩체계 개발하기",
            "2) 관찰내용을 사전 정의된 범주에 코딩하기",
        ],
    ),
    (
        "④ 통계기법",
        "①행동범주별 빈도·비율 산출 →②관찰자간 신뢰도(일치율, Kappa) 검정",
        [
            "1) 행동범주별 빈도·비율 산출하기",
            "2) 관찰자간 신뢰도(일치율, Kappa) 검정하기",
        ],
    ),
    (
        "⑤ 결과도출",
        "①범주별 빈도 패턴 요약 →②행동유형의 경향성 해석",
        [
            "1) 범주별 빈도 패턴을 요약하기",
            "2) 행동유형의 경향성 해석하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "관찰결과 기술 및 범주 분류",
        ["▶ 산출물: 관찰결과 기술 및 범주 분류"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 0 and ti == 2:
        return [figs[0]] if figs else []
    if si == 3 and ti == 0:
        return [figs[1]] if len(figs) > 1 else []
    if si == 3 and ti == 1:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["장면·대상", topic["participants"]],
        ["관찰 모드", topic["mode"]],
        ["초점", topic["focus"]],
        ["분석 키", topic["analysis_key"]],
        ["관찰 명제", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    n_sess = 8 + s % 5
    is_event = "사건" in topic["mode"]
    is_struct = "구조" in topic["mode"] or "체크" in topic["mode"]

    if si == 0 and ti == 0:
        proc = [
            f"【장면선정】「{topic['title']}」 — {topic['participants']}.",
            "【단위】개인/집단/사건 단위와 관찰 가능 시간을 정한다.",
            f"【윤리】동의·가시성·민감정보. ({method_name}/{method_en})",
        ]
        del_r = [
            ["항목", "내용"],
            ["장면", topic["participants"]],
            ["목표회차", str(n_sess)],
            ["초점행동", topic["focus"]],
            ["접근/동의", "확보"],
        ]
        return STAGES[0][2][0], proc, (["회차", str(n_sess)], [["윤리", "Y"]]), ["관찰장면 선정 명세"], del_r

    if si == 0 and ti == 1:
        obs_type = "구조적(체크리스트)" if is_struct or "구조적" in topic["mode"] else "비구조적→이후 구조화"
        proc = [
            f"【관찰유형】{obs_type}로 실시한다. 모드={topic['mode']}.",
            "【훈련】관찰자 매뉴얼·예비관찰 1–2회.",
            f"【초점】{topic['focus']}.",
        ]
        del_r = [
            ["회차", "유형", "길이", "비고"],
            ["S01", "파일럿", "20분", "양식개정"],
            ["S02", "본관찰", "30분", ""],
            ["S03", "본관찰", "30분", "이중관찰"],
            ["…", "—", "—", f"총 {n_sess}회"],
        ]
        return STAGES[0][2][1], proc, (["유형", obs_type[:10]], [["훈련", "Y"]]), ["관찰 실시 로그"], del_r

    if si == 0 and ti == 2:
        samp = "사건표집" if is_event else "시간표집"
        proc = [
            f"【표집방식】{samp} 규칙을 적용한다.",
            "【정의】사건 시작/종료 또는 시간 간격·길이를 고정한다.",
            f"【기록】타임스탬프·장면ID. 「{topic['title']}」.",
        ]
        if is_event:
            del_r = [
                ["규칙", "내용"],
                ["사건정의", topic["focus"][:16]],
                ["기록시점", "사건 발생 시"],
                ["종료기준", "행동 종료/전환"],
                ["예", "손위생 기회, 갈등개시"],
            ]
        else:
            del_r = [
                ["규칙", "내용"],
                ["간격", "3분 또는 5분"],
                ["관측창", "10–20초"],
                ["기록", "그 순간의 범주"],
                ["예외", "중단·가림 표기"],
            ]
        return STAGES[0][2][2], proc, (["표집", samp], [["규칙", "문서화"]]), ["표집방식 명세"], del_r

    if si == 1 and ti == 0:
        proc = [
            "【기록지정리】장면ID·시간·관찰자·원기록·메모를 표준화한다.",
            "【누락】불명·가림 구간을 표시한다.",
            "【백업】원본 체크리스트/노트를 보관한다.",
        ]
        del_r = [
            ["장면ID", "일시", "관찰자", "완결", "비고"],
            ["SC01", "오전", "A", "Y", ""],
            ["SC02", "오후", "A", "Y", "가림2분"],
            ["SC03", "오전", "B", "Y", "이중"],
        ]
        return STAGES[1][2][0], proc, (["정리", "Y"], [["백업", "Y"]]), ["관찰기록지 정리표"], del_r

    if si == 1 and ti == 1:
        proc = [
            "【이중관찰】동일 장면의 A/B 기록을 짝지어 정리한다.",
            "【목적】관찰자간 일치도(Kappa) 산출용.",
            f"【쟁점】초점={topic['focus']} 관련 불일치 범주를 표시한다.",
        ]
        del_r = [
            ["장면", "A기록", "B기록", "짝짓기", "불일치메모"],
            ["SC03", "Y", "Y", "Y", "B02/B03"],
            ["SC05", "Y", "Y", "Y", ""],
            ["SC07", "Y", "Y", "Y", "B04정의"],
        ]
        return STAGES[1][2][1], proc, (["이중자료", "Y"], [["재합의", "예정"]]), ["이중 관찰자료 정리"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc_scheme(topic, method_name, s)
        del_r = scheme_rows(topic, s)
        return STAGES[2][2][0], proc, (["코드북", "v0.9"], [["상호배타", "Y"]]), ["행동범주 코딩체계"], del_r

    if si == 2 and ti == 1:
        proc = coding_proc_apply(topic, method_name)
        del_r = coded_rows(topic, s)
        return STAGES[2][2][1], proc, (["코딩", "적용"], [["QC", "이중"]]), ["범주 코딩표"], del_r

    if si == 3 and ti == 0:
        proc = [
            f"【빈도·비율】행동범주별 빈도·비율을 산출한다. 「{topic['title']}」.",
            "【교차】시간대·조건·장면별 비교표를 만든다.",
            f"【방법】{method_name}.",
        ]
        del_r = freq_rows(topic, s)
        return STAGES[3][2][0], proc, (["빈도표", "Y"], [["교차표", "Y"]]), ["범주 빈도·비율"], del_r

    if si == 3 and ti == 1:
        kappa = 0.62 + (s % 10) * 0.02
        agree = 78 + s % 12
        proc = [
            "【신뢰도】일치율과 Cohen's Kappa를 산출한다.",
            f"【결과(도식)】일치율≈{agree}%, κ≈{kappa:.2f}.",
            "【개선】κ 낮은 범주 정의·재훈련.",
        ]
        del_r = [
            ["지표", "값", "기준", "판정"],
            ["일치율", f"{agree}%", "≥75%", "양호"],
            ["Cohen κ", f"{kappa:.2f}", "≥.60", "양호/경계"],
            ["문제범주", "B02/B03", "재정의", "진행"],
        ]
        return STAGES[3][2][1], proc, (["κ", f"{kappa:.2f}"], [["재훈련", "Y"]]), ["관찰자간 신뢰도"], del_r

    if si == 4 and ti == 0:
        proc = [
            "【패턴요약】범주별 빈도 패턴을 요약한다.",
            f"【포인트】초점={topic['focus']} 관련 고빈도/저빈도.",
            "【시각화】막대·시계열로 제시한다.",
        ]
        del_r = [
            ["패턴", "요지", "근거"],
            ["고빈도", "온과제/핵심범주", "B01/B04"],
            ["저빈도", "기타", "B05"],
            ["조건차", "질문직후↑ 등", "교차표"],
            ["명제연결", topic["hypothesis"][:28], ""],
        ]
        return STAGES[4][2][0], proc, (["요약", "Y"], [["그림", "Y"]]), ["빈도 패턴 요약"], del_r

    if si == 4 and ti == 1:
        proc = [
            f"【경향해석】행동유형의 경향을 해석한다. 명제: {topic['hypothesis']}",
            "【맥락】장면·역할·시간대 조건을 함께 서술한다.",
            "【한계】반응성·표집·관찰자효과를 명시한다.",
        ]
        del_r = [
            ["경향", "해석", "함의"],
            ["참여↑구간", "촉발 단서 존재", "개입타이밍"],
            ["이탈↑구간", "대기·공백", "구조개선"],
            ["초점행동", topic["focus"][:12], "우선관찰"],
            ["종합", topic["hypothesis"][:30], "실무연결"],
        ]
        return STAGES[4][2][1], proc, (["해석", "Y"], [["한계", "명시"]]), ["행동 경향 해석"], del_r

    # si == 5
    proc = [
        "【산출물】관찰결과 기술 및 범주 분류 패키지.",
        f"【핵심】{topic['hypothesis']}",
        "【구성】장면설계·표집규칙·코드북·코딩·빈도·Kappa·해석·한계.",
        f"【모드】{topic['mode']} · {method_name}({method_en}).",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["observation_report.docx", "관찰결과·범주분류", topic["title"][:20]],
        ["codebook.xlsx", "행동범주 체계", ""],
        ["freq_kappa.xlsx", "빈도·일치도", ""],
        ["sheets/", "기록지원본", ""],
        ["figA–C.png", "표집·빈도·κ", ""],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "observation report"]]), ["관찰결과·범주 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」을(를) {topic['mode']}로 관찰할 때, "
        f"{method_name}({method_en})의 범주코딩·빈도·Kappa 절차로 "
        f"{topic['hypothesis']}을(를) 근거 있게 제시할 수 있는가?"
    )
    doc.add_paragraph(f"관찰 명제: {topic['hypothesis']}")
    doc.add_paragraph(
        f"장면 / 모드 / 초점: {topic['participants']} / {topic['mode']} / {topic['focus']}"
    )

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    doc.add_heading("【관찰기법 설계 명세】", level=3)
    doc.add_paragraph("장면·모드·초점 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "블록", "ID", "그룹", "파일", "항목", "발췌ID", "범주", "주제", "관계",
        "회차", "규칙", "장면ID", "장면", "범주코드", "구간/사건", "지표", "패턴", "경향", "구분",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str, no: int) -> str:
    base = name_ko
    en_c = en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"관찰기법 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 12"
        )
        doc.add_paragraph(
            "본 문서는 관찰기법 아키타입의 6단계·12태스크에 따라 "
            "장면선정·구조/비구조관찰·사건/시간표집·기록지정리·이중관찰·"
            "행동범주체계·코딩·빈도·Kappa·패턴요약·경향해석·관찰결과 패키지를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"], no)
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
