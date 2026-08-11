# -*- coding: utf-8 -*-
"""Generate Nursing Theory Construction DOCX: 1 × 5 × 11 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from ntc_coding_examples import (
    cat_rows,
    coding_proc_cat,
    coding_proc_open,
    coding_proc_theme,
    open_rows,
    theme_rows,
)
from ntc_docx_profiles import METHOD_PROFILES
from ntc_figures import figures_for_analysis
from ntc_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-25.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\간호이론구축 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "간호이론구축 아키타입(Nursing Theory Construction Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①목적표집으로 참여자 선정 →②심층면담·참여관찰·문서수집 실시",
        [
            "1) 목적표집으로 참여자 선정하기",
            "2) 심층면담·참여관찰·문서수집 실시하기",
        ],
    ),
    (
        "② 전처리",
        "①전사 및 자료정리 →②분석 소프트웨어에 입력",
        [
            "1) 전사 및 자료정리하기",
            "2) 분석 소프트웨어에 입력하기",
        ],
    ),
    (
        "③ 코딩",
        "①개방코딩으로 초기코드 생성 →②범주화 →③주제 도출",
        [
            "1) 개방코딩으로 초기코드 생성하기",
            "2) 범주화하기",
            "3) 주제 도출하기",
        ],
    ),
    (
        "④ 통계기법",
        "해당없음(필요시 기초빈도만 보조적 제시)",
        ["1) 해당없음(필요시 기초빈도만 보조적 제시)"],
    ),
    (
        "⑤ 결과도출",
        "①주제·범주를 해석적으로 통합 →②서술적 결과로 제시",
        [
            "1) 주제·범주를 해석적으로 통합하기",
            "2) 서술적 결과로 제시하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "질적 분석결과(주제·이론)",
        ["▶ 산출물: 질적 분석결과(주제·이론)"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 0 and ti == 1:
        return [figs[0]] if figs else []
    if si == 2 and ti == 2:
        return [figs[2]] if len(figs) > 2 else []
    if si == 4 and ti == 0:
        return [figs[1]] if len(figs) > 1 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["참여자", topic["participants"]],
        ["자료 모드", topic["mode"]],
        ["초점", topic["focus"]],
        ["분석 키", topic["analysis_key"]],
        ["이론 명제(초안)", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    n_p = 12 + s % 5

    if si == 0 and ti == 0:
        proc = [
            f"【목적표집】「{topic['title']}」 — {topic['participants']}.",
            "【기준】포함·제외·최대변이(경험강도·역할·맥락)를 문서화한다.",
            f"【목표】약 {n_p}명(+이론표집 여지). ({method_name}/{method_en})",
        ]
        del_r = [
            ["ID", "역할", "포함이유", "동의"],
            ["P01", "핵심경험", "정보풍부", "Y"],
            ["P02", "가족/동료", "대비관점", "Y"],
            ["P03", "경계사례", "이론확장", "Y"],
        ]
        return STAGES[0][2][0], proc, (["n가안", str(n_p)], [["윤리", "Y"]]), ["목적표집 로그"], del_r

    if si == 0 and ti == 1:
        proc = [
            f"【복수자료】심층면담·참여관찰·문서를 수집한다. 모드={topic['mode']}.",
            f"【초점】{topic['focus']}가 드러나는 장면·발화를 우선한다.",
            "【메모】이론적 메모·현장노트를 즉시 남긴다.",
        ]
        del_r = [
            ["자료", "분량", "역할", "비고"],
            ["면담", f"{n_p}회", "의미·과정", "60–90분"],
            ["관찰", "6–10회", "돌봄실천", "병동/가정"],
            ["문서", "간호기록·지침", "맥락", "허가범위"],
        ]
        return STAGES[0][2][1], proc, (["삼각검증", "계획"], [["메모", "Y"]]), ["자료수집 로그"], del_r

    if si == 1 and ti == 0:
        proc = [
            "【전사·정리】면담 축어전사, 관찰노트·문서를 정리·익명화한다.",
            "【QC】난해 구간 재생 대조, 파일명 표준화.",
            f"【적용】「{topic['title']}」.",
        ]
        del_r = [
            ["파일", "유형", "상태", "익명"],
            ["P01.docx", "면담", "확정", "Y"],
            ["OBS03.docx", "관찰", "확정", "Y"],
            ["DOC_A.pdf", "문서", "허가범위", "Y"],
        ]
        return STAGES[1][2][0], proc, (["전사", "진행"], [["QC", "Y"]]), ["전사·정리 체크"], del_r

    if si == 1 and ti == 1:
        proc = [
            "【SW입력】분석 소프트웨어에 사례·자료 단위로 입력한다.",
            "【메타】참여자·일시·자료유형·동의상태를 태그한다.",
            "【백업】암호화·접근권한 분리.",
        ]
        del_r = [
            ["항목", "설정"],
            ["SW", "NVivo/ATLAS.ti/기타"],
            ["프로젝트", topic["analysis_key"]],
            ["케이스", f"P01–P{n_p:02d}"],
            ["백업", "이중저장"],
        ]
        return STAGES[1][2][1], proc, (["입력", "Y"], [["태그", "Y"]]), ["분석SW 입력 명세"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc_open(topic, method_name, s)
        del_r = open_rows(topic, s)
        return STAGES[2][2][0], proc, (["초기코드", str(35 + s % 20)], [["메모", "Y"]]), ["개방코딩 표"], del_r

    if si == 2 and ti == 1:
        proc = coding_proc_cat(topic, method_name)
        del_r = cat_rows(topic)
        return STAGES[2][2][1], proc, (["범주", "4"], [["정의", "고정"]]), ["범주화 표"], del_r

    if si == 2 and ti == 2:
        proc = coding_proc_theme(topic, method_name)
        del_r = theme_rows(topic)
        return STAGES[2][2][2], proc, (["주제", "3–5"], [["명제초안", "Y"]]), ["주제·이론요소표"], del_r

    if si == 3 and ti == 0:
        proc = [
            "【원칙】통계적 가설검정은 사용하지 않는다.",
            "【보조】필요 시 주제/코드 출현 기초빈도만 제시한다.",
            f"【주제】「{topic['title']}」 — 핵심은 간호이론 서술.",
        ]
        del_r = [
            ["항목", "결정"],
            ["가설검정", "미사용"],
            ["주근거", "주제·인용·이론통합"],
            ["보조빈도", "선택(증명용 아님)"],
        ]
        return STAGES[3][2][0], proc, (["stats", "N/A"], [["보조빈도", "선택"]]), ["통계 미사용·보조빈도 선언"], del_r

    if si == 4 and ti == 0:
        proc = [
            f"【해석적 통합】주제·범주를 중범위 간호이론으로 통합한다. 「{topic['title']}」.",
            f"【명제】{topic['hypothesis']}",
            f"【틀】조건–과정–결과. ({method_name}/{method_en})",
        ]
        del_r = [
            ["이론요소", "내용", "근거주제"],
            ["조건", "지지공백·가치충돌", "T1"],
            ["과정", "협상·조율·재구성", "T3"],
            ["결과", "적응/고뇌완화/정체성", "T4"],
            ["중심명제", topic["hypothesis"][:36], "통합"],
        ]
        return STAGES[4][2][0], proc, (["이론통합", "초안"], [["도식", "Y"]]), ["이론 통합 요지"], del_r

    if si == 4 and ti == 1:
        proc = [
            "【서술적 결과】주제별 인용+해석으로 서술하고 이론 명제를 제시한다.",
            f"【함의】초점={topic['focus']}에 대한 간호중재·교육 함의.",
            "【한계】표집·전이가능성·연구자 위치를 명시한다.",
        ]
        del_r = [
            ["주제", "대표인용", "해석요지"],
            ["T1", "「혼자…」", "비가시 부담"],
            ["T2", "「몸이…」", "체현 신호"],
            ["T3", "「맞춰요」", "협상적 돌봄"],
            ["이론", "—", topic["hypothesis"][:28]],
        ]
        return STAGES[4][2][1], proc, (["서술", "Y"], [["함의", "구분"]]), ["서술적 결과 요지"], del_r

    # si == 5
    proc = [
        "【산출물】질적 분석결과(주제·이론) 패키지.",
        f"【핵심】{topic['hypothesis']}",
        "【구성】표집·자료·코드북·주제·이론도식·인용록·성찰·함의.",
        f"【모드】{topic['mode']} · {method_name}({method_en}).",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["nursing_theory.docx", "주제·이론 결과", topic["title"][:20]],
        ["codebook.xlsx", "개방·범주·주제", ""],
        ["theory_map.png", "조건–과정–결과", ""],
        ["quotes.xlsx", "인용록", "가명"],
        ["figA–C.png", "자료·이론·주제", ""],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "themes+theory"]]), ["주제·이론 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」에 대해 목적표집·복수자료로 탐구할 때, "
        f"{method_name}({method_en})의 개방코딩–범주–주제 절차로 "
        f"{topic['hypothesis']}을(를) 간호 중범위 이론으로 제시할 수 있는가?"
    )
    doc.add_paragraph(f"이론 명제(초안): {topic['hypothesis']}")
    doc.add_paragraph(
        f"참여자 / 모드 / 초점: {topic['participants']} / {topic['mode']} / {topic['focus']}"
    )

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    doc.add_heading("【간호이론구축 설계 명세】", level=3)
    doc.add_paragraph("참여자·자료·초점 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "블록", "ID", "그룹", "파일", "항목", "발췌ID", "범주", "주제", "관계",
        "자료", "이론요소", "구분",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str, no: int) -> str:
    base = name_ko
    en_c = en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"간호이론구축 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 11"
        )
        doc.add_paragraph(
            "본 문서는 간호이론구축 아키타입의 6단계·11태스크에 따라 "
            "목적표집·복수자료·전사·SW입력·개방코딩·범주·주제·"
            "이론통합·서술결과·주제·이론 패키지를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"], no)
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
