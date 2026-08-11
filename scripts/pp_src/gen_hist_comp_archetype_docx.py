# -*- coding: utf-8 -*-
"""Generate Historical–Comparative Source Archetype DOCX: 15 methods × 5 topics × 12 tasks."""
from __future__ import annotations

import json
import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from hist_comp_coding_examples import (
    codebook_rows,
    coding_application_rows,
    coding_procedure_paras,
    double_coding_rows,
    mechanism_procedure_paras,
    process_evidence_rows,
)
from hist_comp_docx_profiles import METHOD_PROFILES
from hist_comp_figures import figures_for_analysis
from hist_comp_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-02.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\역사·비교사료 아키타입-연구방법별 샘플_정교화")

ARCHETYPE = "역사·비교사료 아키타입(Historical–Comparative Source Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①1차 사료(공문서, 기록물)와 2차 문헌 수집 →②HC-001의 경우 증언자 면담 및 녹음 →③비교연구는 복수 사례(국가/시대)의 자료를 병행 수집",
        [
            "1) 1차 사료(공문서, 기록물)와 2차 문헌 수집하기",
            "2) HC-001의 경우 증언자 면담 및 녹음하기",
            "3) 비교연구는 복수 사례(국가/시대)의 자료를 병행 수집하기",
        ],
    ),
    (
        "② 전처리",
        "①사료의 출처·신뢰성 비평(사료비평) →②연대기(chronology) 재구성 및 사건 목록화",
        [
            "1) 사료의 출처·신뢰성 비평(사료비평)하기",
            "2) 연대기(chronology) 재구성 및 사건 목록화하기",
        ],
    ),
    (
        "③ 코딩",
        "①주제적·개념적 코딩으로 사건·행위자·맥락 범주화 →②과정추적(process tracing)의 경우 인과기제(causal mechanism)를 단계별로 도식화",
        [
            "1) 주제적·개념적 코딩으로 사건·행위자·맥락 범주화하기",
            "2) 과정추적(process tracing)의 경우 인과기제(causal mechanism)를 단계별로 도식화하기",
        ],
    ),
    (
        "④ 통계기법",
        "①원칙적으로 통계 미사용 →②계량사학·비교연구에서는 사건빈도·시계열 기초통계를 보조적으로 활용",
        [
            "1) 원칙적으로 통계 미사용하기",
            "2) 계량사학·비교연구에서는 사건빈도·시계열 기초통계를 보조적으로 활용하기",
        ],
    ),
    (
        "⑤ 결과도출",
        "①사건연쇄와 맥락을 통합해 역사적 해석 서술 →②사례 간 비교를 통해 유사점·차이점에 대한 명제 도출",
        [
            "1) 사건연쇄와 맥락을 통합해 역사적 해석 서술하기",
            "2) 사례 간 비교를 통해 유사점·차이점에 대한 명제 도출하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "역사적 서술 및 비교분석 결과",
        ["▶ 산출물: 역사적 서술 및 비교분석 결과"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc: Document, headers: list[str], rows: list[list[str]]):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = h
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc: Document, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.6))
    doc.add_paragraph()


def figures_for_task(stage_i: int, task_i: int, figs: list[tuple[str, Path]]) -> list[tuple[str, Path]]:
    # A on chronology / stats-aux / deliverable; B on process tracing / compare / deliverable
    if stage_i == 1 and task_i == 1:
        return [figs[0]] if figs else []
    if stage_i == 2 and task_i == 1:
        return [figs[1]] if len(figs) > 1 else []
    if stage_i == 3 and task_i == 1:
        return figs[:2]
    if stage_i == 4 and task_i == 1:
        return [figs[1]] if len(figs) > 1 else []
    if stage_i == 5 and task_i == 0:
        return figs
    return []


def build_source_inventory(topic: dict, method_name: str) -> list[list[str]]:
    seed = _seed(topic["title"], method_name)
    akey = topic["analysis_key"]
    rows = [
        ["SRC-01", "1차", "공문서/기록물철", f"{topic['period']} 생산", "국립·시·도 기록원", f"청구기호 H-{seed % 900 + 100}", "열람"],
        ["SRC-02", "1차", "신문·잡지", topic["period"], "디지털아카이브", f"DB-{seed % 50}", "전문"],
        ["SRC-03", "2차", "연구서·논문", "후대 연구", "도서관", f"ISBN/DOI-{seed}", "배경"],
    ]
    if akey.startswith("oral") or akey.startswith("ohp"):
        rows.insert(
            0,
            ["OH-01", "1차", "구술녹음·전사", "증언자 면담", "연구자 아카이브", f"OH-{seed % 200}", "동의 완료"],
        )
    if "비교" in topic["approach"] or akey.startswith("comp") or akey.startswith("cha") or akey.startswith("pt"):
        rows.append(["SRC-04", "1차", "비교사례 B 사료묶음", topic["cases"], "해외·지방 아카이브", f"CMP-{seed % 80}", "대칭수집"])
    if akey.startswith("arch"):
        rows.append(["ARC-01", "1차", "기록물 시리즈 전체목록", "생산기관 이관", "중앙기록원", f"SER-{seed}", "시리즈 단위"])
    return rows


def task_content(
    stage_i: int, task_i: int, topic: dict, method_name: str, method_en: str
) -> tuple[str, list[str], tuple[list[str], list[list[str]]], list[str], list[list[str]]]:
    seed = _seed(method_name, topic["title"], stage_i, task_i)
    akey = topic["analysis_key"]
    n_src = 18 + seed % 40
    n_oral = 8 + seed % 16

    if stage_i == 0 and task_i == 0:
        inv = build_source_inventory(topic, method_name)
        proc = [
            f"【수집 설계】주제 「{topic['title']}」·시기 {topic['period']}·사례 {topic['cases']}. "
            f"초점: {topic['focus']}. 접근: {topic['approach']} ({method_name}/{method_en}).",
            f"【1차 사료】{topic['sources']}를 우선 수집. 소장처·청구기호·열람일·복사 범위를 수집로그에 기록. 목표 목록 {n_src}건.",
            f"【2차 문헌】연구사 정리용 논문·단행본 12–20편. 주장의 핵심 근거는 1차로 재확인하는 원칙을 명시.",
        ]
        del_r = [["ID", "구분", "유형", "시기/생산", "소장", "식별자", "상태"]] + inv
        return (
            STAGES[0][2][0],
            proc,
            (["항목", "내용"], [["주제", topic["title"]], ["시기", topic["period"]], ["자료", topic["sources"]]]),
            ["사료 수집목록(Inventory)"],
            del_r,
        )

    if stage_i == 0 and task_i == 1:
        oralish = akey.startswith("oral") or akey.startswith("ohp") or "구술" in method_name
        if oralish:
            proc = [
                f"【증언자】목표 {n_oral}명. 포함: {topic['period']} 경험자, 역할·지역 다양성. 제외: 대리진술만 가능한 경우(별도 표시).",
                "【면담】반구조화 가이드(생애/사건/의미) 45–90분. 동의서(사용범위·익명·철회)·장비 점검·백업.",
                "【산출】원녹음 WAV, 전사본(타임코드), 면담일지. 민감구간 마스킹 규칙 적용.",
            ]
            del_r = [
                ["증언ID", "역할", "면담일", "분", "동의", "전사"],
                ["OH01", "노동자", "2026-03-12", "78", "Y", "완료"],
                ["OH02", "사무직", "2026-03-18", "62", "Y", "진행"],
                ["OH03", "가족", "2026-03-22", "55", "Y(가명)", "완료"],
            ]
        else:
            proc = [
                f"【해당 범위】본 주제({method_name})는 문헌·기록·개념 중심이므로 신규 구술면담은 필수 아님.",
                "【대체】필요시 기존 구술아카이브 2차 이용 또는 N/A 명시. 윤리 심의 해당 없음(문헌만) 체크.",
                "【기록】‘구술 미실시 사유’를 방법란에 1문단으로 남김.",
            ]
            del_r = [["항목", "기록"], ["구술", "N/A"], ["사유", "문헌·기록 중심 설계"], ["대안", "기존 아카이브 인용(해당시)"]]
        return STAGES[0][2][1], proc, (["지표", "값"], [["계획 면담", str(n_oral) if oralish else "0"]]), ["면담 로그/ N/A 기록"], del_r

    if stage_i == 0 and task_i == 2:
        multi = (
            akey.startswith("comp")
            or akey.startswith("cha")
            or akey.startswith("pt")
            or akey.startswith("kkv")
            or "비교" in topic["approach"]
            or "비교" in method_name
        )
        if multi:
            proc = [
                f"【복수 사례】{topic['cases']}. 공통 수집 항목: 제도문서·언론·핵심 행위자 기록.",
                "【대칭성】사례 A/B에 동일 유형 사료를 맞춤. 결측은 ‘자료 공백’으로 표기.",
                f"【비교틀】초점({topic['focus']})에 따른 공통 변수표 초안 작성.",
            ]
            del_r = [
                ["사례", "핵심사료", "건수", "공백"],
                ["A", "법령·회의록", str(40 + seed % 20), "일부 비공개"],
                ["B", "법령·신문", str(35 + seed % 15), "지방지 누락"],
            ]
        else:
            proc = [
                "【단일·심화】본 주제는 단일 사례/텍스트 심화이므로 ‘복수 국가 병행’은 축소 적용.",
                "【내부 복수】시기 구분(전기/후기) 또는 매체 복수(공문서+사적기록)로 비교 가능성 확보.",
                "【명시】방법란에 ‘단일사례 설계’와 일반화 한계를 기입.",
            ]
            del_r = [["설계", "내용"], ["사례수", "1(+시기 구분)"], ["내부비교", "전기/후기·매체"], ["한계", "외적 일반화 제한"]]
        return STAGES[0][2][2], proc, (["항목", "내용"], [["cases", topic["cases"]]]), ["복수사례 수집표"], del_r

    if stage_i == 1 and task_i == 0:
        proc = [
            f"【외증】생산기관·직인·필체·용지·이관경로(provenance) 점검. 대상 {min(n_src, 25)}건 중 위조 의심 {seed % 2}건 → 잠정 보류.",
            "【내증】날짜·인명·지명 교차확인. 작성자 이해관계·검열·회고 시점 왜곡 가능성 메모.",
            "【등급】A(교차확인)·B(단일)·C(회고/이차). 인용 강도에 반영.",
        ]
        del_r = [
            ["사료ID", "외증", "내증", "등급", "비고"],
            ["SRC-01", "합격", "일부 충돌", "B", "날짜 하루 차이"],
            ["SRC-02", "합격", "합격", "A", "타지와 일치"],
            ["OH-01", "해당", "기억 재구성", "B", "문서와 교차"],
        ]
        return STAGES[1][2][0], proc, (["규칙", "적용"], [["보류", "외증 실패"], ["주의", "C등급 단독인용 금지"]]), ["사료비평표"], del_r

    if stage_i == 1 and task_i == 1:
        proc = [
            f"【사건카드】시기 {topic['period']} 내 사건 {12 + seed % 20}건. 필드: 일자·행위자·장소·요약·출처ID.",
            "【충돌】동일 사건의 상이한 날짜·인원 서술은 병기. circa/추정 표기.",
            "【산출】연대기표 + 그림 A(타임라인). 전환점 3곳에 표시.",
        ]
        del_r = [
            ["일자", "사건", "행위자", "출처", "확실도"],
            ["1948-08-15", "정부수립 관련 행사", "정부·시민", "SRC-02", "확실"],
            ["1950-06-25", "전쟁 발발(맥락)", "국가행위자", "SRC-01", "확실"],
            ["1987-06-10", "민주화 국면 사건", "시민·학생", "OH/SRC", "교차"],
        ]
        return STAGES[1][2][1], proc, (["항목", "수"], [["사건카드", str(12 + seed % 20)]]), ["연대기·사건목록"], del_r

    if stage_i == 2 and task_i == 0:
        proc = coding_procedure_paras(akey, topic, method_name, seed)
        cb = codebook_rows(akey, topic, seed)
        del_r = coding_application_rows(akey, topic, seed)
        del_p = [
            "결과물 ① 코드북(아래 상세 실행 절차의 표: 정의·포함·제외).",
            "결과물 ② 코딩 적용 예시표(원문 발췌 → 부여 코드 → 근거 1문장) — 아래 표.",
            "결과물 ③ 이중코딩 합의표 — 본 태스크 결과물 샘플 다음에 제시.",
            f"코딩 단위·초점: {topic['focus']} / analysis_key={akey}",
        ]
        return (
            STAGES[2][2][0],
            proc,
            (cb[0], cb[1:]),
            del_p,
            del_r,
        )

    if stage_i == 2 and task_i == 1:
        use_pt = (
            akey.startswith("pt_")
            or "과정" in method_name
            or akey.startswith("cha")
            or akey.startswith("comp")
            or akey.startswith("kkv")
        )
        proc = mechanism_procedure_paras(topic, method_name, use_pt)
        del_r = process_evidence_rows(topic, seed, concrete=use_pt)
        del_p = [
            "결과물: 기제단계×증거×코드 대응표(또는 개념·주장 연결표).",
            "약한 단계(Insufficient)는 추가 수집 과제와 그림 B 점선으로 남긴다.",
            f"중심 명제: {topic['hypothesis']}",
        ]
        return (
            STAGES[2][2][1],
            proc,
            (
                ["도구", "산출"],
                [["기제도식/개념지도", "figB"], ["증거표", "아래"], ["MECH_* 코드", "코드북과 연동"]],
            ),
            del_p,
            del_r,
        )

    if stage_i == 3 and task_i == 0:
        proc = [
            f"【원칙】{ARCHETYPE}는 가설 확인을 통계검정(p값)에 의존하지 않는다. 「{method_name}」 역시 서술·비교·기제 추론이 본령.",
            "【방법란 문구】‘본 연구는 추론통계를 사용하지 않으며, 사료비평·코딩·비교/과정추적으로 주장을 뒷받침한다.’",
            "【금지】표본 평균 차이 t검정 등으로 역사적 인과를 ‘증명’했다고 쓰지 않음.",
        ]
        del_r = [["항목", "결정"], ["추론통계", "미사용"], ["사유", "아키타입·방법 본령"], ["대체", "사료근거 서술"]]
        return STAGES[3][2][0], proc, (["선언", "예"], [["통계미사용", "Y"]]), ["통계미사용 선언문"], del_r

    if stage_i == 3 and task_i == 1:
        use_aux = akey.startswith(("lsb_", "comp", "cha", "hs_", "pt_", "ihr", "ih_"))
        if use_aux or seed % 3 == 0:
            proc = [
                "【보조 통계】사건빈도·연도별 건수·문헌성장 등 기초집계만 제시. 인과 주장의 주근거는 아님.",
                f"【표·그림】연도별 코딩 사건 빈도표 + 그림 B(또는 성장곡선). 모집단=본 연구 코딩 코퍼스({30 + seed % 50}건).",
                "【해석 주의】‘보도량↑ ≠ 실제발생↑’. 표본(신문종·검색어)을 각주에 명시.",
            ]
            del_r = [
                ["연도", "사건수", "비고"],
                ["1985", str(4 + seed % 5), "코딩 코퍼스"],
                ["1986", str(7 + seed % 6), ""],
                ["1987", str(14 + seed % 8), "정점"],
                ["1988", str(9 + seed % 5), ""],
            ]
        else:
            proc = [
                "【보조 통계 생략】본 주제는 단일 텍스트·철학·윤리 실무 중심으로 빈도표를 생략하거나 최소 제시.",
                "【대체】인용 분포(장·절) 정도만 부록에 둘 수 있음.",
                "【명시】‘보조통계 N/A’ 체크.",
            ]
            del_r = [["보조통계", "N/A"], ["사유", "주제 특성"], ["부록", "인용분포(선택)"]]
        return STAGES[3][2][1], proc, (["유형", "사용"], [["빈도", "Y" if use_aux else "선택"]]), ["빈도·시계열 보조표"], del_r

    if stage_i == 4 and task_i == 0:
        proc = [
            f"【해석 서술】「{topic['title']}」: {topic['hypothesis']}",
            f"【구성】배경({topic['period']}) → 사건연쇄(연대기) → 행위자·제도 맥락 → 사료 인용(등급 A/B) → 종합 해석.",
            "【문체】주장·근거·한계를 문단 단위로 분리. 과잉 일반화 금지.",
        ]
        del_r = [
            ["절", "내용 요약"],
            ["1", "문제제기·연구사"],
            ["2", "사료와 방법"],
            ["3", "사건연쇄 해석"],
            ["4", "논의·한계"],
        ]
        return STAGES[4][2][0], proc, (["분량", "초안"], [["본문", "4,000–7,000자"]]), ["역사적 해석 초고"], del_r

    if stage_i == 4 and task_i == 1:
        proc = [
            f"【비교 축】사례 {topic['cases']} · 초점 {topic['focus']}.",
            "【유사·차이표】제도·행위자·결과·기제 4행 이상. 근거 사료ID 병기.",
            f"【명제】검증·유보를 표기한 명제 3개. 중심 명제: {topic['hypothesis']}",
        ]
        del_r = [
            ["명제", "지지근거", "판정"],
            [topic["hypothesis"][:40] + "…", "SRC/OH 교차", "부분지지"],
            ["대안설명 가능성", "경쟁기제", "유보"],
            ["범위조건(scope)", "시기·지역", "한정"],
        ]
        return STAGES[4][2][1], proc, (["표", "비교행렬"], [["완료", "Y"]]), ["비교명제표"], del_r

    # deliverable
    proc = [
        f"【요약(Abstract)】「{topic['title']}」({topic['period']}). 사례: {topic['cases']}. "
        f"자료: {topic['sources']}. 방법: {method_name}({method_en}) · {topic['approach']}. "
        f"주장: {topic['hypothesis']}. 산출: 서술·비교표·연대기·도식.",
        "【패키지】1) 본문 서술 2) 사료목록·비평표 3) 연대기·그림A 4) 코드북·기제도식/그림B 5) 비교명제 6) 윤리·접근 기록.",
        f"【재현】청구기호·타임코드·버전 {date.today().isoformat()} · audit trail 포함.",
        "【한계】사료 공백·기억 재구성·단일/소수사례 일반화 한계를 명시.",
    ]
    del_r = [
        ["유형", "파일/산출", "내용"],
        ["서술", "narrative_v1.docx", "역사적 해석 본문"],
        ["표", "Table_sources.xlsx", "수집·비평"],
        ["표", "Table_chrono.xlsx", "연대기"],
        ["표", "Table_propositions.docx", "비교 명제"],
        ["그림", "FigA_timeline.png", "타임라인"],
        ["그림", "FigB_mechanism.png", "기제/비교/빈도"],
        ["부록", "codebook.pdf", "코드 정의"],
        ["윤리", "consent_log.pdf", "해당시"],
        ["결론(1문장)", "—", topic["hypothesis"]],
    ]
    return STAGES[5][2][0], proc, (["산출", "버전"], [["package_v1.0", date.today().isoformat()]]), ["출판용 역사·비교 패키지"], del_r


def write_topic_section(doc: Document, topic_idx: int, topic: dict, method_name: str, method_en: str, method_no: int):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」에 대해 {method_name}({method_en})의 표준 사료·비교·해석 절차를 적용할 때, "
        f"{topic['hypothesis']}을(를) 사료 근거로 뒷받침할 수 있는가?"
    )
    doc.add_paragraph(f"중심 주장(명제): {topic['hypothesis']}")
    doc.add_paragraph(f"시기·사례·자료: {topic['period']} / {topic['cases']} / {topic['sources']}")
    doc.add_paragraph(f"접근·초점: {topic['approach']} · {topic['focus']}")

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    inv = build_source_inventory(topic, method_name)
    doc.add_heading("【확정 사료·자료 목록(전문 샘플)】", level=3)
    doc.add_paragraph(
        f"본 주제의 핵심 자료 구성입니다. 태스크 1) 수집 결과물과 동일 체계이며, 전체 수집목표는 약 {18 + fig_seed % 40}건 규모로 확장합니다."
    )
    add_table(doc, ["ID", "구분", "유형", "시기/생산", "소장", "식별자", "상태"], inv)

    header_rows = {
        "변수",
        "가설",
        "집단",
        "단계",
        "일자",
        "코드",
        "명제",
        "유형",
        "항목",
        "절",
        "연도",
        "사례",
        "사료ID",
        "증언ID",
        "주장요소",
        "설계",
        "보조통계",
        "ID",
        "발췌ID",
        "기제단계",
        "표본",
        "도구",
    }

    for si, (stage_label, stage_sum, tasks) in enumerate(STAGES):
        doc.add_heading(stage_label, level=3)
        doc.add_paragraph(stage_sum)
        for ti, task_label in enumerate(tasks):
            _, proc_paras, proc_tbl, del_paras, del_rows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(task_label, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(task_label)

            step_notes = task_step_annotations(si, ti, proc_paras, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for step_i, (para, note) in enumerate(zip(proc_paras, step_notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({step_i}) 실행 절차 요지: ").bold = True
                preview = para
                head.add_run(preview)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc_paras:
                doc.add_paragraph(para)
            if proc_tbl[1]:
                add_table(doc, proc_tbl[0], proc_tbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in del_paras:
                doc.add_paragraph(dp)
            if del_rows:
                if del_rows[0][0] in header_rows or (len(del_rows) > 1 and len(del_rows[0]) == len(del_rows[1])):
                    add_table(doc, del_rows[0], del_rows[1:])
                else:
                    ncols = max(len(r) for r in del_rows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], del_rows)

            # ③ 코딩 태스크1: 이중코딩 합의표 추가
            if si == 2 and ti == 0:
                doc.add_paragraph().add_run("■ 결과물 샘플 보충: 이중코딩 합의표").bold = True
                doc.add_paragraph(
                    "같은 발췌를 코더 A·B가 독립 코딩한 뒤 불일치를 유형별로 합의·코드북에 반영한 기록이다."
                )
                dc = double_coding_rows(_seed(method_name, topic["title"], si, ti))
                add_table(doc, dc[0], dc[1:])

            for cap, fpath in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fpath, cap)

            doc.add_paragraph()


def safe_filename(name_ko: str, en: str, no: int) -> str:
    base = name_ko
    en_clean = en.replace("/", "-").replace("\\", "-")
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_clean = en_clean.replace(ch, "")
    return f"역사·비교사료 아키타입-{base}({en_clean}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    methods = bundle["methods"]

    for m in methods:
        no = m["no"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip missing profile", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        title = doc.add_heading(f"{ARCHETYPE}", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {m['name']}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 연구방법 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 12"
        )
        doc.add_paragraph(
            "본 문서는 역사·비교사료 아키타입의 6단계·12태스크 절차에 따라, "
            "각 연구방법의 표준 주제·사료목록·사료비평·연대기·코딩·과정추적·비교명제·그림을 "
            "출판(학술보고) 수준의 형식으로 제시한 실행 샘플이다."
        )
        doc.add_page_break()

        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, m["name"], prof["en"], no)
            if i < 5:
                doc.add_page_break()

        path = OUT_DIR / safe_filename(m["name"], prof["en"], no)
        tmp_path = path.with_suffix(".docx.generating")
        doc.save(tmp_path)
        try:
            if path.exists():
                path.unlink()
            tmp_path.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp_path.replace(alt)
            print("Wrote (locked original)", alt.name)
            continue
        print("Wrote", path.name)

    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
