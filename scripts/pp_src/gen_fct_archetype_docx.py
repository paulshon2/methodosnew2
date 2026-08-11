# -*- coding: utf-8 -*-
"""Generate Foucauldian–Critical Theory DOCX: 5 × 5 × 11 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from fct_coding_examples import (
    apparatus_rows,
    coding_proc_apparatus,
    coding_proc_formation,
    coding_proc_gov,
    formation_rows,
    governmentality_rows,
)
from fct_docx_profiles import METHOD_PROFILES
from fct_figures import figures_for_analysis
from fct_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-17.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\푸코·비판이론 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "푸코·비판이론 아키타입(Foucauldian–Critical Theory Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①제도적 문헌·정책자료·역사적 기록 수집 →②담론이 형성된 시대적 맥락 자료(계보학적 자료) 수집",
        [
            "1) 제도적 문헌·정책자료·역사적 기록 수집하기",
            "2) 담론이 형성된 시대적 맥락 자료(계보학적 자료) 수집하기",
        ],
    ),
    (
        "② 전처리",
        "①문헌 선정 및 시대순 배치 →②주요 개념·용어의 변화 추적을 위한 자료 정리",
        [
            "1) 문헌 선정 및 시대순 배치하기",
            "2) 주요 개념·용어의 변화 추적을 위한 자료 정리하기",
        ],
    ),
    (
        "③ 코딩",
        "①개념적·계보학적 코딩으로 권력/지식 장치 식별 →②담론형성체 및 규율기제 범주화 →③통치성 관련 개념 도출",
        [
            "1) 개념적·계보학적 코딩으로 권력/지식 장치 식별하기",
            "2) 담론형성체(discursive formation) 및 규율기제 범주화하기",
            "3) 통치성(governmentality) 관련 개념 도출하기",
        ],
    ),
    (
        "④ 통계기법",
        "해당없음 — 통계적 검정을 사용하지 않음",
        ["1) 해당없음 — 통계적 검정을 사용하지 않음"],
    ),
    (
        "⑤ 결과도출",
        "①권력-지식 관계의 계보학적 해석 →②역사적 단절과 연속성을 개념적으로 서술",
        [
            "1) 권력-지식 관계의 계보학적 해석하기",
            "2) 역사적 단절과 연속성을 개념적으로 서술하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "비판적 개념분석 및 계보학적 해석",
        ["▶ 산출물: 비판적 개념분석 및 계보학적 해석"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 0 and ti == 1:
        return [figs[0]] if figs else []
    if si == 2 and ti == 1:
        return [figs[1]] if len(figs) > 1 else []
    if si == 4 and ti == 1:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["코퍼스·아카이브", topic["participants"]],
        ["분석 모드", topic["mode"]],
        ["초점", topic["focus"]],
        ["분석 키", topic["analysis_key"]],
        ["중심 주장", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    n_docs = 25 + s % 20
    is_de = topic["analysis_key"].startswith("de_")

    if si == 0 and ti == 0:
        proc = [
            f"【제도·정책 아카이브】「{topic['title']}」에 대해 제도적 문헌·정책자료·역사기록을 수집한다.",
            f"【범위】코퍼스: {topic['participants']}. 목표 약 {n_docs}건(+).",
            f"【로그】출처·판본·접근일·키워드. ({method_name}/{method_en})",
        ]
        del_r = [
            ["ID", "유형", "시기", "출처", "비고"],
            ["D01", "법령/지침", "2010", "관보/부처", "핵심"],
            ["D02", "내부규정", "2015", "기관", ""],
            ["D03", "정책브리프", "2020", "위원회", ""],
            ["D04", "언론·민원", "2022", "공개자료", "대항담론"],
        ]
        return STAGES[0][2][0], proc, (["수집n", str(n_docs)], [["로그", "Y"]]), ["제도·정책 문헌 목록"], del_r

    if si == 0 and ti == 1:
        proc = [
            "【계보자료】담론이 형성된 시대적 맥락·선행규정·사건 자료를 보강한다.",
            f"【초점】{topic['focus']} 개념의 출현·변형 시점 후보를 표시한다.",
            "【우선】1차 기록 > 2차 해설. 급변 시점(개정·사건)을 연표에 표시.",
        ]
        del_r = [
            ["시점", "사건/개정", "담론변화", "자료"],
            ["T1", "초기 규정", "문제 출현", "D01"],
            ["T2", "개정1", "용어 교체", "D02"],
            ["T3", "사건X", "위기·강화", "D04"],
            ["T4", "최근", "디지털화/관리화", "D03"],
        ]
        return STAGES[0][2][1], proc, (["계보보강", "Y"], [["연표", "초안"]]), ["계보학적 맥락 자료표"], del_r

    if si == 1 and ti == 0:
        proc = [
            "【선정】포함·제외 기준을 문서화하고 최종 문헌 세트를 확정한다.",
            "【시대순】선정 문헌을 연표·폴더에 시대순 배치한다.",
            f"【층위】법–지침–현장문서 층을 구분. 「{topic['title']}」.",
        ]
        del_r = [
            ["문헌", "연도", "층위", "포함", "이유"],
            ["D01", "2010", "법/지침", "Y", "기초규범"],
            ["D02", "2015", "기관", "Y", "실행규칙"],
            ["D05", "2018", "해설", "N", "2차과다"],
            ["D03", "2020", "정책", "Y", "전환점"],
        ]
        return STAGES[1][2][0], proc, (["선정", "확정"], [["시대순", "Y"]]), ["문헌선정·시대순 배치표"], del_r

    if si == 1 and ti == 1:
        proc = [
            f"【용어추적】핵심 개념({topic['focus']})의 출현·의미변화를 표로 정리한다.",
            "【연결】동의어·인접개념·대립개념을 맵핑한다.",
            "【메모】권력효과 가설은 보류 메모로만 남기고 성급히 확정하지 않는다.",
        ]
        del_r = [
            ["개념", "초기의미", "전환의미", "최근의미", "자료"],
            [topic["focus"][:10], "도덕/질서", "효율/관리", "자기책임", "D01→D03"],
            ["정상", "일탈대비", "지표화", "데이터정상", "D02"],
            ["책임", "의무", "위험관리", "자기계발", "D03"],
        ]
        return STAGES[1][2][1], proc, (["용어표", "v1"], [["맵", "Y"]]), ["개념·용어 변화 추적표"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc_apparatus(topic, method_name, s)
        del_r = apparatus_rows(topic, s)
        return STAGES[2][2][0], proc, (["장치코드", str(40 + s % 20)], [["계보메모", "Y"]]), ["권력/지식 장치 코드표"], del_r

    if si == 2 and ti == 1:
        proc = coding_proc_formation(topic, method_name)
        del_r = formation_rows(topic)
        return STAGES[2][2][1], proc, (["범주", "4"], [["장치연결", "Y"]]), ["담론형성체·규율기제 범주표"], del_r

    if si == 2 and ti == 2:
        proc = coding_proc_gov(topic, method_name)
        del_r = governmentality_rows(topic)
        return STAGES[2][2][2], proc, (["통치개념", "3–5"], [["반례", "점검"]]), ["통치성 개념 도출표"], del_r

    if si == 3 and ti == 0:
        proc = [
            "【원칙】통계적 가설검정을 사용하지 않는다.",
            "【이유】핵심은 권력-지식의 개념분석·계보학적 해석이다.",
            f"【주제】「{topic['title']}」.",
        ]
        del_r = [
            ["항목", "결정"],
            ["가설검정", "미사용"],
            ["주근거", "장치·담론·통치성 해석"],
            ["보조", "출현빈도 표는 선택(증명용 아님)"],
        ]
        return STAGES[3][2][0], proc, (["검정", "N/A"], [["원칙", "비판·계보"]]), ["통계 미사용 선언"], del_r

    if si == 4 and ti == 0:
        if is_de:
            proc = [
                f"【계몽 변증】「{topic['title']}」에서 합리성·해방 언어가 만드는 지배·물화 효과를 해석한다.",
                f"【주장】{topic['hypothesis']}",
                f"【방법】{method_name}({method_en}).",
            ]
            del_r = [
                ["축", "해방약속", "지배효과", "자료"],
                ["합리성", "효율·객관", "도구이성", "평가/지표"],
                ["선택", "자율", "표준화", "문화산업"],
                ["진보", "개선", "비판봉쇄", "기술담론"],
            ]
        else:
            proc = [
                f"【계보해석】권력-지식 관계가 「{topic['title']}」에서 어떻게 출현·변형했는지 서술한다.",
                f"【효과】규범이 가능하게/불가능하게 하는 주체·실천. 주장: {topic['hypothesis']}",
                f"【방법】{method_name}({method_en}).",
            ]
            del_r = [
                ["관계", "지식형태", "권력효과", "대표자료"],
                ["규범화", "지침·진단", "말할자격 제한", "D01"],
                ["가시화", "지표·기록", "서열·비교", "D02"],
                ["자기통치", "책임·자발", "내적규율", "D03"],
                ["초점", topic["focus"][:12], topic["hypothesis"][:18], "통합"],
            ]
        return STAGES[4][2][0], proc, (["해석", "초안"], [["증거", "연결"]]), ["권력-지식 계보 해석 요지"], del_r

    if si == 4 and ti == 1:
        proc = [
            "【단절·연속】무엇이 끊기고 무엇이 남는지를 개념적으로 서술한다.",
            f"【전환】초점={topic['focus']}의 급변 시점과 잔존 기제를 대비한다.",
            "【함의】비판적 개입·대안 담론의 여지와 한계를 명시한다.",
        ]
        del_r = [
            ["유형", "내용", "시기", "함의"],
            ["단절", "용어·제도 교체", "T2/T3", "새 개입권한"],
            ["연속", "시선·기록 논리", "전 기간", "규율 잔존"],
            ["변형", "디지털화", "T4", "가시성 강화"],
            ["잔존", "위계·정상화", "현재", "비판 지점"],
        ]
        return STAGES[4][2][1], proc, (["단절연속", "Y"], [["한계", "명시"]]), ["역사적 단절·연속 서술표"], del_r

    # si == 5
    proc = [
        "【산출물】비판적 개념분석 및 계보학적 해석 패키지.",
        f"【핵심】{topic['hypothesis']}",
        "【구성】문헌목록·연표·용어추적·장치코드·범주·통치성·해석문·단절/연속·한계.",
        f"【모드】{topic['mode']} · {method_name}({method_en}).",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["critical_genealogy.docx", "비판적 개념분석·계보해석", topic["title"][:20]],
        ["archive.xlsx", "문헌·연표", ""],
        ["codebook.xlsx", "장치·담론·통치성", ""],
        ["dispositif.png", "장치맵", ""],
        ["figA–C.png", "연표·장치·단절연속", ""],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "critical genealogy"]]), ["비판적 개념분석·계보 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」에 대해 제도·정책·계보 자료를 {topic['mode']}로 분석할 때, "
        f"{method_name}({method_en})의 장치·담론·통치성 절차로 "
        f"{topic['hypothesis']}을(를) 비판적으로 제시할 수 있는가?"
    )
    doc.add_paragraph(f"중심 주장: {topic['hypothesis']}")
    doc.add_paragraph(
        f"코퍼스 / 모드 / 초점: {topic['participants']} / {topic['mode']} / {topic['focus']}"
    )

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    doc.add_heading("【푸코·비판이론 분석 명세】", level=3)
    doc.add_paragraph("코퍼스·모드·초점 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "블록", "ID", "그룹", "파일", "항목", "발췌ID", "범주", "주제", "관계",
        "코드ID", "시점", "문헌", "개념", "유형", "축", "구분",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str, no: int) -> str:
    base = name_ko
    en_c = en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"푸코·비판이론 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 11"
        )
        doc.add_paragraph(
            "본 문서는 푸코·비판이론 아키타입의 6단계·11태스크에 따라 "
            "제도·정책 아카이브·계보자료·시대순 배치·용어추적·장치코딩·담론/규율범주·통치성·"
            "권력-지식 해석·단절/연속·비판적 개념분석 패키지를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"], no)
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
