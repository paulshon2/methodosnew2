# -*- coding: utf-8 -*-
"""Generate Forecasting–System Modeling DOCX: 7 × 5 × 12 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from forecast_sys_coding_examples import coding_proc_na, feature_codebook, metrics_rows
from forecast_sys_docx_profiles import METHOD_PROFILES
from forecast_sys_figures import figures_for_analysis
from forecast_sys_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-09.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\예측·시스템모형 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "예측·시스템모형 아키타입(Forecasting–System Modeling Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①시계열/패널/시뮬레이션 입력자료 수집 →②변수 정의 및 데이터 소스 확보(관측자료 또는 시뮬레이션 파라미터)",
        [
            "1) 시계열/패널/시뮬레이션 입력자료 수집하기",
            "2) 변수 정의 및 데이터 소스 확보(관측자료 또는 시뮬레이션 파라미터)하기",
        ],
    ),
    (
        "② 전처리",
        "①결측치 보간·이상치 제거 →②변수 정규화·스케일링 →③훈련/검증/테스트 데이터 분할",
        [
            "1) 결측치 보간·이상치 제거하기",
            "2) 변수 정규화·스케일링하기",
            "3) 훈련/검증/테스트 데이터 분할하기",
        ],
    ),
    (
        "③ 코딩",
        "해당없음",
        ["1) 해당없음(질적 코딩 미수행 · 피처·파라미터 명명만)"],
    ),
    (
        "④ 통계기법",
        "①기저모형(회귀/시계열) 적합 →②ML 예측 또는 ABM 구축 →③교차검증·성능지표(RMSE, MAE, 정확도)",
        [
            "1) 기저모형(회귀/시계열) 적합하기",
            "2) 머신러닝 예측모형(랜덤포레스트 등) 또는 행위자기반모형(ABM) 구축하기",
            "3) 교차검증 및 성능지표(RMSE, MAE, 정확도) 산출하기",
        ],
    ),
    (
        "⑤ 결과도출",
        "①모형 적합도·예측성능 비교 →②최적모형 선정 및 민감도 분석",
        [
            "1) 모형 적합도·예측성능 비교하기",
            "2) 최적모형 선정 및 민감도 분석하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "예측모형 및 시뮬레이션 결과",
        ["▶ 산출물: 예측모형 및 시뮬레이션 결과"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 3 and ti == 0:
        return [figs[0]] if figs else []
    if si == 3 and ti == 1:
        return [figs[1]] if len(figs) > 1 else []
    if si == 4 and ti == 0:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["자료 유형", topic["data_type"]],
        ["예측·시뮬 지평", topic["horizon"]],
        ["타깃/출력", topic["target"]],
        ["모형 계열", topic["model_family"]],
        ["분석 키", topic["analysis_key"]],
        ["중심 주장", topic["hypothesis"]],
    ]


def _is_sim(akey: str) -> bool:
    return akey.startswith(("bd_", "as_", "abm_"))


def _is_class(akey: str) -> bool:
    return akey in ("edu_drop", "pm_credit", "pm_readmit", "pm_churn") or "탈락" in akey


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    akey = topic["analysis_key"]

    if si == 0 and ti == 0:
        proc = [
            f"【자료】{topic['data_type']} 수집. 지평={topic['horizon']}, 타깃={topic['target']}.",
            f"【범위】「{topic['title']}」에 필요한 시계열/패널/시뮬 입력 범위를 확정.",
            "【로그】출처·기간·갱신일·라이선스를 수집 로그에 기록.",
            f"【방법】{method_name}({method_en}) 표준 입력 요건에 맞춤.",
        ]
        del_r = [
            ["소스", "유형", "기간", "비고"],
            ["공공/행정통계", topic["data_type"][:12], topic["horizon"], "버전고정"],
            ["로그·설문", "보조", "—", "동의"],
            ["문헌파라미터", "시뮬용", "—", "교정후보"],
        ]
        return STAGES[0][2][0], proc, (["소스수", str(3 + s % 3)], [["완결", "Y"]]), ["입력자료 수집로그"], del_r

    if si == 0 and ti == 1:
        proc = [
            f"【타깃】{topic['target']}을(를) 조작적으로 정의(단위·집계).",
            f"【인자】예측인자·외생변수 또는 시뮬 파라미터. 계열={topic['model_family']}.",
            "【소스구분】관측자료 vs 교정 파라미터를 코드북에 분리.",
            "【결측코드】NA/-999 등 규칙을 명시.",
        ]
        del_r = [
            ["변수", "역할", "단위", "소스"],
            ["y", "타깃", topic["target"][:16], "관측/출력"],
            ["x1..xk", "피처", "혼합", "관측"],
            ["theta", "파라미터", "비율/지연", "문헌·교정"],
        ]
        return STAGES[0][2][1], proc, (["변수정의", "확정"], [["사전", "data_dict"]]), ["변수·파라미터 사전"], del_r

    if si == 1 and ti == 0:
        miss = 2 + s % 6
        proc = [
            f"【결측】결측률≈{miss}%. 시계열 보간 / 패널 다중대체 / 시뮬 초기화 재시도 중 규칙 선택.",
            "【이상치】도메인 불가능·급변 포인트를 표시 후 제거 또는 윈저화.",
            "【보고】처리 전후 요약통계·n을 표로 남김.",
        ]
        del_r = [
            ["항목", "처리전", "처리후", "규칙"],
            ["n", str(1000 + s), str(980 + s % 20), "이상치제거"],
            ["결측%", f"{miss}%", f"{max(0, miss - 2)}%", "보간/대체"],
            ["이상치", str(10 + s % 8), "0–3", "윈저/삭제"],
        ]
        return STAGES[1][2][0], proc, (["규칙", "문서화"], [["재현", "스크립트"]]), ["결측·이상치 처리로그"], del_r

    if si == 1 and ti == 1:
        proc = [
            "【스케일】z-score / min-max / log1p 중 선택. 트리모형은 선택적.",
            "【누수금지】Scaler는 train에만 fit → val/test transform.",
            f"【역변환】보고 단위={topic['target']}로 되돌리기.",
        ]
        del_r = [
            ["변수", "변환", "fit셋", "비고"],
            ["연속피처", "StandardScaler", "train", ""],
            ["타깃", "log1p(선택)", "train", "역변환보고"],
            ["범주", "one-hot/target", "train", "희소주의"],
        ]
        return STAGES[1][2][1], proc, (["scaler", "train-only"], [["저장", "joblib"]]), ["스케일링 명세"], del_r

    if si == 1 and ti == 2:
        if _is_sim(akey):
            proc = [
                "【시뮬 설계】난수시드·반복(run)·워밍업 기간을 훈련/검증에 대응.",
                "【교정/검증】관측 시계열로 파라미터 교정, 홀드아웃 기간으로 검증.",
                "【금지】검증 구간의 정보로 규칙을 사후 끼워 맞추지 않음.",
            ]
            del_r = [
                ["용도", "설정", "n_runs", "시드"],
                ["교정", "관측정합", str(50 + s % 20), str(1000 + s)],
                ["검증", "홀드아웃", str(30 + s % 10), str(2000 + s)],
                ["정책실험", "시나리오", str(100), "고정목록"],
            ]
        else:
            proc = [
                "【분할】시간순 hold-out 또는 walk-forward (시계열/패널 누수 방지).",
                f"【비율】train/val/test ≈ 70/{15}/{15} (예시). 지평={topic['horizon']}.",
                "【기록】분할 인덱스·날짜 컷오프를 파일로 저장.",
            ]
            del_r = [
                ["셋", "비율", "기간/n", "용도"],
                ["train", "70%", str(700 + s), "적합"],
                ["val", "15%", str(150 + s % 30), "튜닝"],
                ["test", "15%", str(150 + s % 20), "최종보고"],
            ]
        return STAGES[1][2][2], proc, (["누수점검", "pass"], [["컷오프", "저장"]]), ["분할·실험설계"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc_na(topic, method_name)
        del_r = feature_codebook(topic)
        return STAGES[2][2][0], proc, (["코딩", "피처명명"], [["질적코딩", "N/A"]]), ["피처·파라미터 코드북"], del_r

    if si == 3 and ti == 0:
        if _is_sim(akey):
            proc = [
                "【기저】평균회귀·랜덤매칭·무정책 시나리오 등 단순 기준 런.",
                f"【적합】관측 패턴과 기준 런 출력을 정성·정량 비교. 타깃={topic['target']}.",
                f"【문서】{method_name} 기저 가정 명시.",
            ]
            del_r = [
                ["기저", "출력", "관측대비", "비고"],
                ["무정책/랜덤", "기준궤적", "부분일치", ""],
                ["평균유지", "평탄", "과소적합", ""],
            ]
        else:
            proc = [
                f"【기저】나이브(마지막값/계절)·선형회귀·ARIMA/ETS 적합. 타깃={topic['target']}.",
                "【진단】잔차 ACF, Ljung-Box, 잔차플롯.",
                f"【보고】train RMSE 기록. ({method_name})",
            ]
            del_r = [
                ["기저모형", "RMSE_train", "MAE_train", "비고"],
                ["Naive", f"{2.4 + s/100:.2f}", f"{1.9 + s/120:.2f}", ""],
                ["ARIMA/Reg", f"{1.8 + s/110:.2f}", f"{1.4 + s/130:.2f}", "선택후보"],
            ]
        return STAGES[3][2][0], proc, (["기저확정", "Y"], [["진단", "잔차"]]), ["기저모형 적합표"], del_r

    if si == 3 and ti == 1:
        if akey.startswith("bd_"):
            proc = [
                f"【SD】스톡-플로·피드백 루프로 「{topic['title']}」 구축.",
                "【교정】관측 시계열에 맞춰 지연·계수 교정.",
                "【실행】시나리오(정책) 런 설계.",
            ]
            del_r = [
                ["모듈", "내용", "상태"],
                ["재고/인력/구독 등", topic["target"][:14], "구축"],
                ["피드백루프", "+/− 표시", "문서화"],
                ["교정오차", f"RMSE≈{1.3 + s/100:.2f}", "수용"],
            ]
        elif _is_sim(akey):
            proc = [
                f"【ABM】에이전트 규칙·공간/네트워크로 「{topic['title']}」 구축.",
                "【스케일】에이전트 수·tick·시드 고정.",
                "【출력】거시 지표 집계 규칙 정의.",
            ]
            del_r = [
                ["요소", "설정", "비고"],
                ["agents", str(500 + s % 500), ""],
                ["ticks", topic["horizon"], ""],
                ["seed", str(42 + s % 10), "재현"],
                ["출력", topic["target"], "집계"],
            ]
        else:
            proc = [
                f"【ML】랜덤포레스트/GBM/Prophet 등 주모형. 계열={topic['model_family']}.",
                "【튜닝】검증셋·TimeSeriesSplit으로 하이퍼파라미터 탐색.",
                "【해석】중요도·부분의존(가능 시) 기록.",
            ]
            del_r = [
                ["모형", "핵심설정", "val지표"],
                ["RandomForest", f"trees={200 + s % 100}", f"RMSE={1.35 + s/200:.2f}"],
                ["GBM/XGB", "lr·depth 튜닝", f"RMSE={1.25 + s/220:.2f}"],
                ["선택후보", topic["model_family"][:16], "상위"],
            ]
        return STAGES[3][2][1], proc, (["주모형", "구축"], [["시드", "고정"]]), ["주모형·시뮬 구축로그"], del_r

    if si == 3 and ti == 2:
        kind = "sim" if _is_sim(akey) else ("class" if _is_class(akey) else "forecast")
        proc = [
            "【CV】k-fold 또는 time-series/walk-forward 교차검증.",
            "【지표】RMSE·MAE·MAPE 또는 Accuracy·AUC 등 과제 적합 지표.",
            "【분리보고】train / val / test 성능을 섞지 않음.",
        ]
        del_r = metrics_rows(s, kind)
        return STAGES[3][2][2], proc, (["CV", "완료"], [["test보고", "Y"]]), ["교차검증·성능표"], del_r

    if si == 4 and ti == 0:
        proc = [
            "【비교】기저 vs 주모형 vs 대안모형의 테스트 성능을 표·그림으로 비교.",
            f"【적합】「{topic['hypothesis']}」을(를) 성능·궤적 정합으로 평가.",
            "【시각】실측-예측, 학습곡선, 또는 시뮬 시나리오 궤적.",
        ]
        kind = "sim" if _is_sim(akey) else ("class" if _is_class(akey) else "forecast")
        del_r = metrics_rows(s + 1, kind)
        return STAGES[4][2][0], proc, (["승자후보", "1–2"], [["차이", "실질판단"]]), ["성능 비교표"], del_r

    if si == 4 and ti == 1:
        proc = [
            f"【선정】사전 기준(성능·해석가능성·운영비용)으로 최적모형 확정. 계열={topic['model_family']}.",
            "【민감도】파라미터 ±10%, 피처 제거, 시드/분할 변경 재실행.",
            "【한계】구조변화·외삽·가정 민감성을 명시.",
            f"【결론】{topic['hypothesis']}",
        ]
        del_r = [
            ["점검", "결과", "결론변화"],
            ["파라미터±10%", "방향 유지", "없음"],
            ["피처축소", "성능 소폭↓", "핵심피처 유지"],
            ["시드/분할", "순위 유지", "없음"],
            ["선정모형", topic["model_family"][:18], "최종"],
        ]
        return STAGES[4][2][1], proc, (["최종모형", "확정"], [["민감도", "통과"]]), ["최적모형·민감도"], del_r

    # si == 5
    proc = [
        "【산출물】예측모형 및 시뮬레이션 결과 패키지(모형파일·성능·시나리오·코드).",
        f"【핵심결과】{topic['hypothesis']}",
        f"【운영】재학습/재교정 주기, 모니터링 지표({topic['target']}).",
        f"【재현】{method_name}({method_en}) 환경·시드·버전 고정.",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["model.*", "학습가중/SD·ABM", topic["model_family"][:14]],
        ["metrics.xlsx", "RMSE·MAE·CV", "test"],
        ["forecast_or_runs.csv", "예측·시뮬출력", topic["horizon"]],
        ["figA–C.png", "궤적·구조·비교", ""],
        ["code/", "파이프라인", "재현"],
        ["data_dict.md", "변수사전", ""],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "forecast+sim"]]), ["예측·시뮬 결과 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」에서 {topic['data_type']}와 {topic['model_family']} "
        f"({method_name}/{method_en})로 {topic['hypothesis']}을(를) 예측·모의할 수 있는가?"
    )
    doc.add_paragraph(f"중심 주장: {topic['hypothesis']}")
    doc.add_paragraph(
        f"자료·지평·타깃: {topic['data_type']} / {topic['horizon']} / {topic['target']}"
    )

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    doc.add_heading("【예측·시스템모형 설계 명세】", level=3)
    doc.add_paragraph("자료형·지평·타깃·모형 계열 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "소스", "변수", "항목", "셋", "용도", "이름", "기저모형", "기저", "모형", "모듈", "요소",
        "시나리오", "점검", "파일",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str) -> str:
    base, en_c = name_ko, en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"예측·시스템모형 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 12"
        )
        doc.add_paragraph(
            "본 문서는 예측·시스템모형 아키타입의 6단계·12태스크(③코딩=해당없음)에 따라 "
            "입력자료·변수정의·전처리·분할·기저·ML/ABM·성능·최적모형·예측·시뮬 패키지를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"])
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
