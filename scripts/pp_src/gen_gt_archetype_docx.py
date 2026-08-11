# -*- coding: utf-8 -*-
"""Generate Grounded Theory DOCX: 6 × 5 × 18 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from gt_coding_examples import (
    axial_paradigm_rows,
    coding_proc_axial,
    coding_proc_compare,
    coding_proc_open,
    coding_proc_selective,
    open_code_rows,
    propositions,
    saturation_rows,
    selective_rows,
    storyline_paras,
)
from gt_docx_profiles import METHOD_PROFILES
from gt_figures import figures_for_analysis
from gt_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-10.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\QL-008 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "QL-008 아키타입(Grounded Theory Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①이론적 표본추출 →②반구조화 심층면담·참여관찰 →③수집·분석 순환(포화까지) →④메모잉",
        [
            "1) 이론적 표본추출(theoretical sampling)로 초기 참여자 선정하기",
            "2) 반구조화 심층면담·참여관찰 실시하기",
            "3) 자료수집과 분석을 순환적으로 반복(포화 시까지 표본 추가)하기",
            "4) 메모잉(memoing)으로 초기 아이디어 기록하기",
        ],
    ),
    (
        "② 전처리",
        "①축어록 전사 →②전사본 정확성 검토 →③분석 소프트웨어 입력·사례별 정리",
        [
            "1) 면담 녹음의 축어록(verbatim) 전사하기",
            "2) 전사본 정확성 검토(재생 대조)하기",
            "3) 분석 소프트웨어(NVivo 등)에 자료 입력 및 사례별 파일 정리하기",
        ],
    ),
    (
        "③ 코딩",
        "①개방코딩 →②축코딩(패러다임) →③선택코딩(핵심범주) →④지속적 비교·이론적 포화",
        [
            "1) 개방코딩(open coding): 문장·구 단위로 초기 개념(in-vivo code) 부여하기",
            "2) 축코딩(axial coding): 속성·차원 범주화 및 패러다임 모형 연결하기",
            "3) 선택코딩(selective coding): 핵심범주 선정 후 통합하기",
            "4) 지속적 비교분석과 이론적 포화 확인하기",
        ],
    ),
    (
        "④ 통계기법",
        "①통계적 검정 원칙적 미사용 →②코드·범주 빈도 매트릭스·참여자 특성은 보조 가능",
        [
            "1) 통계적 검정은 원칙적으로 사용하지 않기",
            "2) 코드·범주 출현빈도 매트릭스·참여자 특성 기술통계는 보조적으로 제시하기",
        ],
    ),
    (
        "⑤ 결과도출",
        "①패러다임 도식 →②이야기 윤곽 →③문헌 대조(이론적 통합) →④실체이론 서술",
        [
            "1) 핵심범주 중심 범주 관계 다이어그램(패러다임 모형)으로 도식화하기",
            "2) 이야기 윤곽(story line) 작성하기",
            "3) 잠정이론을 문헌과 대조(이론적 통합)하기",
            "4) 실체이론(substantive theory)으로 서술하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "QL-008(실체이론) 모형 및 이론적 명제",
        ["▶ 산출물: QL-008(실체이론) 모형 및 이론적 명제"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 0 and ti == 2:
        return [figs[0]] if figs else []
    if si == 2 and ti == 1:
        return [figs[1]] if len(figs) > 1 else []
    if si == 3 and ti == 1:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["탐구 과정", topic["process"]],
        ["현장", topic["site"]],
        ["GT 전통", topic["tradition"]],
        ["핵심 초점", topic["core_focus"]],
        ["분석 키", topic["analysis_key"]],
        ["중심 주장(실체이론 방향)", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    n0 = 6 + s % 4

    if si == 0 and ti == 0:
        proc = [
            f"【이론적 표집】「{topic['process']}」에 정통한 초기 참여자 {n0}명부터 선정.",
            f"【기준】현장={topic['site']}, 초점={topic['core_focus']}. 이론적 이유로 포함·제외.",
            "【유연】분석 후 개념 공백에 따라 표집 변경(전수 고정 금지).",
            f"【윤리】동의·가명. ({method_name}/{method_en})",
        ]
        del_r = [
            ["ID", "선정이유(이론적)", "자료", "동의"],
            ["P01", f"{topic['core_focus']} 경험자", "면담", "Y"],
            ["P02", "동일과정·다른역할", "면담+관찰", "Y"],
            ["P03", "부정사례 후보", "면담", "Y"],
        ]
        return STAGES[0][2][0], proc, (["초기n", str(n0)], [["고정n", "N"]]), ["이론적 표집 로그"], del_r

    if si == 0 and ti == 1:
        proc = [
            f"【면담】반구조 가이드: 「{topic['title']}」 과정을 시간순·조건·전략·결과로 묻기.",
            f"【관찰】{topic['site']}에서 관련 장면·인공물 보완(가능 시).",
            "【탐침】‘그때 무엇이 달랐나?’ ‘누가 개입했나?’ 등.",
            "【기록】녹음 60–90분 + 필드노트.",
        ]
        del_r = [
            ["회차", "유형", "분", "핵심질문"],
            ["1", "면담", "75", "과정·전환점"],
            ["2", "관찰", "120", "규범·상호작용"],
            ["3", "후속면담", "45", "개념확인"],
        ]
        return STAGES[0][2][1], proc, (["가이드", "v1"], [["관찰", "병행"]]), ["면담·관찰 로그"], del_r

    if si == 0 and ti == 2:
        proc = [
            "【순환】수집→개방/축 코딩→메모→다음 표집 결정을 반복.",
            f"【추가】「{topic['core_focus']}」 속성·부정사례를 메울 참여자 추가.",
            "【중단】포화 근거가 문서화될 때까지(미리 n 확정 금지).",
        ]
        del_r = [
            ["라운드", "누적n", "추가이유", "새개념"],
            ["R1", str(n0), "초기", "많음"],
            ["R2", str(n0 + 4), "중재조건 공백", "일부"],
            ["R3", str(n0 + 6), "부정사례", "거의없음"],
        ]
        return STAGES[0][2][2], proc, (["순환", "3+"], [["포화", "근접"]]), ["순환수집 일지"], del_r

    if si == 0 and ti == 3:
        proc = [
            "【메모잉】코드 정의·비교·가설·다음 질문을 즉시 기록.",
            "【구분】코드메모 / 이론메모 / 운영메모.",
            f"【연결】이후 이야기 윤곽·명제 재료. 초점={topic['core_focus']}.",
        ]
        del_r = [
            ["메모ID", "유형", "요지", "다음행동"],
            ["M01", "코드", "‘버티기’≠수동성?", "속성비교"],
            ["M02", "이론", "중재=동료발화?", "표집추가"],
            ["M03", "운영", "관찰허가", "일정"],
        ]
        return STAGES[0][2][3], proc, (["메모수", str(20 + s % 15)], [["이론메모", "8+"]]), ["메모잉 샘플"], del_r

    if si == 1 and ti == 0:
        proc = [
            "【축어전사】추임새·침묵·겹말 표기.",
            "【관찰노트】면담과 별도 파일.",
            f"【주제】「{topic['title']}」.",
        ]
        del_r = [
            ["파일", "분량", "유형", "상태"],
            ["P01.docx", "12p", "면담", "초안"],
            ["P02.docx", "10p", "면담", "초안"],
            ["OBS01.md", "4p", "관찰", "초안"],
        ]
        return STAGES[1][2][0], proc, (["전사완료", f"{n0}"], [["관찰", "Y"]]), ["전사 체크"], del_r

    if si == 1 and ti == 1:
        proc = [
            "【재생대조】난해 구간 녹음 재청취·수정.",
            "【수정로그】일시·담당·구간.",
            "【회원확인】핵심 인용(가능 시).",
        ]
        del_r = [
            ["파일", "대조", "수정", "확정"],
            ["P01", "Y", "3구간", "Y"],
            ["P02", "Y", "1구간", "Y"],
            ["OBS01", "Y", "0", "Y"],
        ]
        return STAGES[1][2][1], proc, (["QC", "pass"], [["로그", "Y"]]), ["전사 정확성 검토"], del_r

    if si == 1 and ti == 2:
        proc = [
            "【입력】NVivo/MAXQDA/Atlas.ti 또는 체계적 폴더 구조.",
            "【정리】참여자별·자료유형별 케이스 파일.",
            "【버전】코드북·메모 버전 관리.",
        ]
        del_r = [
            ["케이스", "소스", "메모링크", "비고"],
            ["P01", "면담×2", "M01–M05", ""],
            ["P02", "면담+관찰", "M06–M09", ""],
            ["프로젝트", "코드북 v0.3", "—", method_name[:12]],
        ]
        return STAGES[1][2][2], proc, (["SW", "NVivo등"], [["백업", "Y"]]), ["사례 파일 구조"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc_open(topic, method_name, s)
        del_r = open_code_rows(topic, s)
        return STAGES[2][2][0], proc, (["초기코드", str(40 + s % 30)], [["인비보", "포함"]]), ["개방코딩 표"], del_r

    if si == 2 and ti == 1:
        proc = coding_proc_axial(topic, method_name)
        del_r = axial_paradigm_rows(topic)
        return STAGES[2][2][1], proc, (["패러다임", "초안"], [["근거", "발췌"]]), ["축코딩·패러다임"], del_r

    if si == 2 and ti == 2:
        proc = coding_proc_selective(topic, method_name)
        del_r = selective_rows(topic)
        return STAGES[2][2][2], proc, (["핵심범주", topic["core_focus"][:10]], [["통합", "Y"]]), ["선택코딩·핵심범주"], del_r

    if si == 2 and ti == 3:
        proc = coding_proc_compare(topic, method_name)
        del_r = saturation_rows(topic, s)
        return STAGES[2][2][3], proc, (["포화", "문서화"], [["추가표집", "중단/보완"]]), ["비교·포화 점검"], del_r

    if si == 3 and ti == 0:
        proc = [
            "【원칙】통계적 가설검정(t, ANOVA, 회귀 등)을 사용하지 않는다.",
            "【이유】GT 타당성은 범주관계·인용·포화·이론적 표집 논증에 있다.",
            f"【주제】「{topic['title']}」.",
        ]
        del_r = [
            ["항목", "결정"],
            ["가설검정", "미사용"],
            ["이론근거", "범주·인용·포화"],
            ["보고문", "‘유의하다’ 표현 금지"],
        ]
        return STAGES[3][2][0], proc, (["검정", "N/A"], [["원칙", "질적"]]), ["통계 미사용 선언"], del_r

    if si == 3 and ti == 1:
        proc = [
            "【보조】코드·범주 출현빈도 매트릭스, 참여자 특성 기술통계.",
            "【금지】빈도로 실체이론을 증명한다고 서술하지 않음.",
            f"【제시】핵심={topic['core_focus']} 관련 코드 분포(참고).",
        ]
        del_r = [
            ["범주", "출현(파일수)", "참여자수", "비고"],
            [topic["core_focus"][:10], str(18 + s % 5), str(n0 + 2), "핵심"],
            ["지지/중재", str(14 + s % 4), str(n0), ""],
            ["규범/맥락", str(12 + s % 3), str(n0 - 1), ""],
            ["인구특성", f"여{n0//2}/남{n0 - n0//2}", "연령대", "기술만"],
        ]
        return STAGES[3][2][1], proc, (["보조표", "Y"], [["증명용", "N"]]), ["빈도·특성 보조표"], del_r

    if si == 4 and ti == 0:
        proc = [
            f"【도식】핵심범주 「{topic['core_focus']}」 중심 패러다임/과정 다이어그램.",
            "【관계】각 화살표에 발췌ID·메모ID 연결.",
            f"【전통】{topic['tradition']}에 맞는 도식 표기.",
        ]
        del_r = axial_paradigm_rows(topic)
        return STAGES[4][2][0], proc, (["도식", "v1"], [["근거링크", "Y"]]), ["패러다임 도식 표"], del_r

    if si == 4 and ti == 1:
        proc = storyline_paras(topic) + [f"【작성】{method_name} 서술 톤 유지."]
        del_r = [
            ["구간", "내용"],
            ["시작조건", "정보·자원 공백"],
            ["전개", f"{topic['core_focus']} 전략 전개"],
            ["전환", "중재·규범 충돌"],
            ["결말", "적응/학습 또는 소진"],
        ]
        return STAGES[4][2][1], proc, (["윤곽", "초안"], [["인용", "삽입"]]), ["이야기 윤곽"], del_r

    if si == 4 and ti == 2:
        proc = [
            "【문헌대조】잠정이론 vs 선행연구(확증·확장·반박·공백).",
            "【금지】문헌을 자료에 끼워 맞추기.",
            f"【위치】「{topic['hypothesis']}」의 기여 한 문장.",
        ]
        del_r = [
            ["문헌영역", "관계", "본연구", "비고"],
            ["유사과정이론", "부분확증", "중재 강조 확장", ""],
            ["조직/돌봄/학습", "부분반박", "개인성격설 기각", ""],
            ["방법문헌", "정합", topic["tradition"][:14], ""],
        ]
        return STAGES[4][2][2], proc, (["통합메모", "Y"], [["기여", "명시"]]), ["이론적 통합 표"], del_r

    if si == 4 and ti == 3:
        proc = [
            f"【실체이론】{topic['site']} 맥락의 「{topic['process']}」에 대한 한정 이론으로 서술.",
            f"【중심】{topic['hypothesis']}",
            "【명제화】관계 문장으로 정리(아래 표).",
            f"【범위】일반이론 주장 자제. ({method_name})",
        ]
        del_r = propositions(topic)
        return STAGES[4][2][3], proc, (["이론초안", "Y"], [["명제", "3+"]]), ["실체이론·명제"], del_r

    # si == 5
    proc = [
        "【산출물】QL-008(실체이론) 모형 및 이론적 명제 패키지.",
        f"【핵심명제】{topic['hypothesis']}",
        "【구성】도식·이야기윤곽·코드북·포화근거·인용록·성찰.",
        f"【방법】{method_name}({method_en}) · {topic['tradition']}.",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["theory_model.png", "패러다임/과정", topic["core_focus"]],
        ["propositions.md", "이론적 명제", "P1–P3"],
        ["storyline.docx", "이야기 윤곽", ""],
        ["codebook.xlsx", "개방·축·선택", ""],
        ["saturation.md", "포화근거", ""],
        ["quotes.xlsx", "인용록", "가명"],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "substantive theory"]]), ["실체이론 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」의 「{topic['process']}」 과정은 어떻게 전개되며, "
        f"{method_name}({method_en})으로 {topic['hypothesis']}을(를) "
        f"실체이론으로 구축할 수 있는가?"
    )
    doc.add_paragraph(f"중심 주장: {topic['hypothesis']}")
    doc.add_paragraph(
        f"현장·전통·핵심: {topic['site']} / {topic['tradition']} / {topic['core_focus']}"
    )

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    doc.add_heading("【QL-008 설계 명세】", level=3)
    doc.add_paragraph("과정·현장·전통·핵심범주 후보 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "ID", "회차", "라운드", "메모ID", "파일", "케이스", "발췌ID", "패러다임 요소", "핵심범주",
        "항목", "범주", "구간", "문헌영역", "명제ID",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str) -> str:
    base, en_c = name_ko, en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"QL-008 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 18"
        )
        doc.add_paragraph(
            "본 문서는 QL-008 아키타입의 6단계·18태스크에 따라 "
            "이론적 표집·면담·순환수집·메모·전사·개방·축·선택코딩·포화·패러다임·이야기윤곽·실체이론·명제를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"])
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
