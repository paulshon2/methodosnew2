# -*- coding: utf-8 -*-
"""Generate Action–Participatory Research DOCX: 6 × 5 × 11 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from action_res_coding_examples import (
    coding_proc_outcome,
    coding_proc_theme,
    outcome_track_rows,
    reflection_rows,
    theme_code_rows,
)
from action_res_docx_profiles import METHOD_PROFILES
from action_res_figures import figures_for_analysis
from action_res_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-13.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\실행·QL-046 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "실행·QL-046 아키타입(Action–Participatory Research Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①현장 문제 확인 및 실행계획 수립 →②실행 사이클별 관찰·면담·설문·실행일지 수집(계획-실행-관찰-반성)",
        [
            "1) 현장 문제 확인 및 실행계획 수립하기",
            "2) 실행 사이클별로 관찰·면담·설문·실행일지 등 자료 수집하기",
        ],
    ),
    (
        "② 전처리",
        "①사이클별 자료를 시간순으로 정리 →②질적 자료 전사 및 양적 자료(설문 등) 코딩",
        [
            "1) 사이클별 자료를 시간순으로 정리하기",
            "2) 질적 자료 전사 및 양적 자료(설문 등) 코딩하기",
        ],
    ),
    (
        "③ 코딩",
        "①실행일지·관찰기록 주제코딩 →②실행성과 지표 범주화·변화과정 추적",
        [
            "1) 실행일지·관찰기록에 대한 주제코딩하기",
            "2) 실행성과 지표를 범주화하여 변화과정 추적하기",
        ],
    ),
    (
        "④ 통계기법",
        "①설문 등 기술통계 →②사전-사후 비교 시 t검정 등 보조 가능",
        [
            "1) 설문 등 양적자료 활용 시 기술통계(평균, 빈도) 산출하기",
            "2) 사전-사후 비교시 t검정 등 보조적으로 사용하기",
        ],
    ),
    (
        "⑤ 결과도출",
        "①사이클별 성찰 종합 →②개선안 재설계·다음 사이클 반영",
        [
            "1) 사이클별 성찰(reflection) 결과를 종합하기",
            "2) 개선안을 재설계하여 다음 사이클에 반영(순환적 성찰)하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "실행 개선안 및 변화보고서",
        ["▶ 산출물: 실행 개선안 및 변화보고서"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 0 and ti == 1:
        return [figs[0]] if figs else []
    if si == 2 and ti == 1:
        return [figs[1]] if len(figs) > 1 else []
    if si == 3 and ti == 1:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["현장", topic["site"]],
        ["문제", topic["problem"]],
        ["사이클", topic["cycles"]],
        ["개입", topic["intervention"]],
        ["분석 키", topic["analysis_key"]],
        ["중심 주장", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    n_surv = 20 + s % 30

    if si == 0 and ti == 0:
        proc = [
            f"【문제】현장「{topic['site']}」의 「{topic['problem']}」을(를) 관찰가능 지표로 정의.",
            f"【계획】개입={topic['intervention']}, 사이클={topic['cycles']}, 성공지표·일정·역할.",
            "【기준선】개입 전 상태(관찰·간단설문) 수집.",
            f"【윤리】동의·부담최소화. ({method_name}/{method_en})",
        ]
        del_r = [
            ["항목", "내용"],
            ["문제진술", topic["problem"]],
            ["목표", "문제 완화·개선"],
            ["개입", topic["intervention"]],
            ["지표", "행동/품질/인식"],
            ["일정", topic["cycles"]],
        ]
        return STAGES[0][2][0], proc, (["계획서", "v1"], [["기준선", "Y"]]), ["문제·실행계획서"], del_r

    if si == 0 and ti == 1:
        proc = [
            "【순환】각 사이클: 계획→실행→관찰→반성.",
            f"【자료】관찰·면담·설문(n≈{n_surv})·실행일지. 주제 「{topic['title']}」.",
            "【로그】사이클ID·날짜·충실도·사건.",
        ]
        del_r = [
            ["사이클", "실행요지", "자료", "성찰메모"],
            ["C1", "개입 착수", "일지+관찰", "저항 확인"],
            ["C2", "조정 실행", "설문+면담", "촉진요인"],
            ["C3", "강화/확산", "일지+지표", "잔여과제"],
        ]
        return STAGES[0][2][1], proc, (["사이클", topic["cycles"][:6]], [["다중자료", "Y"]]), ["사이클 수집 로그"], del_r

    if si == 1 and ti == 0:
        proc = [
            "【시간순】C1→C2→C3 폴더·인덱스로 정리.",
            "【대조】계획 vs 실제 실행 차이 표기.",
            "【완전성】누락 관찰·미회수 설문 표시.",
        ]
        del_r = [
            ["사이클", "일지", "관찰", "면담", "설문"],
            ["C1", "Y", "Y", "N", "Y"],
            ["C2", "Y", "Y", "Y", "Y"],
            ["C3", "Y", "Y", "Y", "Y"],
        ]
        return STAGES[1][2][0], proc, (["인덱스", "완료"], [["갭", "기록"]]), ["사이클별 자료철"], del_r

    if si == 1 and ti == 1:
        proc = [
            "【질】면담·성찰 녹음 축어 전사.",
            f"【양】설문 문항코딩·결측처리(n≈{n_surv}).",
            "【키】참여자/일시로 질·양 연결.",
        ]
        del_r = [
            ["자료", "처리", "상태"],
            ["면담전사", "축어", "확정"],
            ["설문", "Likert코딩", "분석용"],
            ["일지", "날짜정렬", "코딩대기"],
        ]
        return STAGES[1][2][1], proc, (["전사", "Y"], [["설문코딩", "Y"]]), ["전사·설문코딩 로그"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc_theme(topic, method_name, s)
        del_r = theme_code_rows(topic, s)
        return STAGES[2][2][0], proc, (["주제코드", "5+"], [["근거", "발췌"]]), ["주제코딩 표"], del_r

    if si == 2 and ti == 1:
        proc = coding_proc_outcome(topic, method_name)
        del_r = outcome_track_rows(topic, s)
        return STAGES[2][2][1], proc, (["지표범주", "5"], [["추적", "C1–C3"]]), ["성과·변화 추적표"], del_r

    if si == 3 and ti == 0:
        m1 = 3.2 + (s % 8) / 20
        proc = [
            f"【기술통계】설문 평균·SD·빈도. 예: 만족 M={m1:.2f}.",
            "【층화】집단·사이클별 요약.",
            "【병치】주제코딩·관찰과 함께 해석.",
        ]
        del_r = [
            ["문항/지표", "n", "M", "SD", "비고"],
            ["참여/만족", str(n_surv), f"{m1:.2f}", "0.7", "C3"],
            ["부담감", str(n_surv), f"{2.8 - (s % 5)/20:.2f}", "0.8", "↓바람직"],
            ["개입도움", str(n_surv), f"{3.5 + (s % 4)/20:.2f}", "0.6", ""],
        ]
        return STAGES[3][2][0], proc, (["응답률", f"{80 + s % 15}%"], [["결측", "처리"]]), ["기술통계 표"], del_r

    if si == 3 and ti == 1:
        tstat = 2.1 + (s % 10) / 10
        proc = [
            "【보조】사전–사후 대응표본 t검정(또는 비모수) — 해당 시.",
            f"【예시】t≈{tstat:.2f}, p≈{0.04 if tstat > 2 else 0.12:.2f} (예시수치).",
            "【주의】무선배정 없음 → 인과 단정 금지. 실행연구 보조증거.",
            f"【보고】{method_name} 맥락에서 질적 성찰과 통합.",
        ]
        del_r = [
            ["비교", "Pre M", "Post M", "t/p", "해석"],
            ["핵심지표", "3.1", "3.6", f"{tstat:.2f}/<.05", "보조지지"],
            ["한계", "—", "—", "비실험", "일반화 제한"],
        ]
        return STAGES[3][2][1], proc, (["검정", "보조"], [["주증거", "사이클성찰"]]), ["사전–사후 보조검정"], del_r

    if si == 4 and ti == 0:
        proc = [
            f"【종합】사이클별 성찰 통합. 주장 점검: {topic['hypothesis']}",
            "【삼각화】일지·관찰·설문·면담 일치/불일치.",
            "【학습】실무 지식 문장화.",
        ]
        del_r = reflection_rows(topic)
        return STAGES[4][2][0], proc, (["성찰", "종합"], [["증거", "삼각화"]]), ["사이클 성찰 종합"], del_r

    if si == 4 and ti == 1:
        proc = [
            f"【재설계】다음 사이클/제도화 개입 수정. 현재개입={topic['intervention']}.",
            "【유지】성공 요소를 역할·문서·루틴으로.",
            "【공유】이해관계자 피드백 반영.",
            f"【방법】{method_name}.",
        ]
        del_r = [
            ["구분", "내용"],
            ["유지", "효과 입증된 절차"],
            ["수정", "부담 큰 절차 간소화"],
            ["추가", "확산·훈련"],
            ["폐기", "효과 없는 요소"],
            ["다음계획", "C4 또는 제도화"],
        ]
        return STAGES[4][2][1], proc, (["개선안", "v1"], [["합의", "Y"]]), ["재설계·다음계획"], del_r

    # si == 5
    proc = [
        "【산출물】실행 개선안 및 변화보고서.",
        f"【핵심】{topic['hypothesis']}",
        "【구성】문제·계획·사이클증거·성찰·개선안·한계·지속계획.",
        f"【현장】{topic['site']} · {method_name}({method_en}).",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["change_report.docx", "변화보고서", topic["title"][:20]],
        ["improvement_plan.md", "실행 개선안", topic["intervention"][:16]],
        ["cycle_evidence.xlsx", "지표·코딩", topic["cycles"]],
        ["reflection.md", "성찰종합", ""],
        ["figA–C.png", "사이클·지표·사전사후", ""],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "improvement+report"]]), ["개선안·변화보고 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['site']}」의 「{topic['problem']}」에 대해 {topic['intervention']} "
        f"개입을 {topic['cycles']}로 순환할 때, {method_name}({method_en})으로 "
        f"{topic['hypothesis']}을(를) 증거 기반으로 개선할 수 있는가?"
    )
    doc.add_paragraph(f"중심 주장: {topic['hypothesis']}")
    doc.add_paragraph(
        f"현장·문제·개입: {topic['site']} / {topic['problem']} / {topic['intervention']}"
    )

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    doc.add_heading("【실행·QL-046 설계 명세】", level=3)
    doc.add_paragraph("문제·현장·사이클·개입 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "항목", "사이클", "코드", "지표범주", "문항/지표", "비교", "구분", "파일",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str, no: int) -> str:
    base = name_ko
    if no in (23, 51) and name_ko == "교육 실행연구 절차 절차":
        base = f"{name_ko}{no}"
    en_c = en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"실행·QL-046 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 11"
        )
        doc.add_paragraph(
            "본 문서는 실행·QL-046 아키타입의 6단계·11태스크에 따라 "
            "문제·실행계획·사이클수집·정리·주제코딩·성과추적·기술통계·보조t검정·성찰·개선안·변화보고서를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"], no)
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
