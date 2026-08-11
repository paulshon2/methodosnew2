# -*- coding: utf-8 -*-
"""Generate Discourse–Critical Discourse DOCX: 5 × 5 × 12 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from discourse_coding_examples import (
    coding_proc_inter,
    coding_proc_ling,
    coding_proc_strat,
    collocation_rows,
    intertext_ideo_rows,
    linguistic_rows,
    strategy_rows,
)
from discourse_docx_profiles import METHOD_PROFILES
from discourse_figures import figures_for_analysis
from discourse_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-14.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\담화·비판담화 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "담화·비판담화 아키타입(Discourse–Critical Discourse Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①분석대상 텍스트/발화자료 선정 →②현장 녹음 또는 기존 텍스트 자료 수집",
        [
            "1) 분석대상 텍스트/발화자료 선정(연설, 대화, 미디어 텍스트 등)하기",
            "2) 녹음이 있는 경우 현장 녹음 또는 기존 텍스트 자료 수집하기",
        ],
    ),
    (
        "② 전처리",
        "①발화단위 표기를 포함한 정밀전사 →②텍스트를 절/문장 단위로 분절",
        [
            "1) 발화단위 표기를 포함한 정밀전사하기",
            "2) 텍스트를 절/문장 단위로 분절하기",
        ],
    ),
    (
        "③ 코딩",
        "①언어적·화용적 단위 코딩 →②상호텍스트성·이데올로기 장치 →③담화전략 범주화",
        [
            "1) 언어적·화용적 단위별 코딩(어휘선택, 수사적 장치, 함축)하기",
            "2) 상호텍스트성·이데올로기적 장치 식별하기",
            "3) 담화전략(정당화, 타자화 등) 범주화하기",
        ],
    ),
    (
        "④ 통계기법",
        "①원칙적으로 통계 미사용 →②코퍼스 시 어휘빈도·연어 보조",
        [
            "1) 원칙적으로 통계 미사용하기",
            "2) 코퍼스 기반 담화분석의 경우 어휘빈도·연어(collocation) 빈도 분석 보조 활용하기",
        ],
    ),
    (
        "⑤ 결과도출",
        "①담화전략 주제별 통합 해석 →②사회적 맥락(권력·이데올로기) 연결 해석",
        [
            "1) 범주화된 담화전략을 주제별로 통합 해석하기",
            "2) 사회적 맥락(권력, 이데올로기)과 연결하여 해석하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "담화구조·이데올로기 해석 결과",
        ["▶ 산출물: 담화구조·이데올로기 해석 결과"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 2 and ti == 2:
        return [figs[0]] if figs else []
    if si == 4 and ti == 1:
        return [figs[1]] if len(figs) > 1 else []
    if si == 3 and ti == 1:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["코퍼스", topic["corpus"]],
        ["장르", topic["genre"]],
        ["분석 초점", topic["focus"]],
        ["분석 키", topic["analysis_key"]],
        ["중심 주장", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)

    if si == 0 and ti == 0:
        proc = [
            f"【선정】「{topic['corpus']}」({topic['genre']})를 분석 대상으로 확정.",
            f"【범위】초점={topic['focus']}. 전형·대비·경계 텍스트 포함.",
            f"【윤리】공개/비공개·동의. ({method_name}/{method_en})",
        ]
        del_r = [
            ["항목", "결정"],
            ["장르", topic["genre"]],
            ["코퍼스", topic["corpus"][:28]],
            ["포함", "관련 핵심 텍스트"],
            ["제외", "비관련·중복·저품질"],
        ]
        return STAGES[0][2][0], proc, (["선정", "Y"], [["목록", "작성"]]), ["분석대상 선정서"], del_r

    if si == 0 and ti == 1:
        proc = [
            "【수집】현장 녹음 또는 기존 텍스트/아카이브 확보.",
            "【메타】출처·날짜·화자·매체·편집여부 로그.",
            f"【주제】「{topic['title']}」.",
        ]
        del_r = [
            ["ID", "유형", "출처", "분량"],
            ["D01", "녹음/텍스트", "현장/DB", "—"],
            ["D02", "텍스트", "언론/문서", "—"],
            ["D03", "보조", "맥락자료", "—"],
        ]
        return STAGES[0][2][1], proc, (["수집완료", "Y"], [["메타", "완비"]]), ["자료 수집 로그"], del_r

    if si == 1 and ti == 0:
        proc = [
            "【정밀전사】발화단위·침묵·겹침·비언어 표기(대화자료).",
            "【기텍스트】문장·레이아웃 보존, OCR 오류 교정.",
            "【QC】난해 구간 재생/원문 대조.",
        ]
        del_r = [
            ["파일", "표기규약", "상태"],
            ["tr_D01", "발화단위+침묵", "확정"],
            ["tr_D02", "문장보존", "확정"],
            ["메모", "잡음/편집", "기록"],
        ]
        return STAGES[1][2][0], proc, (["전사", "Y"], [["규약", "문서화"]]), ["정밀전사 체크"], del_r

    if si == 1 and ti == 1:
        proc = [
            "【분절】절/문장(또는 분석단위)으로 분할·ID 부여.",
            f"【정렬】화자·차례와 맞춤. 초점={topic['focus']}.",
            "【산출】분절 테이블.",
        ]
        del_r = [
            ["단위ID", "화자/위치", "텍스트요지", "비고"],
            ["T01-01", "S1", "개시", ""],
            ["T01-02", "S2", "응답/반박", ""],
            ["T01-03", "S1", "정당화", "핵심"],
        ]
        return STAGES[1][2][1], proc, (["단위수", str(80 + s % 40)], [["ID체계", "Txx-yy"]]), ["분절 단위표"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc_ling(topic, method_name, s)
        del_r = linguistic_rows(topic, s)
        return STAGES[2][2][0], proc, (["코드유형", "LEX/RHET/IMP"], [["근거", "발췌"]]), ["언어·화용 코딩"], del_r

    if si == 2 and ti == 1:
        proc = coding_proc_inter(topic, method_name)
        del_r = intertext_ideo_rows(topic)
        return STAGES[2][2][1], proc, (["장치", "식별"], [["효과메모", "Y"]]), ["상호텍스트·이데올로기"], del_r

    if si == 2 and ti == 2:
        proc = coding_proc_strat(topic, method_name)
        del_r = strategy_rows(topic)
        return STAGES[2][2][2], proc, (["전략수", "4–6"], [["패턴", "공출현"]]), ["담화전략 범주"], del_r

    if si == 3 and ti == 0:
        proc = [
            "【원칙】통계적 가설검정은 사용하지 않는다.",
            "【이유】핵심 근거는 담화 기능·전략·이데올로기 해석이다.",
            f"【주제】「{topic['title']}」.",
        ]
        del_r = [
            ["항목", "결정"],
            ["가설검정", "미사용"],
            ["주근거", "코딩·발췌·맥락해석"],
            ["보조", "선택적 빈도/연어"],
        ]
        return STAGES[3][2][0], proc, (["검정", "N/A"], [["원칙", "질적"]]), ["통계 미사용 선언"], del_r

    if si == 3 and ti == 1:
        proc = [
            "【보조】코퍼스 규모가 충분하면 어휘빈도·연어 산출.",
            "【금지】빈도로 이데올로기·전략을 단독 증명하지 않음.",
            "【병치】질적 발췌와 함께 보고.",
        ]
        del_r = collocation_rows(s)
        return STAGES[3][2][1], proc, (["보조분석", "실시/생략"], [["해석우선", "질적"]]), ["빈도·연어 보조표"], del_r

    if si == 4 and ti == 0:
        proc = [
            f"【통합】전략·장치를 주제별로 묶어 「{topic['title']}」 해석.",
            f"【주장】{topic['hypothesis']}",
            "【예시】대표 연쇄 2–3개를 두껍게 제시.",
            f"【방법】{method_name}.",
        ]
        del_r = [
            ["주제묶음", "주요전략", "대표발췌", "해석"],
            ["편가르기", "타자화+결속", "「그들/우리」", "경계구성"],
            ["필연만들기", "정당화+자연화", "「불가피」", "대안봉쇄"],
            ["책임흐리기", "전가+수동", "상황탓", "행위자은폐"],
        ]
        return STAGES[4][2][0], proc, (["통합해석", "초안"], [["예시", "삽입"]]), ["전략 통합 해석"], del_r

    if si == 4 and ti == 1:
        proc = [
            "【맥락】권력·제도·이데올로기 장과 연결.",
            f"【효과】담화가 구성하는 주체·관계: {topic['focus']}.",
            "【위치성】연구자 해석의 한계·대안읽기.",
            f"【결론】{topic['hypothesis']}",
        ]
        del_r = [
            ["층위", "연결", "증거"],
            ["제도/매체", topic["genre"], "코퍼스 메타"],
            ["권력관계", "누가 말하며 누가 침묵하나", "차례·명명"],
            ["이데올로기", "자연화된 전제", "장치표"],
            ["사회효과", "포함/배제 재생산", "전략표"],
        ]
        return STAGES[4][2][1], proc, (["맥락해석", "Y"], [["한계", "명시"]]), ["사회·이데올로기 해석"], del_r

    # si == 5
    proc = [
        "【산출물】담화구조·이데올로기 해석 결과 패키지.",
        f"【핵심】{topic['hypothesis']}",
        "【구성】해석문·전략표·장치표·전사/분절·(선택)빈도·성찰.",
        f"【방법】{method_name}({method_en}).",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["interpretation.docx", "담화·이데올로기 해석", topic["title"][:20]],
        ["strategies.xlsx", "담화전략", ""],
        ["devices.xlsx", "상호텍스트·이데올로기", ""],
        ["transcript/", "정밀전사·분절", ""],
        ["collocation.xlsx", "보조빈도", "선택"],
        ["figA–C.png", "전략·층위·빈도", ""],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "discourse+ideology"]]), ["해석 결과 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」({topic['genre']})에서 {topic['focus']} 관련 담화는 "
        f"어떤 언어·전략·이데올로기 장치로 조직되며, {method_name}({method_en})으로 "
        f"{topic['hypothesis']}을(를) 해석할 수 있는가?"
    )
    doc.add_paragraph(f"중심 주장: {topic['hypothesis']}")
    doc.add_paragraph(f"코퍼스·장르·초점: {topic['corpus']} / {topic['genre']} / {topic['focus']}")

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    doc.add_heading("【담화·비판담화 설계 명세】", level=3)
    doc.add_paragraph("코퍼스·장르·초점 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "항목", "ID", "파일", "단위ID", "장치", "전략", "노드", "주제묶음", "층위",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str) -> str:
    base, en_c = name_ko, en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"담화·비판담화 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 12"
        )
        doc.add_paragraph(
            "본 문서는 담화·비판담화 아키타입의 6단계·12태스크에 따라 "
            "자료선정·수집·정밀전사·분절·언어화용코딩·이데올로기장치·담화전략·보조빈도·통합·사회해석 패키지를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"])
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
