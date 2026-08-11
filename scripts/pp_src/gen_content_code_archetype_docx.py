# -*- coding: utf-8 -*-
"""Generate Content Analysis–Coding DOCX: 6 × 5 × 14 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from content_code_coding_examples import (
    chisq_rows,
    codebook_rows,
    coding_proc_main,
    coding_proc_rel,
    coding_proc_resolve,
    coding_sheet_rows,
    disagreement_rows,
    freq_rows,
    pilot_rows,
    reliability_rows,
)
from content_code_docx_profiles import METHOD_PROFILES
from content_code_figures import figures_for_analysis
from content_code_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-11.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\TX-002·코딩 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "TX-002·코딩 아키타입(Content Analysis–Coding Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①분석대상 문헌/미디어/텍스트 모집단 정의 →②표본추출(전수 또는 확률표집)로 분석자료 수집",
        [
            "1) 분석대상 문헌/미디어/텍스트 모집단 정의하기",
            "2) 표본추출(전수 또는 확률표집)로 분석자료 수집하기",
        ],
    ),
    (
        "② 전처리",
        "①분석단위 설정 →②코드북·코딩지침서 개발 →③코더 훈련 및 예비코딩",
        [
            "1) 분석단위(단어, 문장, 기사 등) 설정하기",
            "2) 코드북(codebook) 및 코딩지침서 개발하기",
            "3) 코더 훈련 및 예비코딩(pilot coding) 실시하기",
        ],
    ),
    (
        "③ 코딩",
        "①코드북 기반 범주코딩(2인 이상 독립) →②불일치 논의·조정 →③코더간 신뢰도 산출",
        [
            "1) 코드북 기반 범주코딩 수행(2인 이상 독립 코딩 권장)하기",
            "2) 코더 간 불일치 사례 논의·조정하기",
            "3) 코더간 신뢰도(inter-coder reliability) 산출하기",
        ],
    ),
    (
        "④ 통계기법",
        "①코드 빈도·백분율 →②Cohen's Kappa 또는 Krippendorff's α →③필요시 카이제곱",
        [
            "1) 코드 빈도·백분율 산출(기술통계)하기",
            "2) Cohen's Kappa 또는 Krippendorff's alpha로 코더간 신뢰도 검정하기",
            "3) 필요시 카이제곱검정으로 범주 간 연관성 검정하기",
        ],
    ),
    (
        "⑤ 결과도출",
        "①범주별 패턴 요약 →②범주 간 관계 통계적 해석",
        [
            "1) 빈도·신뢰도 결과를 바탕으로 범주별 패턴 요약하기",
            "2) 범주 간 관계를 통계적으로 해석하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "범주·주제별 빈도표 및 해석",
        ["▶ 산출물: 범주·주제별 빈도표 및 해석"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 3 and ti == 0:
        return [figs[0]] if figs else []
    if si == 3 and ti == 1:
        return [figs[1]] if len(figs) > 1 else []
    if si == 4 and ti == 0:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["코퍼스", topic["corpus"]],
        ["분석단위", topic["unit"]],
        ["범주체계", topic["categories"]],
        ["표집", topic["sampling"]],
        ["분석 키", topic["analysis_key"]],
        ["중심 주장", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    n = 200 + s % 150

    if si == 0 and ti == 0:
        proc = [
            f"【모집단】{topic['corpus']}. 기간·언어·형식 경계 문서화.",
            f"【포함·제외】분석 대상/비대상 목록. 주제 「{topic['title']}」.",
            f"【접근】아카이브·DB·크롤 가능성·저작권. ({method_name}/{method_en})",
        ]
        del_r = [
            ["경계", "결정"],
            ["매체/출처", topic["corpus"][:24]],
            ["기간", "연구 proto 명시"],
            ["제외예", "광고·중복·비관련"],
            ["윤리", "공개자료/개인정보"],
        ]
        return STAGES[0][2][0], proc, (["모집단정의", "Y"], [["문서", "protocol"]]), ["모집단 정의서"], del_r

    if si == 0 and ti == 1:
        proc = [
            f"【표집】{topic['sampling']}.",
            f"【목표n】약 {n} 단위({topic['unit']}). 대체규칙 명시.",
            "【로그】ID·출처·수집일 기록.",
        ]
        del_r = [
            ["단계", "n", "방법", "비고"],
            ["표집틀", str(n * 2), "목록구축", ""],
            ["표본", str(n), topic["sampling"][:16], ""],
            ["최종분석", str(n - s % 10), "QC통과", ""],
        ]
        return STAGES[0][2][1], proc, (["표본n", str(n)], [["전수여부", "표집에따름"]]), ["표집·수집 로그"], del_r

    if si == 1 and ti == 0:
        proc = [
            f"【분석단위】{topic['unit']}로 고정.",
            "【복수코딩】한 단위 다코드 허용 여부 결정.",
            "【경계예시】애매 단위 처리 규칙.",
        ]
        del_r = [
            ["항목", "결정"],
            ["단위", topic["unit"]],
            ["복수코드", "허용/상호배타"],
            ["분할규칙", "문장경계 등"],
        ]
        return STAGES[1][2][0], proc, (["단위확정", "Y"], [["지침반영", "Y"]]), ["분석단위 명세"], del_r

    if si == 1 and ti == 1:
        proc = [
            f"【코드북】범주={topic['categories']}. 정의·포함·제외·예시.",
            "【지침서】결정트리·FAQ·개정절차.",
            f"【버전】v0.9. ({method_name})",
        ]
        del_r = codebook_rows(topic)
        return STAGES[1][2][1], proc, (["코드수", str(4 + s % 3)], [["버전", "v0.9"]]), ["코드북·지침"], del_r

    if si == 1 and ti == 2:
        proc = [
            "【훈련】코더 교육·질의응답·연습단위.",
            f"【파일럿】약 {30 + s % 20}단위 예비코딩 → 정의 개정.",
            "【목표】파일럿 κ≥.70(예시) 후 본코딩.",
        ]
        del_r = pilot_rows(s)
        return STAGES[1][2][2], proc, (["파일럿", "완료"], [["개정", "Y"]]), ["예비코딩·훈련 로그"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc_main(topic, method_name, s)
        del_r = coding_sheet_rows(topic, s)
        return STAGES[2][2][0], proc, (["이중%", f"{20 + s % 10}%"], [["코더", "2+"]]), ["본코딩 시트"], del_r

    if si == 2 and ti == 1:
        proc = coding_proc_resolve(topic, method_name)
        del_r = disagreement_rows()
        return STAGES[2][2][1], proc, (["불일치", str(8 + s % 7)], [["합의", "완료"]]), ["불일치·합의 로그"], del_r

    if si == 2 and ti == 2:
        proc = coding_proc_rel(topic, method_name)
        del_r = reliability_rows(s)
        return STAGES[2][2][2], proc, (["κ", "보고"], [["α", "보고"]]), ["코더간 신뢰도"], del_r

    if si == 3 and ti == 0:
        proc = [
            f"【빈도】합의코드 기준 범주별 n·%. 단위={topic['unit']}.",
            "【층화】시기·매체 등 교차빈도(해당 시).",
            f"【주장점검】{topic['hypothesis']}",
        ]
        del_r = freq_rows(topic, s)
        return STAGES[3][2][0], proc, (["총단위", str(n)], [["결측코드", "제외"]]), ["빈도·백분율 표"], del_r

    if si == 3 and ti == 1:
        proc = [
            "【검정】Cohen's κ 및/또는 Krippendorff's α 산출·보고.",
            "【해석】기준값과 비교하되 맥락적으로 판단.",
            "【범주별】문제 범주 식별.",
        ]
        del_r = reliability_rows(s + 3)
        return STAGES[3][2][1], proc, (["신뢰도", "양호/재검토"], [["표본", "이중코딩"]]), ["κ/α 검정표"], del_r

    if si == 3 and ti == 2:
        proc = [
            "【χ²】범주×시기/매체 등 연관성 검정(필요시).",
            "【효과크기】Cramer's V·잔차.",
            "【가정】기대빈도 점검. 불충분 시 Fisher/기술만.",
        ]
        del_r = chisq_rows(s)
        return STAGES[3][2][2], proc, (["검정", "실시/생략"], [["유의", "일부"]]), ["카이제곱 결과"], del_r

    if si == 4 and ti == 0:
        proc = [
            "【패턴】우세·소수·공백 범주, 시기 추세를 요약.",
            f"【신뢰도】κ/α가 해석을 지지하는지 명시. 주제 「{topic['title']}」.",
            "【예시】대표 단위 인용 2–3개.",
        ]
        del_r = [
            ["패턴", "관찰", "근거"],
            ["우세범주", topic["categories"].split("/")[0][:12], "빈도1위"],
            ["추세", "시기별 변화", "히트맵"],
            ["공백", "Other/미분류 낮음", "코드북"],
            ["신뢰", "κ 양호", "신뢰도표"],
        ]
        return STAGES[4][2][0], proc, (["요약", "Y"], [["예시인용", "Y"]]), ["범주 패턴 요약"], del_r

    if si == 4 and ti == 1:
        proc = [
            f"【관계해석】χ²·교차표로 「{topic['hypothesis']}」을(를) 평가.",
            "【실질】통계적 유의와 실질적 차이를 구분.",
            "【한계】표집·단위·코더·매체 편향.",
            f"【함의】{method_name} 관점의 결론 한 문단.",
        ]
        del_r = [
            ["질문", "결과", "해석"],
            ["주가설", "지지/부분/기각", topic["hypothesis"][:28]],
            ["교차연관", "일부유의", "Cramer's V 참고"],
            ["한계", "외적타당도", "코퍼스 특수성"],
        ]
        return STAGES[4][2][1], proc, (["결론", "작성"], [["한계", "명시"]]), ["관계 해석"], del_r

    # si == 5
    proc = [
        "【산출물】범주·주제별 빈도표 및 해석 패키지.",
        f"【핵심】{topic['hypothesis']}",
        "【첨부】코드북·신뢰도·코딩시트·표집로그.",
        f"【방법】{method_name}({method_en}).",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["freq_table.xlsx", "범주 n·%", topic["unit"]],
        ["interpretation.md", "패턴·관계해석", topic["hypothesis"][:30]],
        ["codebook.pdf", "정의·예시", "v1.0"],
        ["reliability.xlsx", "κ·α", ""],
        ["coding_sheet.csv", "단위×코드", ""],
        ["figA–C.png", "빈도·신뢰도·교차", ""],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "freq+interpretation"]]), ["빈도표·해석 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」에서 {topic['unit']} 단위·범주({topic['categories']}) "
        f"TX-002({method_name}/{method_en})으로 {topic['hypothesis']}을(를) "
        f"빈도·신뢰도 근거로 밝힐 수 있는가?"
    )
    doc.add_paragraph(f"중심 주장: {topic['hypothesis']}")
    doc.add_paragraph(
        f"코퍼스·단위·표집: {topic['corpus']} / {topic['unit']} / {topic['sampling']}"
    )

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    doc.add_heading("【TX-002·코딩 설계 명세】", level=3)
    doc.add_paragraph("코퍼스·단위·범주·표집 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "경계", "단계", "항목", "코드", "단위ID", "지표", "범주", "교차", "패턴", "질문", "파일", "요약",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str) -> str:
    base, en_c = name_ko, en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"TX-002·코딩 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 14"
        )
        doc.add_paragraph(
            "본 문서는 TX-002·코딩 아키타입의 6단계·14태스크에 따라 "
            "모집단·표집·분석단위·코드북·파일럿·본코딩·합의·신뢰도·빈도·κ/α·카이제곱·패턴해석 패키지를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"])
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
