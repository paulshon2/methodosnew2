# -*- coding: utf-8 -*-
"""Generate Phenomenology–Lived Experience DOCX: 8 × 5 × 14 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from phenom_coding_examples import (
    coding_proc_horiz,
    coding_proc_imag,
    coding_proc_mu,
    coding_proc_theme,
    essence_sample_paras,
    horizontalization_rows,
    imaginative_variation_rows,
    meaning_unit_rows,
    theme_cluster_rows,
)
from phenom_docx_profiles import METHOD_PROFILES
from phenom_figures import figures_for_analysis
from phenom_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-08.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\현상학·체험탐구 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "현상학·체험탐구 아키타입(Phenomenology–Lived Experience Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①경험을 가진 참여자 목적표집(purposive sampling) →②개방형 심층면담(1~2회, 60~90분) →③면담 중 판단중지(epoché/bracketing)",
        [
            "1) 경험을 가진 참여자 목적표집(purposive sampling)하기",
            "2) \"그 경험을 상세히 기술해달라\"는 개방형 질문의 심층면담(1~2회, 60~90분) 실시하기",
            "3) 면담 중 판단중지(epoché/bracketing)로 연구자 선입견 배제 노력하기",
        ],
    ),
    (
        "② 전처리",
        "①면담 전사 및 반복 정독 →②전체적 인상(sense of the whole) 파악 →③의미단위(meaning unit)로 텍스트 구획",
        [
            "1) 면담 전사 및 반복 정독하기",
            "2) 전체적 인상(sense of the whole) 파악하기",
            "3) 의미단위(meaning unit)로 텍스트 구획하기",
        ],
    ),
    (
        "③ 코딩",
        "①의미단위 추출 →②수평화(horizontalization) →③주제/구성요소 군집화 →④상상적 변형으로 본질적 구조 도출",
        [
            "1) 의미단위 추출하기",
            "2) 각 의미단위를 연구자의 언어로 변형(수평화, horizontalization)하기",
            "3) 변형된 의미를 주제(theme)/구성요소로 군집화하기",
            "4) 상상적 변형(imaginative variation)을 통해 본질적 구조 도출하기",
        ],
    ),
    (
        "④ 통계기법",
        "해당없음 — 통계적 검정을 사용하지 않음",
        ["1) 해당없음 — 통계적 검정을 사용하지 않음"],
    ),
    (
        "⑤ 결과도출",
        "①구조적·직물적 기술 작성 →②본질(essence)에 대한 일반적 기술로 종합",
        [
            "1) 주제들을 종합하여 경험의 구조적(structural)·직물적(textural) 기술 작성하기",
            "2) 개별사례 기술을 통합해 본질(essence)에 대한 일반적 기술로 종합하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "경험의 본질(essence) 기술문 및 구조적 서술",
        ["▶ 산출물: 경험의 본질(essence) 기술문 및 구조적 서술"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 0 and ti == 2:
        return [figs[0]] if figs else []
    if si == 2 and ti == 2:
        return [figs[1]] if len(figs) > 1 else []
    if si == 4 and ti == 0:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["체험 국면", topic["lived"]],
        ["참여자", topic["participants"]],
        ["전통·접근", topic["tradition"]],
        ["초점 차원", topic["focus"]],
        ["분석 키", topic["analysis_key"]],
        ["중심 주장(본질 가설)", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    n_p = 8 + s % 5

    if si == 0 and ti == 0:
        proc = [
            f"【목적표집】「{topic['lived']}」을(를) 직접 겪은 참여자만 모집. 목표 n≈{n_p}명({topic['participants']}).",
            f"【포함】체험 기간·언어·동의 가능. 초점={topic['focus']}.",
            "【제외】대리 진술만 있는 경우, 체험이 너무 먼 과거로 회상만 가능한 경우(연구설계에 따라 조정).",
            f"【모집】공고·소개·눈덩이. ({method_name}/{method_en})",
        ]
        del_r = [
            ["ID", "포함기준", "모집경로", "동의"],
            ["P01", f"{topic['lived'][:14]} 경험", "병원/커뮤니티", "Y"],
            ["P02", "동일 체험", "소개", "Y"],
            ["P03", "동일 체험", "공고", "Y"],
            ["제외예", "체험 없음", "—", "—"],
        ]
        return STAGES[0][2][0], proc, (["목표n", str(n_p)], [["확정", str(n_p - 1)]]), ["목적표집 로그"], del_r

    if si == 0 and ti == 1:
        proc = [
            f"【개방형】「{topic['title']}」 경험을 가능한 한 상세히 기술해 달라.",
            "【면담】1–2회, 60–90분. 몸·시간·공간·관계 차원을 따라 촉진 질문.",
            f"【전통】{topic['tradition']}에 맞게 해석 개입 강도를 조절(기술 vs 해석/IPA).",
            "【기록】녹음·비언어 메모. 후속 면담으로 모호 구절 확인.",
        ]
        del_r = [
            ["회차", "분", "핵심 촉진질문", "산출"],
            ["1", "75", "그때 몸/시간은?", "녹음+메모"],
            ["2", "60", "빠진 장면이 있나?", "보완전사"],
            ["가이드", "—", f"초점:{topic['focus']}", "프로토콜"],
        ]
        return STAGES[0][2][1], proc, (["회차", "1–2"], [["평균분", "70"]]), ["면담 프로토콜·로그"], del_r

    if si == 0 and ti == 2:
        proc = [
            "【epoché】면담 전 선입견 목록을 괄호에 넣는다(이론·진단·조언 유예).",
            "【중】참여자 말을 교정·라벨링하지 않는다. 유도 질문을 성찰일지에 기록.",
            f"【후】「{topic['title']}」해석이 내 가정에 끌렸는지 점검. ({method_name})",
        ]
        del_r = [
            ["시점", "선입견/가정", "괄호조치", "메모"],
            ["사전", "‘힘들면 우울’", "유예", "진단어 금지"],
            ["중", "빨리 주제화", "보류", "개방 유지"],
            ["사후", topic["hypothesis"][:20], "재검토", "과잉적합 주의"],
        ]
        return STAGES[0][2][2], proc, (["성찰일지", "Y"], [["유도질문", "기록"]]), ["판단중지·성찰 로그"], del_r

    if si == 1 and ti == 0:
        proc = [
            "【전사】축어 전사. 침묵·웃음·울음 표기.",
            "【정독】코딩 전 전체 2–3회 반복 정독.",
            f"【QC】핵심 구절 회원검토(가능 시). 주제 「{topic['title']}」.",
        ]
        del_r = [
            ["파일", "분량", "정독횟수", "상태"],
            ["P01.docx", "14p", "3", "확정"],
            ["P02.docx", "11p", "2", "확정"],
            ["P03.docx", "16p", "3", "보완중"],
        ]
        return STAGES[1][2][0], proc, (["전사완료", f"{n_p - 1}"], [["검수", "Y"]]), ["전사·정독 체크"], del_r

    if si == 1 and ti == 1:
        proc = [
            "【전체 인상】참여자별로 이야기 전체의 분위기·흐름을 한 단락으로 요약.",
            f"【초점】{topic['focus']}가 전체에서 어떻게 드러나는지 메모.",
            "【금지】이 단계에서 주제명을 확정하지 않는다(인상만).",
        ]
        del_r = [
            ["참여자", "전체 인상(한줄)", "감정분위기", "메모"],
            ["P01", "익숙한 몸이 깨지는 이야기", "불안·혼란", ""],
            ["P02", "관계에서 멀어지는 흐름", "고립", ""],
            ["P03", "시간이 멈춘 듯한 국면", "압도", ""],
        ]
        return STAGES[1][2][1], proc, (["인상문", str(n_p - 1)], [["확정주제", "아직N"]]), ["전체 인상 요약"], del_r

    if si == 1 and ti == 2:
        proc = [
            "【구획】뜻이 전환되는 지점에서 의미단위로 텍스트를 나눈다.",
            "【ID】P0k-MUxx 부여. 원문 위치를 남긴다.",
            f"【주의】「{topic['lived']}」 서술과 무관한 잡담은 보류 단위로.",
        ]
        del_r = [
            ["참여자", "단위수(예)", "보류", "비고"],
            ["P01", str(12 + s % 6), "2", ""],
            ["P02", str(10 + s % 5), "1", ""],
            ["P03", str(14 + s % 4), "3", ""],
        ]
        return STAGES[1][2][2], proc, (["구획규칙", "의미전환"], [["도구", "한글/NVivo"]]), ["의미단위 구획 로그"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc_mu(topic, method_name, s)
        del_r = meaning_unit_rows(topic, s)
        return STAGES[2][2][0], proc, (["단위총", str(40 + s % 20)], [["관련높음", "다수"]]), ["의미단위 추출표"], del_r

    if si == 2 and ti == 1:
        proc = coding_proc_horiz(topic, method_name)
        del_r = horizontalization_rows(topic)
        return STAGES[2][2][1], proc, (["수평화", "완료"], [["변형", "연구자언어"]]), ["수평화·변형 목록"], del_r

    if si == 2 and ti == 2:
        proc = coding_proc_theme(topic, method_name)
        del_r = theme_cluster_rows(topic)
        return STAGES[2][2][2], proc, (["주제수", "4–6"], [["교차사례", "Y"]]), ["주제 군집표"], del_r

    if si == 2 and ti == 3:
        proc = coding_proc_imag(topic, method_name)
        del_r = imaginative_variation_rows(topic)
        return STAGES[2][2][3], proc, (["본질후보", "2–4"], [["검증", "사례회귀"]]), ["상상적 변형·본질 판정"], del_r

    if si == 3 and ti == 0:
        proc = [
            "【해당없음】통계적 가설검정(t, ANOVA, 회귀 등)을 사용하지 않는다.",
            "【이유】본 아키타입의 타당성은 본질 기술의 설득력·인용·절차 투명성에 있다.",
            f"【보조(선택)】참여자 인구특성·주제 출현 건수만 기술. 주제 「{topic['title']}」.",
            "【금지】빈도로 본질을 ‘유의하게 증명’한다고 서술하지 않는다.",
        ]
        del_r = [
            ["항목", "내용"],
            ["통계검정", "미사용"],
            ["보조기술", f"n≈{n_p}, 여성/남성·연령대 요약"],
            ["주제출현", "T1–T5 건수(참고)"],
            ["해석원칙", "질적 본질 > 빈도"],
        ]
        return STAGES[3][2][0], proc, (["검정", "N/A"], [["보조표", "선택"]]), ["통계 미사용 선언"], del_r

    if si == 4 and ti == 0:
        proc = [
            f"【직물적】「{topic['title']}」에서 무엇이 어떻게 느껴졌는지 생생히 기술(인용 풍부).",
            "【구조적】그 체험을 가능하게 한 조건·관계·시간 구조를 기술.",
            f"【종합】주제 T1–T5를 직물·구조 축으로 배치. 초점={topic['focus']}.",
            f"【전통】{topic['tradition']}의 서술 톤 유지.",
        ]
        del_r = [
            ["층위", "내용 요지", "근거주제", "인용예"],
            ["직물적", "몸 공백·시간 지연의 결", "T1,T2", "「머리가 하얘…」"],
            ["구조적", "익숙한 세계 소격→재편", "T3,T5", "「남의 일 같…」"],
            ["관계", "고립/지지의 간격", "T4", "「말이 안 들…」"],
        ]
        return STAGES[4][2][0], proc, (["초안", "Y"], [["인용수", "12+"]]), ["직물적·구조적 기술 초안"], del_r

    if si == 4 and ti == 1:
        proc = essence_sample_paras(topic)[:3] + [
            f"【통합】개별 사례 기술을 넘어 일반적 본질 문장으로 종합. ({method_name})",
        ]
        del_r = [
            ["구분", "서술"],
            ["본질 문장", topic["hypothesis"]],
            ["공통구조", "소격 → 현전 변화 → 의미 재편"],
            ["변주", "강도·지지·지속의 개인차"],
            ["한계", "맥락 특수성·소표본"],
        ]
        return STAGES[4][2][1], proc, (["본질문", "확정초안"], [["변주표", "Y"]]), ["본질 일반 기술"], del_r

    # si == 5
    proc = [
        "【산출물】경험의 본질(essence) 기술문 및 구조적·직물적 서술 패키지.",
        f"【핵심】{topic['hypothesis']}",
        "【구성】방법(표집·면담·epoché)·분석(단위→수평화→주제→상상적변형)·결과·성찰·인용록.",
        f"【윤리】가명·회원검토 여부 명시. ({method_name}/{method_en})",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["essence.md", "본질 기술문", topic["hypothesis"][:40]],
        ["textural_structural.docx", "직물·구조 서술", topic["focus"]],
        ["theme_matrix.xlsx", "주제–단위", "T1–T5"],
        ["quotes.xlsx", "인용록", "가명"],
        ["reflexivity.md", "성찰·epoché", ""],
        ["protocol.pdf", "면담·표집", topic["tradition"]],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "essence"]]), ["본질 기술문 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」의 체험 구조는 무엇이며, {method_name}({method_en})의 "
        f"{topic['tradition']} 절차로 {topic['hypothesis']}을(를) 드러낼 수 있는가?"
    )
    doc.add_paragraph(f"중심 주장(본질): {topic['hypothesis']}")
    doc.add_paragraph(
        f"체험·참여자·초점: {topic['lived']} / {topic['participants']} / {topic['focus']}"
    )

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    doc.add_heading("【현상학·체험 설계 명세】", level=3)
    doc.add_paragraph("체험 국면·참여자·전통·초점 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "ID", "회차", "시점", "파일", "참여자", "단위ID", "주제", "요소", "층위", "구분", "항목", "파일",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str, no: int) -> str:
    # 동명: 현상학적연구방법(5) vs 현상학적 연구방법(53)
    base = name_ko
    if no in (5, 53) and "현상학" in name_ko:
        base = f"{name_ko}{no}"
    en_c = en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"현상학·체험탐구 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 14"
        )
        doc.add_paragraph(
            "본 문서는 현상학·체험탐구 아키타입의 6단계·14태스크(④통계=해당없음)에 따라 "
            "목적표집·심층면담·판단중지·전사·의미단위·수평화·주제군집·상상적변형·직물·구조·본질 기술문을 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"], no)
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
