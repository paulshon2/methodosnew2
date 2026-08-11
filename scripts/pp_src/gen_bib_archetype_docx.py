# -*- coding: utf-8 -*-
"""Generate Bibliometrics–Citation Analysis DOCX: 2 × 5 × 9 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from bib_coding_examples import (
    coding_na_rows,
    indicator_rows,
    network_rows,
    stats_proc_indicators,
    stats_proc_network,
)
from bib_docx_profiles import METHOD_PROFILES
from bib_figures import figures_for_analysis
from bib_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-20.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\계량서지·인용분석 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "계량서지·인용분석 아키타입(Bibliometrics–Citation Analysis Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①서지 데이터베이스(Web of Science, Scopus 등)에서 출판·인용 메타데이터 수집",
        [
            "1) 서지 데이터베이스(Web of Science, Scopus 등)에서 출판·인용 메타데이터 수집하기",
        ],
    ),
    (
        "② 전처리",
        "①저자명·기관명 정규화 및 중복제거 →②분석기간·분야 필터링",
        [
            "1) 저자명·기관명 정규화 및 중복제거하기",
            "2) 분석기간·분야 필터링하기",
        ],
    ),
    (
        "③ 코딩",
        "해당없음",
        ["1) 해당없음"],
    ),
    (
        "④ 통계기법",
        "①인용지수·h-지수 등 계량지표 산출 →②시계열 회귀분석 및 공저·인용 네트워크 분석(중심성 지표)",
        [
            "1) 인용지수·h-지수 등 계량지표 산출하기",
            "2) 시계열 회귀분석 및 공저·인용 네트워크 분석(중심성 지표)하기",
        ],
    ),
    (
        "⑤ 결과도출",
        "①계량지표 기반 영향력·연구동향 판단 →②시각화(네트워크 맵, 추세그래프)로 제시",
        [
            "1) 계량지표 기반 영향력·연구동향 판단하기",
            "2) 시각화(네트워크 맵, 추세그래프)로 제시하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "계량서지 지표 및 동향분석",
        ["▶ 산출물: 계량서지 지표 및 동향분석"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 4 and ti == 1:
        return figs[:2] if len(figs) >= 2 else figs
    if si == 3 and ti == 0:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["코퍼스·DB", topic["participants"]],
        ["분석 모드", topic["mode"]],
        ["초점", topic["focus"]],
        ["분석 키", topic["analysis_key"]],
        ["중심 주장", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    n_raw = 2500 + s % 2000
    is_cit = topic["analysis_key"].startswith("cit_")

    if si == 0 and ti == 0:
        db = "WoS+Scopus" if "WoS" in topic["participants"] or "병합" in topic["participants"] else "Scopus/WoS"
        proc = [
            f"【수집】「{topic['title']}」 — {db}에서 출판·인용 메타데이터를 수집한다.",
            f"【코퍼스】{topic['participants']}. 원시 가안 n≈{n_raw}.",
            f"【기록】검색식·다운로드일·필드목록. ({method_name}/{method_en})",
        ]
        del_r = [
            ["항목", "내용"],
            ["DB", db],
            ["기간", "예: 2015–2025"],
            ["검색식", "주제키워드+동의어(문서화)"],
            ["필드", "TI/AU/PY/SO/TC/DE/C1…"],
            ["원시n", str(n_raw)],
            ["초점", topic["focus"][:20]],
        ]
        return STAGES[0][2][0], proc, (["raw_n", str(n_raw)], [["검색식", "Y"]]), ["서지 메타데이터 수집 로그"], del_r

    if si == 1 and ti == 0:
        proc = [
            "【정규화】저자명·기관명 이표기를 통일하고 중복 레코드를 제거한다.",
            "【규칙】DOI 우선, 없으면 제목+연도(+저널) 키.",
            f"【로그】제거 건수·매핑표를 남긴다. 「{topic['title']}」.",
        ]
        removed = 80 + s % 120
        del_r = [
            ["단계", "건수", "규칙", "비고"],
            ["중복제거", str(removed), "DOI/TI+PY", ""],
            ["저자정규화", str(200 + s % 100), "ORCID/수작업매핑", ""],
            ["기관정규화", str(60 + s % 40), "표준기관명", ""],
            ["정제후n", str(n_raw - removed), "분석코퍼스", ""],
        ]
        return STAGES[1][2][0], proc, (["정제", "Y"], [["제거", str(removed)]]), ["정규화·중복제거 로그"], del_r

    if si == 1 and ti == 1:
        proc = [
            f"【필터】분석기간·문서유형·분야를 고정한다. 초점={topic['focus']}.",
            "【재현】필터 조건을 검색식과 함께 보관한다.",
            "【민감도】경계 연도 포함/제외 결과를 비교한다.",
        ]
        del_r = [
            ["필터", "설정", "비고"],
            ["연도", "2015–2025", "민감도 ±1년"],
            ["문서유형", "Article/Review", ""],
            ["언어", "English(+선택)", "DB편향 고지"],
            ["주제", topic["focus"][:18], ""],
            ["최종n", str(n_raw - 150 - s % 80), "분석셋"],
        ]
        return STAGES[1][2][1], proc, (["filter", "확정"], [["재현패키지", "Y"]]), ["기간·분야 필터 명세"], del_r

    if si == 2 and ti == 0:
        proc = [
            "【코딩】해당 없음 — 질적 코딩 단계를 적용하지 않는다.",
            "【대체】계량지표 산출·시계열·네트워크 분석으로 진행한다.",
            f"【주제】「{topic['title']}」.",
        ]
        del_r = coding_na_rows(topic)
        return STAGES[2][2][0], proc, (["coding", "N/A"], [["다음", "통계"]]), ["코딩 해당없음 선언"], del_r

    if si == 3 and ti == 0:
        proc = stats_proc_indicators(topic, method_name, s)
        del_r = indicator_rows(topic, s)
        return STAGES[3][2][0], proc, (["지표", "산출"], [["단위구분", "Y"]]), ["계량지표 요약표"], del_r

    if si == 3 and ti == 1:
        proc = stats_proc_network(topic, method_name)
        del_r = network_rows(topic, s)
        return STAGES[3][2][1], proc, (["시계열", "Y"], [["네트워크", "Y"]]), ["시계열·네트워크 분석 요지"], del_r

    if si == 4 and ti == 0:
        if is_cit:
            proc = [
                f"【인용해석】「{topic['title']}」에서 인용경로·속도·중개 구조를 바탕으로 영향력·동향을 판단한다.",
                f"【주장】{topic['hypothesis']}",
                f"【한계】DB편향·자기인용·시간창. ({method_name}/{method_en})",
            ]
            del_r = [
                ["판단축", "관찰", "함의"],
                ["중개경로", "리뷰/교재 경유", "직접계보 과대평가 주의"],
                ["인용속도", "조기고인용≠장기", "수명분석 필요"],
                ["군집차이", "자기인용·국제성", "저널군별 해석"],
                ["통합", topic["hypothesis"][:28], "지표+망"],
            ]
        else:
            proc = [
                f"【동향판단】계량지표로 「{topic['title']}」의 영향력·연구동향을 판단한다.",
                f"【주장】{topic['hypothesis']}",
                f"【한계】DB편향·용어변화·분야차이. ({method_name}/{method_en})",
            ]
            del_r = [
                ["판단축", "관찰", "함의"],
                ["성장", "출판량 증가", "주제 확장"],
                ["집중", "소수저널/국가", "다양성 한계"],
                ["키워드", "용어교체", "겉보기성장 분해"],
                ["통합", topic["hypothesis"][:28], "지표+추세"],
            ]
        return STAGES[4][2][0], proc, (["해석", "초안"], [["한계", "명시"]]), ["영향력·동향 판단 요지"], del_r

    if si == 4 and ti == 1:
        proc = [
            "【시각화】네트워크 맵·추세그래프로 핵심 결과를 제시한다.",
            "【캡션】기간·노드의미·크기/색 범례를 붙인다.",
            "【재현】그림 원데이터와 생성 스크립트/설정을 보관한다.",
        ]
        del_r = [
            ["그림", "내용", "파일"],
            ["추세", "연도별 출판/인용", "figA_trend.png"],
            ["네트워크", "공저/인용 맵", "figB_net.png"],
            ["지표", "상대 프로파일", "figC_ind.png"],
            ["읽기포인트", topic["focus"][:16], "캡션에 명시"],
        ]
        return STAGES[4][2][1], proc, (["viz", "Y"], [["캡션", "Y"]]), ["시각화 목록"], del_r

    # si == 5
    proc = [
        "【산출물】계량서지 지표 및 동향분석 패키지.",
        f"【핵심】{topic['hypothesis']}",
        "【구성】검색식·정제로그·필터·지표표·시계열/네트워크·시각화·해석·한계.",
        f"【모드】{topic['mode']} · {method_name}({method_en}).",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["bibliometric_report.docx", "지표·동향분석", topic["title"][:20]],
        ["search_string.txt", "검색식·다운로드일", ""],
        ["clean_log.xlsx", "정규화·중복·필터", ""],
        ["indicators.xlsx", "인용·h·생산성", ""],
        ["network.xlsx", "중심성·엣지", ""],
        ["figA–C.png", "추세·망·지표", ""],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "bibliometric trends"]]), ["계량서지 지표·동향 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」에 대해 {topic['mode']}로 WoS/Scopus 메타데이터를 분석할 때, "
        f"{method_name}({method_en})의 계량지표·시계열·네트워크 절차로 "
        f"{topic['hypothesis']}을(를) 동향분석으로 제시할 수 있는가?"
    )
    doc.add_paragraph(f"중심 주장: {topic['hypothesis']}")
    doc.add_paragraph(
        f"코퍼스·DB / 모드 / 초점: {topic['participants']} / {topic['mode']} / {topic['focus']}"
    )

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    doc.add_heading("【계량서지·인용분석 설계 명세】", level=3)
    doc.add_paragraph("코퍼스·DB·초점 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "블록", "ID", "그룹", "파일", "항목", "발췌ID", "범주", "주제", "관계",
        "단계", "필터", "단위", "분석", "판단축", "그림", "구분",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str, no: int) -> str:
    base = name_ko
    en_c = en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"계량서지·인용분석 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 9"
        )
        doc.add_paragraph(
            "본 문서는 계량서지·인용분석 아키타입의 6단계·9태스크에 따라 "
            "서지메타 수집·정규화·필터·(코딩 해당없음)·계량지표·시계열/네트워크·"
            "동향판단·시각화·지표·동향분석 패키지를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"], no)
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
