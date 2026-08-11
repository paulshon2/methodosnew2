# -*- coding: utf-8 -*-
"""Generate Qualitative Design–Naturalistic DOCX: 5 × 5 × 10 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from qdn_coding_examples import coding_proc_peer, coding_proc_strategy, peer_rows, strategy_rows
from qdn_docx_profiles import METHOD_PROFILES
from qdn_figures import figures_for_analysis
from qdn_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-16.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\질적설계·자연주의 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "질적설계·자연주의 아키타입(Qualitative Design–Naturalistic Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①연구설계(현상학/QL-008/QL-020 등)에 따라 면담·관찰·문서 등 질적자료 수집전략 결정 →②적합한 표집전략(목적표집, 이론적 표집 등) 선택",
        [
            "1) 연구설계(현상학/QL-008/QL-020 등)에 따라 면담·관찰·문서 등 질적자료 수집전략 결정하기",
            "2) 적합한 표집전략(목적표집, 이론적 표집 등) 선택하기",
        ],
    ),
    (
        "② 전처리",
        "①자료관리체계 구축(전사, 파일링) →②신뢰성(trustworthiness)확보 절차(감사기록, 삼각검증) 계획",
        [
            "1) 자료관리체계 구축(전사, 파일링)하기",
            "2) 신뢰성(trustworthiness)확보 절차(감사기록, 삼각검증) 계획하기",
        ],
    ),
    (
        "③ 코딩",
        "①선택한 설계 전통에 맞는 코딩전략 적용(주제적/개념적/이론적 코딩) →②코딩의 엄격성을 위한 동료검토(peer debriefing) 실시",
        [
            "1) 선택한 설계 전통에 맞는 코딩전략 적용(주제적/개념적/이론적 코딩)하기",
            "2) 코딩의 엄격성을 위한 동료검토(peer debriefing) 실시하기",
        ],
    ),
    (
        "④ 통계기법",
        "해당없음 —설계서 특성상 통계기법을 규정하지 않음(개별 연구방법에 따라 상이)",
        [
            "1) 해당없음 —설계서 특성상 통계기법을 규정하지 않음(개별 연구방법에 따라 상이)",
        ],
    ),
    (
        "⑤ 결과도출",
        "①신뢰성·전이가능성·확증가능성 기준에 따라 결과 해석 →②연구방법론적 정당성 논의",
        [
            "1) 신뢰성·전이가능성·확증가능성 기준에 따라 결과 해석하기",
            "2) 연구방법론적 정당성 논의하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "연구설계·방법론 지침",
        ["▶ 산출물: 연구설계·방법론 지침"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 0 and ti == 0:
        return [figs[0]] if figs else []
    if si == 4 and ti == 0:
        return [figs[1]] if len(figs) > 1 else []
    if si == 2 and ti == 1:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["참여자·현장", topic["participants"]],
        ["설계 모드", topic["mode"]],
        ["초점", topic["focus"]],
        ["분석 키", topic["analysis_key"]],
        ["중심 주장", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    n_p = 10 + s % 8
    is_mm = topic["analysis_key"].startswith("mm_")
    is_ni = topic["analysis_key"].startswith("ni_")
    is_gl = topic["analysis_key"].startswith("gl_")

    if si == 0 and ti == 0:
        proc = [
            f"【설계선택】「{topic['title']}」에 대해 설계모드={topic['mode']}를 채택하고 "
            f"질문–패러다임–자료의 정합을 표로 고정한다.",
            "【수집전략】면담·관찰·문서(및 혼합 시 양적 자료)의 주·보조 역할을 정한다.",
            f"【정당성】왜 이 전통인가: {topic['hypothesis'][:40]}… ({method_name}/{method_en})",
        ]
        if is_mm:
            del_r = [
                ["층", "자료", "역할", "통합지점"],
                ["양적", "설문/지표", "패턴·효과", "이상치→면담"],
                ["질적", "면담/FGI", "설명·과정", "범주→문항"],
                ["통합", "공동해석", "수렴·발산규칙", "사전합의"],
            ]
        elif is_ni:
            del_r = [
                ["자료", "수집방식", "침습도", "윤리메모"],
                ["관찰", "비참여/참여", "낮음", "공개성 점검"],
                ["면담", "상황면담", "중간", "과정적 동의"],
                ["문서", "현장문서", "낮음", "허가·인용"],
            ]
        else:
            del_r = [
                ["요소", "결정", "근거"],
                ["설계전통", topic["mode"], "질문 적합"],
                ["주자료", "면담/관찰/문서", topic["focus"][:14]],
                ["보조자료", "문서·메모", "삼각검증"],
                ["괄호·성찰", "성찰일지", "위치성"],
            ]
        return STAGES[0][2][0], proc, (["설계버전", "v0.9"], [["정합표", "Y"]]), ["수집전략·설계정합표"], del_r

    if si == 0 and ti == 1:
        samp = "이론적 표집" if "근거" in topic["mode"] else ("층화+목적" if is_mm else "목적표집")
        proc = [
            f"【표집전략】{samp}을(를) 선택. 참여자·현장: {topic['participants']}.",
            "【포함·제외】기준·다양성·충분성/포화 계획을 문서화한다.",
            f"【규모 가안】약 {n_p}(명/사례/현장단위). 모드={topic['mode']}.",
        ]
        del_r = [
            ["기준", "내용", "비고"],
            ["포함", "핵심 경험/현장 접근 가능", ""],
            ["제외", "의사소통 불가·위험 과다", ""],
            ["다양성", "유사·대비·경계 사례", "의도적"],
            ["표집명", samp, "논리명시"],
            ["윤리", "설명·동의·철회", "필수"],
        ]
        return STAGES[0][2][1], proc, (["표집", samp], [["충분성계획", "Y"]]), ["표집전략 명세"], del_r

    if si == 1 and ti == 0:
        proc = [
            "【자료관리】전사 규칙·파일명·폴더·버전·백업을 표준화한다.",
            "【메타】날짜·출처·동의상태·민감도 태그를 붙인다.",
            f"【적용】「{topic['title']}」 자료 세트에 파일링 프로토콜을 적용한다.",
        ]
        del_r = [
            ["항목", "규칙", "예"],
            ["파일명", "날짜_출처_버전", "2026-03-01_P03_v1"],
            ["전사", "축어+메모분리", "본문/연구자메모"],
            ["백업", "이중저장+암호화", "클라우드/HDD"],
            ["권한", "역할별 접근", "연구책임자"],
        ]
        return STAGES[1][2][0], proc, (["DM", "구축"], [["QC", "Y"]]), ["자료관리 체계표"], del_r

    if si == 1 and ti == 1:
        proc = [
            "【신뢰성계획】감사기록·삼각검증·(필요시)멤버체크·동료검토 일정을 사전 계획한다.",
            f"【초점】{topic['focus']}에 대한 주장을 어떤 절차 증거로 뒷받침할지 체크리스트화한다.",
            f"【방법】{method_name}({method_en}) 설계서에 신뢰성 섹션을 삽입한다.",
        ]
        del_r = [
            ["절차", "시기", "산출", "담당"],
            ["감사기록", "전 기간", "결정일지", "연구자"],
            ["삼각검증", "수집·분석", "자료교차표", "연구자"],
            ["동료검토", "코딩 후", "의사록", "동료"],
            ["멤버체크", "초안 후", "피드백", "참여자(선택)"],
        ]
        return STAGES[1][2][1], proc, (["trust plan", "Y"], [["체크리스트", "Y"]]), ["신뢰성 확보 계획"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc_strategy(topic, method_name, s)
        del_r = strategy_rows(topic, s)
        return STAGES[2][2][0], proc, (["전략", "적용"], [["코드북", "v0.9"]]), ["코딩전략 적용표"], del_r

    if si == 2 and ti == 1:
        proc = coding_proc_peer(topic, method_name)
        del_r = peer_rows(topic)
        return STAGES[2][2][1], proc, (["peer", "완료"], [["개정", "v1.0"]]), ["동료검토 의사록"], del_r

    if si == 3 and ti == 0:
        proc = [
            "【원칙】본 아키타입 설계서 특성상 통계기법을 일률 규정하지 않는다.",
            "【개별성】혼합·평가 등 양적 모듈이 있으면 해당 연구방법의 통계를 따른다.",
            f"【주제】「{topic['title']}」 — 주 산출은 설계·방법론 지침이다.",
        ]
        if is_mm:
            del_r = [
                ["구분", "결정"],
                ["질적 모듈 통계", "미규정(해당없음)"],
                ["양적 모듈", "설문 분석은 별도 프로토콜"],
                ["통합 해석", "수치≠의미 증명"],
            ]
        else:
            del_r = [
                ["항목", "결정"],
                ["가설검정", "설계서에서 미규정"],
                ["주근거", "설계정합·신뢰성·해석"],
                ["보조수치", "필요시 기술통계만(선택)"],
            ]
        return STAGES[3][2][0], proc, (["stats", "N/A"], [["설계중심", "Y"]]), ["통계 미규정 선언"], del_r

    if si == 4 and ti == 0:
        proc = [
            f"【4기준 해석】신뢰성·전이가능성·확증가능성으로 「{topic['title']}」 결과를 해석한다.",
            f"【증거】주장 「{topic['hypothesis']}」마다 자료·절차 증거를 연결한다.",
            "【전이】적용 가능 맥락과 적용 불가 맥락을 구분해 적는다.",
        ]
        del_r = [
            ["기준", "본연구 조치", "잔여한계"],
            ["신뢰성", "장기관여·삼각검증", "시간제약"],
            ["전이가능성", "두꺼운기술", "맥락특수"],
            ["확증가능성", "감사기록·성찰", "연구자위치"],
            ["의존가능성", "절차문서화", "팀변동"],
        ]
        return STAGES[4][2][0], proc, (["해석", "4기준"], [["증거표", "Y"]]), ["신뢰성 기준 해석표"], del_r

    if si == 4 and ti == 1:
        proc = [
            f"【정당성】{method_name}({method_en})로서 왜 {topic['mode']} 설계가 질문에 적합한지 논증한다.",
            "【대안】채택하지 않은 설계(예: 순수설문, 실험)와 비교해 한계를 명시한다.",
            f"【함의】초점={topic['focus']}에 대한 실무·후속연구 함의.",
        ]
        if is_gl:
            del_r = [
                ["용어", "정의요지", "오용주의"],
                ["신뢰성", "절차증거와 결합", "주관=신뢰로 혼동 금지"],
                ["표집", "논리+한계", "이름만 붙이기 금지"],
                ["포화", "속성출현", "숫자목표 혼동"],
            ]
        else:
            del_r = [
                ["논증축", "내용"],
                ["질문–설계 정합", topic["mode"]],
                ["자료–분석 정합", topic["focus"][:20]],
                ["대안설계 기각사유", "깊이/맥락/과정 필요"],
                ["한계", "표집·이전가능성"],
                ["중심주장", topic["hypothesis"][:30]],
            ]
        return STAGES[4][2][1], proc, (["정당성", "초안"], [["한계", "명시"]]), ["방법론적 정당성 요지"], del_r

    # si == 5
    proc = [
        "【산출물】연구설계·방법론 지침 패키지.",
        f"【핵심】{topic['hypothesis']}",
        "【구성】설계정합표·표집명세·자료관리·신뢰성계획·코딩전략·동료검토·4기준해석·정당성논증.",
        f"【모드】{topic['mode']} · {method_name}({method_en}).",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["design_guideline.docx", "연구설계·방법론 지침", topic["title"][:20]],
        ["alignment.xlsx", "질문–설계–자료 정합", ""],
        ["sampling.pdf", "표집전략 명세", ""],
        ["trust_plan.xlsx", "신뢰성 확보 계획", ""],
        ["code_strategy.xlsx", "코딩전략·동료검토", ""],
        ["figA–C.png", "정렬·신뢰성·흐름", ""],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "design guideline"]]), ["연구설계·방법론 지침 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」을(를) {topic['mode']} 설계로 추진할 때, "
        f"{method_name}({method_en})의 절차(수집전략·표집·신뢰성·코딩·정당성)로 "
        f"{topic['hypothesis']}을(를) 설계·방법론적으로 뒷받침할 수 있는가?"
    )
    doc.add_paragraph(f"중심 주장: {topic['hypothesis']}")
    doc.add_paragraph(
        f"참여자·현장 / 설계모드 / 초점: {topic['participants']} / {topic['mode']} / {topic['focus']}"
    )

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    doc.add_heading("【질적설계·자연주의 설계 명세】", level=3)
    doc.add_paragraph("질문–설계전통–자료–표집–신뢰성의 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "블록", "ID", "그룹", "파일", "항목", "발췌ID", "범주", "주제", "관계",
        "구분", "요소", "기준", "절차", "쟁점", "층", "자료", "용어", "논증축",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str, no: int) -> str:
    base = name_ko
    en_c = en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"질적설계·자연주의 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 10"
        )
        doc.add_paragraph(
            "본 문서는 질적설계·자연주의 아키타입의 6단계·10태스크에 따라 "
            "수집전략·표집·자료관리·신뢰성계획·코딩전략·동료검토·4기준해석·방법론 정당성·설계지침 패키지를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"], no)
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
