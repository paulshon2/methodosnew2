# -*- coding: utf-8 -*-
"""Generate Conversation Analysis DOCX: 2 × 5 × 9 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from ca_coding_examples import coding_proc_practice, coding_proc_turns, practice_rows, turn_rows
from ca_docx_profiles import METHOD_PROFILES
from ca_figures import figures_for_analysis
from ca_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-22.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\대화·QL-050 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "대화·QL-050 아키타입(Conversation Analysis Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①자연발생적 대화 상황을 음성/영상으로 녹화 →②녹화자료 수집 시 참여자 동의 및 맥락정보 기록",
        [
            "1) 자연발생적 대화 상황을 음성/영상으로 녹화하기",
            "2) 녹화자료 수집 시 참여자 동의 및 맥락정보 기록하기",
        ],
    ),
    (
        "② 전처리",
        "①정밀전사(Jefferson 표기법: 억양, 중첩발화, 침묵시간, 늘임 등 표기)",
        [
            "1) 정밀전사(Jefferson 표기법: 억양, 중첩발화, 침묵시간, 늘임 등 표기)하기",
        ],
    ),
    (
        "③ 코딩",
        "①순차분석으로 발화순서(turn-taking) 구조 코딩 →②인접쌍·수리·선호구조 식별",
        [
            "1) 순차분석(sequence analysis)으로 발화순서(turn-taking) 구조 코딩하기",
            "2) 인접쌍(adjacency pair), 수리(repair), 선호구조(preference organization) 식별하기",
        ],
    ),
    (
        "④ 통계기법",
        "해당없음 — 통계적 검정을 사용하지 않음",
        ["1) 해당없음 — 통계적 검정을 사용하지 않음"],
    ),
    (
        "⑤ 결과도출",
        "①사례별 시퀀스 구조를 상세 기술 →②반복되는 상호작용 규칙을 일반화하여 제시",
        [
            "1) 사례별 시퀀스 구조를 상세 기술하기",
            "2) 반복되는 상호작용 규칙을 일반화하여 제시하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "대화구조 분석결과",
        ["▶ 산출물: 대화구조 분석결과"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 0 and ti == 0:
        return [figs[0]] if figs else []
    if si == 2 and ti == 0:
        return [figs[1]] if len(figs) > 1 else []
    if si == 4 and ti == 1:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["참여자·장면", topic["participants"]],
        ["분석 모드", topic["mode"]],
        ["초점", topic["focus"]],
        ["분석 키", topic["analysis_key"]],
        ["중심 주장", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    n_clip = 8 + s % 6
    is_intro = topic["analysis_key"].startswith("cai_")

    if si == 0 and ti == 0:
        proc = [
            f"【자연대화 녹화】「{topic['title']}」 — {topic['participants']} 상황을 음성/영상으로 녹화한다.",
            f"【목표】분석 클립 약 {n_clip}개(각 1–10분).",
            f"【주의】연출·인터뷰가 아닌 자연발생 상호작용. ({method_name}/{method_en})",
        ]
        del_r = [
            ["클립ID", "길이", "참여자", "매체", "품질"],
            ["CL01", "4:20", "2인", "음성+영상", "양호"],
            ["CL02", "6:10", "3인", "음성", "중첩많음"],
            ["CL03", "3:45", "2인", "음성+영상", "양호"],
            ["CL04", "5:00", "2인+", "음성", "잡음주의"],
        ]
        return STAGES[0][2][0], proc, (["클립n", str(n_clip)], [["자연성", "Y"]]), ["녹화 클립 로그"], del_r

    if si == 0 and ti == 1:
        proc = [
            "【동의】녹화·전사·논문 인용 범위를 설명하고 서면/구두 동의를 받는다.",
            "【맥락】일시·장소·역할·직전사건·제도적 목적을 기록한다.",
            f"【익명】가명 규칙. 초점={topic['focus']}.",
        ]
        del_r = [
            ["항목", "내용"],
            ["동의", "녹화+전사+가명인용"],
            ["일시/장소", "기록"],
            ["역할", topic["participants"][:20]],
            ["철회권", "안내"],
            ["맥락메모", topic["focus"][:16]],
        ]
        return STAGES[0][2][1], proc, (["동의", "Y"], [["맥락카드", "Y"]]), ["동의·맥락정보 카드"], del_r

    if si == 1 and ti == 0:
        proc = [
            "【Jefferson 정밀전사】억양·중첩([)·침묵(0.4)·늘임(:)·강세 등을 표기한다.",
            "【줄번호】분석 인용을 위해 줄번호를 부여한다.",
            f"【QC】난해 구간 재생 대조. 「{topic['title']}」.",
        ]
        del_r = [
            ["기호", "의미", "예"],
            ["[", "중첩 시작", "A: [yeah"],
            ["(0.4)", "침묵(초)", "갭/포즈"],
            [":", "늘임", "we::ll"],
            ["↑↓", "억양", "질문끝↑"],
            ["(( ))", "비언어/주석", "((gaze))"],
        ]
        return STAGES[1][2][0], proc, (["전사", "Jefferson"], [["줄번호", "Y"]]), ["정밀전사 기호·QC"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc_turns(topic, method_name, s)
        del_r = turn_rows(topic, s)
        return STAGES[2][2][0], proc, (["턴코딩", "Y"], [["시퀀스구간", str(8 + s % 6)]]), ["턴/순차 코딩표"], del_r

    if si == 2 and ti == 1:
        proc = coding_proc_practice(topic, method_name)
        del_r = practice_rows(topic)
        return STAGES[2][2][1], proc, (["인접쌍·수리·선호", "식별"], [["비선호표시", "Y"]]), ["인접쌍·수리·선호 식별표"], del_r

    if si == 3 and ti == 0:
        proc = [
            "【원칙】통계적 가설검정을 사용하지 않는다.",
            "【이유】핵심은 전사 근거에 기반한 시퀀스 기술과 규칙 일반화이다.",
            f"【주제】「{topic['title']}」.",
        ]
        del_r = [
            ["항목", "결정"],
            ["가설검정", "미사용"],
            ["주근거", "Jefferson전사+시퀀스"],
            ["보조", "현상 출현 횟수(선택, 증명용 아님)"],
        ]
        return STAGES[3][2][0], proc, (["검정", "N/A"], [["원칙", "CA"]]), ["통계 미사용 선언"], del_r

    if si == 4 and ti == 0:
        if is_intro:
            proc = [
                f"【사례기술】입문 수준에서 대표 시퀀스를 줄단위로 기술한다. 「{topic['title']}」.",
                "【근거】각 주장 옆에 전사 줄번호를 붙인다.",
                f"【방법】{method_name}({method_en}).",
            ]
        else:
            proc = [
                f"【사례기술】대표 시퀀스 구조를 상세 기술한다. 「{topic['title']}」.",
                f"【주장점검】{topic['hypothesis']}",
                f"【방법】{method_name}({method_en}).",
            ]
        del_r = [
            ["시퀀스", "줄범위", "구조요지", "핵심자원"],
            ["Ex1", "12–25", "FPP→지연→SPP", "침묵·완화"],
            ["Ex2", "40–55", "요청–거절", "계정(account)"],
            ["Ex3", "60–72", "수리연쇄", "자기수리"],
            ["초점", "—", topic["focus"][:14], topic["hypothesis"][:18]],
        ]
        return STAGES[4][2][0], proc, (["사례기술", "Y"], [["줄근거", "Y"]]), ["시퀀스 상세 기술 요지"], del_r

    if si == 4 and ti == 1:
        proc = [
            "【일반화】반복되는 상호작용 규칙을 명제 형태로 제시한다.",
            f"【핵심규칙】{topic['hypothesis']}",
            "【범위】적용 제도/장면·예외·한계를 명시한다.",
        ]
        del_r = [
            ["규칙ID", "일반화 명제", "근거사례", "예외"],
            ["R1", "비선호는 지연·완화로 조직", "Ex1/Ex2", "긴급상황"],
            ["R2", "확인쌍 완결이 정보전달 담보", "Ex3", "시간압박"],
            ["R3", topic["hypothesis"][:36], "다수", "문화·제도차"],
        ]
        return STAGES[4][2][1], proc, (["규칙", "2–4"], [["한계", "명시"]]), ["상호작용 규칙 일반화"], del_r

    # si == 5
    proc = [
        "【산출물】대화구조 분석결과 패키지.",
        f"【핵심】{topic['hypothesis']}",
        "【구성】동의·맥락·Jefferson전사·턴코딩·인접쌍/수리/선호·사례기술·규칙·한계.",
        f"【모드】{topic['mode']} · {method_name}({method_en}).",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["ca_report.docx", "대화구조 분석결과", topic["title"][:20]],
        ["transcripts/", "Jefferson 전사", "줄번호"],
        ["coding.xlsx", "턴·인접쌍·수리·선호", ""],
        ["collections.pdf", "사례 시퀀스", ""],
        ["figA–C.png", "타임라인·시퀀스·규칙", ""],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "CA structures"]]), ["대화구조 분석 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」의 자연발생 대화를 Jefferson 전사·순차분석으로 볼 때, "
        f"{method_name}({method_en})의 턴·인접쌍·수리·선호 분석으로 "
        f"{topic['hypothesis']}을(를) 대화구조로 제시할 수 있는가?"
    )
    doc.add_paragraph(f"중심 주장: {topic['hypothesis']}")
    doc.add_paragraph(
        f"참여자·장면 / 모드 / 초점: {topic['participants']} / {topic['mode']} / {topic['focus']}"
    )

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    doc.add_heading("【대화·QL-050 설계 명세】", level=3)
    doc.add_paragraph("장면·모드·초점 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "블록", "ID", "그룹", "파일", "항목", "발췌ID", "범주", "주제", "관계",
        "클립ID", "기호", "줄", "시퀀스", "규칙ID", "구분",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str, no: int) -> str:
    base = name_ko
    en_c = en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"대화·QL-050 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 9"
        )
        doc.add_paragraph(
            "본 문서는 대화·QL-050 아키타입의 6단계·9태스크에 따라 "
            "자연대화 녹화·동의/맥락·Jefferson 전사·턴/순차 코딩·인접쌍·수리·선호·"
            "시퀀스 기술·규칙 일반화·대화구조 분석 패키지를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"], no)
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
