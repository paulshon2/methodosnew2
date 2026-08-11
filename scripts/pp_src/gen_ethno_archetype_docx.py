# -*- coding: utf-8 -*-
"""Generate Ethnography–Participant Observation DOCX: 14 methods × 5 topics × 17 tasks."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from ethno_coding_examples import (
    coding_proc_descriptive,
    coding_proc_integrate,
    coding_proc_pattern,
    coding_proc_thick,
    concept_map_rows,
    descriptive_application,
    descriptive_codebook,
    pattern_rows,
    thick_description_sample,
)
from ethno_docx_profiles import METHOD_PROFILES
from ethno_figures import figures_for_analysis
from ethno_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-03.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\민족지·참여관찰 아키타입-연구방법별 샘플_정교화")

ARCHETYPE = "민족지·참여관찰 아키타입(Ethnography–Participant Observation Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①현장 접근(gatekeeper 섭외) 및 라포 형성 →②장기간 참여관찰 수행 →③비형식 면담·핵심정보제공자 면담 병행 →④사진·문서·유물 등 인공물 수집",
        [
            "1) 현장 접근(gatekeeper 섭외) 및 라포 형성하기",
            "2) 장기간 참여관찰(participant observation) 수행하기",
            "3) 비형식 면담·핵심정보제공자(key informant) 면담 병행하기",
            "4) 사진·문서·유물 등 인공물(artifact) 수집하기",
        ],
    ),
    (
        "② 전처리",
        "①매일 현장노트 작성(기술적 노트+분석적 메모 구분) →②녹음자료 전사 →③자료를 사건·장소·시간별로 정리 및 색인화",
        [
            "1) 매일 현장노트(fieldnotes) 작성(기술적 노트+분석적 메모 구분)하기",
            "2) 녹음자료 전사하기",
            "3) 자료를 사건·장소·시간별로 정리 및 색인화하기",
        ],
    ),
    (
        "③ 코딩",
        "①기술적 코딩으로 초기 범주 부여 →②패턴코딩으로 반복 주제·규범 식별 →③주제 간 관계 개념적 통합 →④두꺼운 기술로 문화적 의미 해석",
        [
            "1) 기술적 코딩(descriptive coding)으로 초기 범주 부여하기",
            "2) 패턴코딩(pattern coding)으로 반복되는 주제·규범 식별하기",
            "3) 주제 간 관계를 개념적으로 통합하기",
            "4) 두꺼운 기술(thick description)로 문화적 의미 해석하기",
        ],
    ),
    (
        "④ 통계기법",
        "①원칙적으로 통계 미사용 →②필요시 참여자 인구통계·행동빈도의 기초빈도/백분율만 보조적으로 제시",
        [
            "1) 원칙적으로 통계 미사용하기",
            "2) 필요시 참여자 인구통계·행동빈도의 기초빈도/백분율만 보조적으로 제시하기",
        ],
    ),
    (
        "⑤ 결과도출",
        "①현장노트·코딩 결과를 주제별로 통합 →②문화적 패턴·상징체계를 해석적으로 재구성 →③참여자 검토로 타당성 확보",
        [
            "1) 현장노트·코딩 결과를 주제별로 통합하기",
            "2) 문화적 패턴·상징체계를 해석적으로 재구성하기",
            "3) 참여자 검토(member checking)로 타당성 확보하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "문화기술적 사례기술(ethnography)",
        ["▶ 산출물: 문화기술적 사례기술(ethnography)"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc: Document, headers: list[str], rows: list[list[str]]):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = h
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc: Document, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(stage_i: int, task_i: int, figs: list[tuple[str, Path]]) -> list[tuple[str, Path]]:
    if stage_i == 0 and task_i == 1:
        return [figs[0]] if figs else []
    if stage_i == 2 and task_i == 2:
        return [figs[1]] if len(figs) > 1 else []
    if stage_i == 3 and task_i == 1:
        return [figs[2]] if len(figs) > 2 else figs[-1:]
    if stage_i == 5 and task_i == 0:
        return figs
    return []


def site_inventory(topic: dict, method_name: str) -> list[list[str]]:
    seed = _seed(topic["title"], method_name)
    akey = topic["analysis_key"]
    rows = [
        ["GK-01", "게이트키퍼", topic["access"].split("·")[0], "소개·허가", "동의Y"],
        ["KI-01", "핵심정보제공자", "현장 고참", "비형식면담", "가명"],
        ["KI-02", "핵심정보제공자", "신규/주변", "교차관점", "가명"],
        ["ART-01", "인공물", "안내문·게시물", f"ID-A{seed % 80}", "촬영동의"],
        ["ART-02", "인공물", "유니폼·공간배치", f"ID-A{seed % 50 + 20}", "스케치"],
    ]
    if akey.startswith("netn"):
        rows.append(["ON-01", "온라인자료", "공개스레드", f"URL-hash-{seed}", "공개만"])
    if akey.startswith("meta_"):
        rows = [
            ["LIT-01", "1차연구", "민족지 논문", "포함", "품질평가A"],
            ["LIT-02", "1차연구", "질적사례", "포함", "품질평가B"],
            ["LIT-03", "제외", "양적설문만", "제외", "설계불일치"],
        ]
    if akey.startswith("argo_"):
        rows = [
            ["TXT-01", "원전", "Argonauts 해당 장", "인용", "쪽수기입"],
            ["TXT-02", "지도·도판", "쿨라 순환도", "재현", "출처표기"],
            ["SEC-01", "2차", "방법론 해설", "대조", "현대비판"],
        ]
    return rows


def task_content(stage_i, task_i, topic, method_name, method_en):
    seed = _seed(method_name, topic["title"], stage_i, task_i)
    akey = topic["analysis_key"]
    hours = 40 + seed % 80
    days = 20 + seed % 40

    if stage_i == 0 and task_i == 0:
        proc = [
            f"【현장】「{topic['title']}」·장소/플랫폼 {topic['site']} · 기간 {topic['duration']}.",
            f"【게이트키퍼】접근 경로: {topic['access']}. 허가 범위(관찰·면담·촬영)·거절 가능 구역을 문서화.",
            f"【라포】초기 {7 + seed % 14}일간 정기 방문·작은 호혜. 목적 재설명. ({method_name}/{method_en})",
            "【윤리】동의서·익명·철회. 네트노그래피는 공개/비공개 구분, 메타민족지는 문헌 인용윤리.",
        ]
        inv = site_inventory(topic, method_name)
        return (
            STAGES[0][2][0],
            proc,
            (["항목", "내용"], [["site", topic["site"]], ["access", topic["access"]]]),
            ["접근·동의 로그 / 현장 진입표"],
            [["유형", "역할", "비고", "식별", "상태"]] + inv,
        )

    if stage_i == 0 and task_i == 1:
        if akey.startswith("meta_") or akey.startswith("argo_"):
            proc = [
                "【조정】본 주제는 문헌·고전 재독 중심 → ‘현장 장기체류’를 텍스트 정밀독해·재구성으로 대체.",
                "【활동】원전/1차연구 정독 로그, 개념·사건 색인. 참여관찰 시간 대신 독해 시간 기록.",
                "【명시】방법란에 대체 사유 기입.",
            ]
            del_r = [["활동", "시간/분량"], ["정독", f"{hours}h"], ["발췌카드", str(30 + seed % 20)], ["지도·도판", "재현"]]
        else:
            proc = [
                f"【참여관찰】목표 체류·접속 {topic['duration']} (누적 약 {hours}시간 / {days}회차).",
                f"【위치】관찰-참여 스펙트럼: {['관찰우세', '균형', '참여우세'][seed % 3]}. 변경 시 노트에 기록.",
                f"【초점】{topic['focus']}. 루틴·예외·갈등 장면을 우선 기록.",
                "【산출】일지·스케치맵(그림 A). 온라인은 접속 로그·스레드 캡처.",
            ]
            del_r = [["회차", "날짜", "시간", "초점장면"], ["#01", "2026-03-03", "2.5h", "입장·인사"], ["#12", "2026-03-20", "3.0h", "갈등/교대"], ["#28", "2026-04-10", "2.0h", "의례·이벤트"]]
        return STAGES[0][2][1], proc, (["지표", "값"], [["누적h", str(hours)]]), ["참여관찰 로그"], del_r

    if stage_i == 0 and task_i == 2:
        if akey.startswith("meta_"):
            proc = ["【면담 대체】저자·원자료에 접근 불가 시 ‘텍스트 내 참여자 목소리 인용’으로 대체.", "【명시】N/A 또는 문헌 인용 면담.", "【품질】1차연구의 인용 맥락을 왜곡하지 않음."]
            del_r = [["항목", "처리"], ["면담", "문헌내 인용으로 대체"], ["KI", "N/A"]]
        elif akey.startswith("argo_"):
            proc = ["【면담 대체】고전 텍스트의 대화·진술 인용을 ‘역사적 목소리’로 다룸.", "【주의】식민 권력·번역 왜곡을 성찰 메모에 적음.", "【산출】인용카드."]
            del_r = [["인용", "쪽", "주제"], ["항해 의례 서술", "ch.X", "기술·의례"], ["추장 연회", "ch.Y", "권위"]]
        else:
            proc = [
                f"【비형식 면담】현장 대화 {8 + seed % 10}건 + 반구조화 {3 + seed % 4}건.",
                "【핵심정보제공자】KI 2–4명. 관점 다양성(고참/신규, 중심/주변).",
                "【기록】직후 요지·가능 인용·관계 맥락. 동의 범위 확인.",
            ]
            del_r = [["ID", "유형", "분", "주제"], ["KI-01", "고참", "45", topic["focus"][:12]], ["KI-02", "신규", "30", "진입경험"], ["INF-05", "비형식", "12", "즉흥"]]
        return STAGES[0][2][2], proc, (["면담", "계획"], [["KI", "2–4"]]), ["면담 로그"], del_r

    if stage_i == 0 and task_i == 3:
        proc = [
            f"【인공물】초점({topic['focus']})과 연결된 물건·텍스트·이미지 수집.",
            "【메타데이터】날짜·장소·생산자·동의·파일명. 시각민족지는 촬영 가이드 준수.",
            "【목록】ART-ID를 현장노트와 링크.",
        ]
        del_r = [["ART-ID", "유형", "설명", "동의"], ["ART-01", "문서/게시", "안내·규칙", "Y"], ["ART-02", "공간/사진", "배치·간판", "Y"], ["ART-03", "디지털", "이모지·리액션 캡처", "공개"]]
        return STAGES[0][2][3], proc, (["건수", "목표"], [["인공물", str(10 + seed % 15)]]), ["인공물 목록"], del_r

    if stage_i == 1 and task_i == 0:
        proc = [
            "【매일】jottings → 확장 fieldnotes. 기술적 노트(관찰) / 분석적 메모(해석) / 성찰(위치성) 3칸.",
            f"【분량】관찰일 기준 확장노트 {1 + seed % 3}–{3 + seed % 3}쪽. 당일 작성 원칙.",
            "【샘플】아래 표처럼 한 장면을 세 칸으로 분리.",
        ]
        del_r = [
            ["시간", "기술적 노트(사실)", "분석적 메모", "성찰"],
            ["09:10", "신입이 문앞에서 3초 멈춤, KI가 끄덕임", "승인 의례?", "내가 먼저 들어가 흐름을 끊음"],
            ["12:40", "교대 시 ‘야간만’ 지시", "비공식 지식 전달", "녹음 요청을 망설임"],
        ]
        return STAGES[1][2][0], proc, (["규칙", "적용"], [["당일작성", "Y"], ["3칸분리", "Y"]]), ["현장노트 샘플"], del_r

    if stage_i == 1 and task_i == 1:
        if akey.startswith("netn"):
            proc = ["【전사 대체】스레드·채팅 로그를 타임스탬프와 함께 아카이브.", "【표기】삭제·수정된 글은 [deleted] 표시.", "【익명】닉네임 해시화."]
            del_r = [["파일", "기간", "상태"], ["thread_0312.md", "3일", "익명화"], ["emoji_log.csv", "1주", "집계용"]]
        elif akey.startswith("meta_") or akey.startswith("argo_"):
            proc = ["【전사 N/A】문헌·원전 인용 카드로 대체.", "【정확성】쪽수·판본 고정.", "【번역】번역본 사용 시 원문 대조 메모."]
            del_r = [["항목", "처리"], ["녹음전사", "N/A"], ["인용카드", f"{20 + seed % 15}매"]]
        else:
            proc = ["【전사】핵심 면담·현장음 선별 전사. 비언어 [웃음][침묵] 규약.", "【품질】난청 구간 (inaudible).", "【보관】원음/전사/익명본 분리."]
            del_r = [["파일", "분", "상태"], ["KI01.wav→.txt", "45", "완료"], ["FN15_ambient", "8", "부분"]]
        return STAGES[1][2][1], proc, (["전사", "선정"], [["우선", "KI면담"]]), ["전사·아카이브 로그"], del_r

    if stage_i == 1 and task_i == 2:
        proc = [
            "【색인】날짜·장소/채널·사건·인물가명·ART-ID·노트파일명 테이블.",
            "【사건카드】코딩 단위로 장면 분할.",
            "【버전】YYYYMMDD_현장_vN 파일명.",
        ]
        del_r = [["사건ID", "날짜", "장소", "인물", "노트"], ["E03", "03-03", topic["site"][:10], "KI-01", "FN-03"], ["E09", "03-20", "교대/스레드", "INF", "FN-09"], ["E15", "04-10", "의례", "다수", "FN-15"]]
        return STAGES[1][2][2], proc, (["색인", "완료"], [["사건카드", str(15 + seed % 20)]]), ["색인·사건목록"], del_r

    if stage_i == 2 and task_i == 0:
        proc = coding_proc_descriptive(topic, method_name, seed)
        cb = descriptive_codebook(topic, seed)
        del_r = descriptive_application(topic, akey)
        return STAGES[2][2][0], proc, (cb[0], cb[1:]), ["코드북 + 기술코딩 적용 예시"], del_r

    if stage_i == 2 and task_i == 1:
        proc = coding_proc_pattern(topic)
        pr = pattern_rows(topic)
        return STAGES[2][2][1], proc, (["패턴수", "값"], [["후보", "5"], ["채택", "4"]]), ["패턴코딩 표"], pr

    if stage_i == 2 and task_i == 2:
        proc = coding_proc_integrate(topic)
        cm = concept_map_rows(topic)
        return STAGES[2][2][2], proc, (["산출", "그림"], [["개념지도", "figB"]]), ["개념 통합표"], cm

    if stage_i == 2 and task_i == 3:
        proc = coding_proc_thick(topic)
        thick = thick_description_sample(topic)
        del_r = [["구분", "텍스트"], ["얇은 기술", thick[1]], ["두꺼운 기술", thick[2][:200] + "…"], ["해석", thick[3]]]
        return STAGES[2][2][3], proc, (["에피소드", "1–2"], [["연결코드", "PATT_*"]]), thick[:1], del_r

    if stage_i == 3 and task_i == 0:
        proc = [
            f"【원칙】{ARCHETYPE}는 추론통계로 문화를 검정하지 않는다.",
            "【방법란】‘본 연구는 통계적 가설검정 없이 참여관찰·코딩·두꺼운 기술로 해석한다.’",
            "【금지】코딩 빈도의 p값화, 소표본 t검정으로 ‘문화 차이 증명’ 서술 금지.",
        ]
        return STAGES[3][2][0], proc, (["선언", "Y"], [["통계미사용", "Y"]]), ["통계미사용 선언"], [["항목", "결정"], ["추론통계", "미사용"], ["대체", "민족지 서술"]]

    if stage_i == 3 and task_i == 1:
        proc = [
            "【보조】참여자 구성·관찰시간·기술코드 장면 빈도만 제시(그림 C).",
            f"【표】관찰 {hours}h, 면담 {8 + seed % 10}건, 사건카드 {15 + seed % 20}.",
            "【주의】빈도=중요도 아님. 의미는 인용·두꺼운 기술로.",
        ]
        del_r = [
            ["항목", "n", "%"],
            ["참여자(가명)", str(12 + seed % 10), "—"],
            ["여성·남성(예시)", "7 / 5", "—"],
            ["DESC_GREET 장면", str(10 + seed % 8), f"{30 + seed % 20}%"],
            ["DESC_SILENCE 장면", str(6 + seed % 7), f"{15 + seed % 15}%"],
        ]
        return STAGES[3][2][1], proc, (["보조통계", "Y"], [["그림C", "Y"]]), ["기초빈도표"], del_r

    if stage_i == 4 and task_i == 0:
        proc = [
            f"【통합】「{topic['title']}」주제로 노트·코드·인공물을 재배열.",
            "【절 구성】진입·일상루틴·갈등/의례·경계·연구자성찰.",
            "【인용】FN/KI/ART ID 병기.",
        ]
        return STAGES[4][2][0], proc, (["초고", "절"], [["5절", "Y"]]), ["주제별 통합 초고 목차"], [["절", "내용"], ["1", "진입과 라포"], ["2", "일상 루틴"], ["3", "규범·위계"], ["4", "갈등·의례"], ["5", "성찰"]]

    if stage_i == 4 and task_i == 1:
        proc = [
            f"【재구성】문화적 패턴: {topic['hypothesis']}",
            "【상징】인공물·몸짓·언어의 의미망(그림 B와 연결).",
            "【긴장】공식 수사 vs 암묵 실천.",
        ]
        return STAGES[4][2][1], proc, (["패턴", "채택"], [["PATT_*", "4"]]), ["문화 해석 요지"], [["패턴", "한줄"], ["암묵서열", "직책보다 눈치"], ["공식↔실제", "매뉴얼과 편법"], ["경계", "공간·멤버십"]]

    if stage_i == 4 and task_i == 2:
        if akey.startswith("meta_") or akey.startswith("argo_"):
            proc = ["【멤버체크 조정】문헌·고전은 ‘동료·지도교수 검토’ 또는 ‘원전 대조’로 대체.", "【기록】피드백 로그.", "【명시】참여자 검토 N/A 사유."]
            del_r = [["검토자", "의견", "반영"], ["동료", "인용 맥락 보강", "Y"], ["원전대조", "쪽수 수정", "Y"]]
        else:
            proc = ["【멤버체크】요지 1–2쪽을 KI·희망 참여자에게 제시.", "【기록】동의·수정·이견. 이견은 본문에 ‘다른 해석’으로.", "【윤리】거부 시 해당 인용 삭제·완화."]
            del_r = [["참여자", "반응", "조치"], ["KI-01", "대체로 동의, 직함 수정", "반영"], ["INF-03", "장면 식별 우려", "세부 일반화"], ["KI-02", "해석 이견", "이견 병기"]]
        return STAGES[4][2][2], proc, (["검토", "실시"], [["Y/대체", "Y"]]), ["멤버체크 로그"], del_r

    # deliverable
    proc = [
        f"【요약】「{topic['title']}」({topic['site']}, {topic['duration']}). 방법: {method_name}({method_en}). "
        f"주장: {topic['hypothesis']}. 산출: 문화기술적 사례기술.",
        "【본문】서론·현장·방법·주제절(두꺼운 기술)·논의·한계·윤리.",
        "【부록】노트샘플·코드북·인공물목록·동의·멤버체크·그림 A–C.",
        f"【버전】{date.today().isoformat()} · 재현 가능한 ID 체계 유지.",
    ]
    del_r = [
        ["유형", "파일", "내용"],
        ["본문", "ethnography_v1.docx", "사례기술"],
        ["부록", "codebook.xlsx", "DESC/PATT"],
        ["부록", "fieldnote_sample.pdf", "3칸노트"],
        ["그림", "FigA_map.png", "현장맵"],
        ["그림", "FigB_concepts.png", "개념도"],
        ["그림", "FigC_freq.png", "빈도"],
        ["윤리", "consent_log.pdf", "동의"],
        ["결론", "—", topic["hypothesis"]],
    ]
    return STAGES[5][2][0], proc, (["패키지", "v1.0"], [["날짜", date.today().isoformat()]]), ["출판용 민족지 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」현장에서 {method_name}({method_en}) 절차를 따를 때, "
        f"{topic['hypothesis']}을(를) 두꺼운 기술·코딩 근거로 제시할 수 있는가?"
    )
    doc.add_paragraph(f"중심 주장: {topic['hypothesis']}")
    doc.add_paragraph(f"현장·기간·접근: {topic['site']} / {topic['duration']} / {topic['access']}")
    doc.add_paragraph(f"관찰 초점: {topic['focus']}")

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    inv = site_inventory(topic, method_name)
    doc.add_heading("【현장·자료 진입 목록(샘플)】", level=3)
    doc.add_paragraph("게이트키퍼·핵심정보제공자·인공물(또는 문헌/원전) 구성. 태스크 1)·4)와 동일 체계로 확장한다.")
    add_table(doc, ["ID", "유형", "역할·내용", "비고", "상태"], inv)

    header_first = {
        "유형", "회차", "ID", "ART-ID", "시간", "파일", "사건ID", "코드", "패턴코드", "개념A", "구분", "항목", "절", "패턴", "참여자", "검토자", "활동", "인용",
    }

    for si, (stage_label, stage_sum, tasks) in enumerate(STAGES):
        doc.add_heading(stage_label, level=3)
        doc.add_paragraph(stage_sum)
        for ti, task_label in enumerate(tasks):
            _, proc_paras, proc_tbl, del_paras, del_rows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(task_label, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(task_label)

            step_notes = task_step_annotations(si, ti, proc_paras, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for step_i, (para, note) in enumerate(zip(proc_paras, step_notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({step_i}) 실행 절차 요지: ").bold = True
                preview = para
                head.add_run(preview)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc_paras:
                doc.add_paragraph(para)
            if proc_tbl[1]:
                add_table(doc, proc_tbl[0], proc_tbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in del_paras:
                doc.add_paragraph(dp)
            if del_rows:
                if del_rows[0][0] in header_first or (len(del_rows) > 1 and len(del_rows[0]) == len(del_rows[1])):
                    add_table(doc, del_rows[0], del_rows[1:])
                else:
                    ncols = max(len(r) for r in del_rows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in del_rows])

            for cap, fpath in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fpath, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str) -> str:
    base, en_clean = name_ko, en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_clean = en_clean.replace(ch, "-")
    return f"민족지·참여관찰 아키타입-{base}({en_clean}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no = m["no"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {m['name']}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 17"
        )
        doc.add_paragraph(
            "본 문서는 민족지·참여관찰 아키타입의 6단계·17태스크에 따라 "
            "현장 접근·참여관찰·현장노트·코딩(기술·패턴·통합·두꺼운 기술)·보조빈도·멤버체크·사례기술을 "
            "출판 수준 샘플로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, m["name"], prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(m["name"], prof["en"])
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
