# -*- coding: utf-8 -*-
"""Generate General Methodology DOCX: 1 × 5 × 7 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from grm_coding_examples import coding_proc, coding_rows
from grm_docx_profiles import METHOD_PROFILES
from grm_figures import figures_for_analysis
from grm_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-31.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\일반연구방법론 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "일반연구방법론 아키타입(General Methodology Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①연구주제에 따라 문헌·면담·설문 등 자료수집 전략을 선택 →②연구설계 단계에서 자료수집 계획 수립",
        [
            "1) 연구주제에 따라 문헌·면담·설문 등 자료수집 전략을 선택하기",
            "2) 연구설계 단계에서 자료수집 계획 수립하기",
        ],
    ),
    (
        "② 전처리",
        "①자료유형에 맞는 표준 전처리(전사 또는 코딩) 수행",
        ["1) 자료유형에 맞는 표준 전처리(전사 또는 코딩) 수행하기"],
    ),
    (
        "③ 코딩",
        "①질적자료는 주제코딩, 양적자료는 변수화 코딩을 연구설계에 따라 선택적으로 적용",
        ["1) 질적자료는 주제코딩, 양적자료는 변수화 코딩을 연구설계에 따라 선택적으로 적용하기"],
    ),
    (
        "④ 통계기법",
        "①연구주제에 따라 기술통계에서 추론통계까지 선택적으로 적용",
        ["1) 연구주제에 따라 기술통계에서 추론통계까지 선택적으로 적용하기"],
    ),
    (
        "⑤ 결과도출",
        "①연구설계 전반의 절차에 따라 결과를 통합적으로 서술",
        ["1) 연구설계 전반의 절차에 따라 결과를 통합적으로 서술하기"],
    ),
    (
        "⑥ 최종 산출물",
        "연구방법론 지침서",
        ["▶ 산출물: 연구방법론 지침서"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 0 and ti == 0:
        return [figs[1]] if len(figs) > 1 else []
    if si == 0 and ti == 1:
        return [figs[0]] if figs else []
    if si == 4 and ti == 0:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["대상", topic["participants"]],
        ["설계 모드", topic["mode"]],
        ["수집 전략", topic["strategy"]],
        ["초점", topic["focus"]],
        ["분석 키", topic["analysis_key"]],
        ["방법론 명제(초안)", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    mode = topic["mode"]
    n = 80 + s % 120

    if si == 0 and ti == 0:
        proc = [
            f"【전략선택】「{topic['title']}」에 대해 문헌·면담·설문 등 수집 전략을 선택한다.",
            f"【결정】{topic['strategy']} (모드={mode}).",
            f"【정합】연구질문·자원·윤리. ({method_name}/{method_en})",
        ]
        del_r = [
            ["옵션", "채택", "이유"],
            ["문헌", "Y", "이론·선행 틀"],
            ["면담", "Y" if "면담" in mode or "면담" in topic["strategy"] else "선택", "깊이"],
            ["설문", "Y" if "설문" in mode or "설문" in topic["strategy"] else "선택", "범위"],
            ["2차자료", "Y" if "2차" in mode else "N/선택", "가용성"],
            ["종합전략", topic["strategy"][:28], topic["focus"][:8]],
        ]
        return STAGES[0][2][0], proc, (["전략", "확정"], [["감사기록", "Y"]]), ["자료수집 전략 선택서"], del_r

    if si == 0 and ti == 1:
        proc = [
            "【수집계획】연구설계 단계에서 표집·도구·일정·윤리·품질기준을 포함한 계획을 수립한다.",
            f"【대상】{topic['participants']}. 예상 n≈{n}(양) / 면담은 설계에 따라.",
            f"【완료기준】프로토콜·동의·로그 준비 완료. 초점={topic['focus']}.",
        ]
        del_r = [
            ["계획항목", "내용", "상태"],
            ["표집", "모집단·표본·접근", "초안"],
            ["도구", "설문/면담가이드/문헌DB", "준비"],
            ["일정", "수집창·파일럿", "Y"],
            ["윤리", "동의·익명·보관", "Y"],
            ["품질", "파일럿·이중점검", "계획"],
        ]
        return STAGES[0][2][1], proc, (["계획", "v1"], [["윤리", "Y"]]), ["자료수집 계획서"], del_r

    if si == 1 and ti == 0:
        prep = "전사 중심" if ("면담" in mode and "설문" not in mode) else (
            "변수정리·코딩준비" if ("2차" in mode or "양적" in mode) else "전사 및/또는 설문정리"
        )
        proc = [
            f"【표준전처리】자료유형에 맞게 {prep}를 수행한다.",
            "【버전】원자료·처리본·로그를 분리 보관한다.",
            f"【QC】결측·중복·전사품질. 「{topic['title']}」.",
        ]
        del_r = [
            ["자료", "전처리", "완료", "비고"],
            ["문헌메모", "서지·요약표", "Y", ""],
            ["면담", "전사" if "면담" in topic["strategy"] or "면담" in mode else "N/A", "—", ""],
            ["설문/2차", "정리·코딩준비", "Y", ""],
            ["파일규칙", "ID·버전", "Y", ""],
        ]
        return STAGES[1][2][0], proc, (["전처리", "Y"], [["QC", "Y"]]), ["표준 전처리 점검표"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc(topic, method_name, s)
        del_r = coding_rows(topic, s)
        return STAGES[2][2][0], proc, (["코딩", "선택적용"], [["코드북", "Y"]]), ["주제코딩·변수화"], del_r

    if si == 3 and ti == 0:
        if "면담" in mode and "설문" not in mode:
            stats = "기술요약(빈도 보조) 위주 · 추론통계 최소"
            est = "—"
        elif "2차" in mode or "양적" in mode:
            stats = "기술통계 + 회귀/비교 등 추론통계"
            est = f"β≈{0.20 + (s % 10) * 0.02:.2f}"
        else:
            stats = "기술통계 + 필요 시 t/χ²/회귀"
            est = f"p≈{0.02 + (s % 8) * 0.01:.2f}"
        proc = [
            f"【통계선택】연구주제에 따라 {stats}를 적용한다.",
            f"【결과(도식)】{est}. 질 중심이면 주제 서술을 우선한다.",
            f"【근거】모드={mode}. ({method_name})",
        ]
        del_r = [
            ["분석", "적용", "산출", "비고"],
            ["기술통계", "Y", "M/SD·빈도", ""],
            ["추론통계", "선택", est, stats[:18]],
            ["질적요약", "설계에 따라", "주제표", ""],
            ["가정·한계", "명시", "—", ""],
        ]
        return STAGES[3][2][0], proc, (["분석", "선택"], [["문서화", "Y"]]), ["선택적 통계 적용"], del_r

    if si == 4 and ti == 0:
        proc = [
            f"【통합서술】연구설계 전반 절차에 따라 결과를 통합적으로 서술한다. 「{topic['title']}」.",
            f"【명제】{topic['hypothesis']}",
            "【구성】질문–설계–수집–분석–발견–한계–함의.",
        ]
        del_r = [
            ["절", "내용 요지", "근거"],
            ["설계정당화", topic["strategy"][:20], "전략서"],
            ["주요발견", topic["focus"][:12], "표/인용"],
            ["방법론함의", "재현·윤리·한계", "체크리스트"],
            ["종합", topic["hypothesis"][:36], "통합"],
        ]
        return STAGES[4][2][0], proc, (["서술", "Y"], [["한계", "명시"]]), ["통합 결과 서술"], del_r

    # si == 5
    proc = [
        "【산출물】연구방법론 지침서 패키지.",
        f"【핵심】{topic['hypothesis']}",
        "【구성】전략선택·수집계획·전처리·코딩규칙·분석선택·통합서술·체크리스트·서식.",
        f"【모드】{mode} · {method_name}({method_en}).",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["methodology_guide.docx", "연구방법론 지침서", topic["title"][:20]],
        ["strategy_plan.docx", "수집전략·계획", ""],
        ["codebook_dict.xlsx", "코드북/변수사전", ""],
        ["checklist.xlsx", "품질·윤리·재현", ""],
        ["figA–C.png", "파이프라인·전략·체크", ""],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "methodology guide"]]), ["방법론 지침서 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」을(를) {topic['mode']}로 설계할 때, "
        f"{method_name}({method_en})의 전략선택·수집계획·전처리·선택코딩·분석·통합서술 절차로 "
        f"{topic['hypothesis']}을(를) 지침서 형태로 제시할 수 있는가?"
    )
    doc.add_paragraph(f"방법론 명제: {topic['hypothesis']}")
    doc.add_paragraph(f"전략 / 모드 / 초점: {topic['strategy']} / {topic['mode']} / {topic['focus']}")
    doc.add_paragraph(f"대상: {topic['participants']}")

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(
        topic["analysis_key"], topic["title"], topic["mode"], fig_dir, fig_seed
    )

    doc.add_heading("【일반연구방법론 설계 명세】", level=3)
    doc.add_paragraph("전략·설계 모드 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "블록", "ID", "그룹", "파일", "항목", "옵션", "계획항목", "자료", "유형",
        "코드", "변수", "분석", "절", "구분",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str, no: int) -> str:
    base = name_ko
    en_c = en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"일반연구방법론 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 7"
        )
        doc.add_paragraph(
            "본 문서는 일반연구방법론 아키타입의 6단계·7태스크에 따라 "
            "수집전략 선택·수집계획·표준전처리·선택적 주제코딩/변수화·선택적 통계·"
            "통합서술·연구방법론 지침서 패키지를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"], no)
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
