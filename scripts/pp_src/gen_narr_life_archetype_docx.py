# -*- coding: utf-8 -*-
"""Generate Narrative–Life History DOCX: 6 × 5 × 12 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from narr_life_coding_examples import (
    chronology_rows,
    coding_proc_plot,
    coding_proc_struct,
    coding_proc_type,
    narrative_sample_paras,
    plot_element_rows,
    structure_rows,
    typology_rows,
)
from narr_life_docx_profiles import METHOD_PROFILES
from narr_life_figures import figures_for_analysis
from narr_life_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-12.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\서사·생애사 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "서사·생애사 아키타입(Narrative–Life History Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①생애사적 경험을 가진 참여자 선정 →②비구조화·개방형 생애사 면담(복수 회기) →③일기·편지·사진 등 개인기록물 병행 수집",
        [
            "1) 생애사적 경험을 가진 참여자 선정하기",
            "2) 비구조화·개방형 생애사 면담(복수 회기 가능) 실시하기",
            "3) 일기, 편지, 사진 등 개인기록물 병행 수집하기",
        ],
    ),
    (
        "② 전처리",
        "①면담 전사 →②사건을 시간순으로 재배열(chronology 재구성)",
        [
            "1) 면담 전사하기",
            "2) 사건을 시간순으로 재배열(chronology 재구성)하기",
        ],
    ),
    (
        "③ 코딩",
        "①스토리라인/플롯 요소 식별 →②서사구조(시작–전개–절정–결말) →③주제별 내러티브 유형화",
        [
            "1) 스토리라인/플롯 요소(사건-등장인물-맥락-전환점) 식별 코딩하기",
            "2) 서사구조 분석(시작-전개-절정-결말)하기",
            "3) 주제별 내러티브 유형화하기",
        ],
    ),
    (
        "④ 통계기법",
        "해당없음 — 통계적 검정을 사용하지 않음",
        ["1) 해당없음 — 통계적 검정을 사용하지 않음"],
    ),
    (
        "⑤ 결과도출",
        "①개별 생애사를 재구성된 내러티브로 통합 서술 →②사회문화적 맥락과 연결한 의미 해석",
        [
            "1) 개별 생애사를 재구성된 내러티브로 통합 서술하기",
            "2) 사회문화적 맥락과 연결한 의미 해석하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "생애사·내러티브 서술",
        ["▶ 산출물: 생애사·내러티브 서술"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 1 and ti == 1:
        return [figs[0]] if figs else []
    if si == 2 and ti == 1:
        return [figs[1]] if len(figs) > 1 else []
    if si == 2 and ti == 2:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["생애 호(arc)", topic["life_arc"]],
        ["참여자", topic["participants"]],
        ["병행 기록물", topic["docs"]],
        ["분석 초점", topic["focus"]],
        ["분석 키", topic["analysis_key"]],
        ["중심 주장", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    n_p = 4 + s % 3

    if si == 0 and ti == 0:
        proc = [
            f"【선정】「{topic['life_arc']}」 경험을 가진 참여자 {topic['participants']}.",
            f"【기준】초점={topic['focus']}. 유사·대비·부정사례 고려.",
            "【윤리】민감 생애 동의·가명·철회·공개범위.",
            f"【방법】{method_name}({method_en}).",
        ]
        del_r = [
            ["ID", "선정이유", "대비점", "동의"],
            ["P01", "핵심 생애국면", "전형", "Y"],
            ["P02", "동일호·다른경로", "대비", "Y"],
            ["P03", "부정/이탈 사례", "경계", "Y"],
            ["예비", f"n≈{n_p}", "추가가능", "—"],
        ]
        return STAGES[0][2][0], proc, (["목표n", str(n_p)], [["확정", str(n_p)]]), ["참여자 선정 로그"], del_r

    if si == 0 and ti == 1:
        proc = [
            f"【개방면담】「{topic['title']}」 삶을 이야기로 들려달라는 요청으로 시작.",
            "【복수회기】R1 연대기 → R2 전환점 → R3 현재 의미(예시).",
            "【촉진】시간·장소·관계·감정 차원을 따라가되 서사를 끊지 않음.",
            "【기록】녹음 60–120분/회 + 청자 성찰메모.",
        ]
        del_r = [
            ["회기", "초점", "분", "산출"],
            ["R1", "살아온 흐름", "90", "전사"],
            ["R2", "전환점·위기", "80", "전사"],
            ["R3", "현재 의미", "60", "확인·보완"],
        ]
        return STAGES[0][2][1], proc, (["회기", "2–3"], [["개방도", "고"]]), ["면담 회기 로그"], del_r

    if si == 0 and ti == 2:
        proc = [
            f"【기록물】{topic['docs']} 등 동의 하에 수집·목록화.",
            "【대조】구술 연대와 기록 날짜·인물 불일치 표시.",
            "【보관】암호화·접근권한·반환 계획.",
        ]
        del_r = [
            ["ID", "유형", "시기", "구술대조"],
            ["D01", "사진", "유년", "일치"],
            ["D02", "편지/일기", "전환기", "부분불일치"],
            ["D03", "증명·기사", "공적사건", "연대보강"],
        ]
        return STAGES[0][2][2], proc, (["기록물", str(5 + s % 6)], [["동의", "Y"]]), ["개인기록물 목록"], del_r

    if si == 1 and ti == 0:
        proc = [
            "【전사】축어. 침묵·웃음·울음 표기.",
            "【회기별】파일 분리, 메타(날짜·장소·참여자).",
            f"【QC】난해 구간 재생 대조. 「{topic['title']}」.",
        ]
        del_r = [
            ["파일", "회기", "쪽수", "상태"],
            ["P01_R1.docx", "R1", "16", "확정"],
            ["P01_R2.docx", "R2", "14", "확정"],
            ["P02_R1.docx", "R1", "12", "검수중"],
        ]
        return STAGES[1][2][0], proc, (["전사", "완료중"], [["검수", "Y"]]), ["전사 체크"], del_r

    if si == 1 and ti == 1:
        proc = [
            f"【연대기】「{topic['life_arc']}」에 맞춰 사건을 시간순 재배열.",
            "【공백·모순】억지 메우지 않고 표시.",
            "【병기】기록물 날짜를 옆에 둠.",
        ]
        del_r = chronology_rows(topic, s)
        return STAGES[1][2][1], proc, (["연대표", "v1"], [["공백표시", "Y"]]), ["연대기 재구성표"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc_plot(topic, method_name, s)
        del_r = plot_element_rows(topic, s)
        return STAGES[2][2][0], proc, (["플롯코드", "EVT/ACT/CTX/TP"], [["근거", "발췌"]]), ["플롯 요소 코딩"], del_r

    if si == 2 and ti == 1:
        proc = coding_proc_struct(topic, method_name)
        del_r = structure_rows(topic)
        return STAGES[2][2][1], proc, (["구조맵", "Y"], [["평가장치", "표시"]]), ["서사구조 분석"], del_r

    if si == 2 and ti == 2:
        proc = coding_proc_type(topic, method_name)
        del_r = typology_rows(topic)
        return STAGES[2][2][2], proc, (["유형수", "3–5"], [["반례", "포함"]]), ["내러티브 유형화"], del_r

    if si == 3 and ti == 0:
        proc = [
            "【해당없음】통계적 가설검정(t, χ², 회귀 등)을 사용하지 않는다.",
            "【이유】설득력은 서사 재구성·인용·맥락 해석·회원검토에 있다.",
            f"【보조(선택)】참여자 특성·유형 건수 기술만. 주제 「{topic['title']}」.",
            "【금지】빈도로 서사 의미를 ‘증명’한다고 쓰지 않는다.",
        ]
        del_r = [
            ["항목", "결정"],
            ["가설검정", "미사용"],
            ["보조기술", f"n≈{n_p}, 유형건수"],
            ["해석원칙", "서사·맥락 > 빈도"],
        ]
        return STAGES[3][2][0], proc, (["검정", "N/A"], [["보조", "선택"]]), ["통계 미사용 선언"], del_r

    if si == 4 and ti == 0:
        proc = narrative_sample_paras(topic)[:3] + [
            f"【통합】연대기+플롯+유형을 읽히는 내러티브로 서술. ({method_name})",
        ]
        del_r = [
            ["장", "내용"],
            ["I", "삶의 자리매김(시작)"],
            ["II", f"압력과 선택({topic['focus']})"],
            ["III", "전환점과 재구성"],
            ["IV", "현재의 의미"],
        ]
        return STAGES[4][2][0], proc, (["초안", "Y"], [["회원검토", "권장"]]), ["재구성 내러티브 요지"], del_r

    if si == 4 and ti == 1:
        proc = [
            f"【맥락】시대·제도·계층·젠더 등과 「{topic['title']}」 연결.",
            f"【의미】{topic['hypothesis']}",
            "【교차】유형별 사회문화적 위치가 어떻게 다른지.",
            "【한계】소표본·회고 편향·연구자 청자 효과.",
        ]
        del_r = [
            ["층위", "연결", "근거"],
            ["제도/정책", "생애 선택 제약", "CTX코드"],
            ["문화규범", "좋은 삶/교사/노동자상", "EVAL"],
            ["관계망", "조력·배제", "ACT"],
            ["주장", topic["hypothesis"][:28], "유형표"],
        ]
        return STAGES[4][2][1], proc, (["해석", "초안"], [["한계", "명시"]]), ["맥락·의미 해석"], del_r

    # si == 5
    proc = [
        "【산출물】생애사·내러티브 서술 패키지.",
        f"【핵심】{topic['hypothesis']}",
        "【구성】재구성 서술·연대기·플롯/구조·유형·기록물목록·성찰·동의범위.",
        f"【방법】{method_name}({method_en}).",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["narrative.docx", "생애사·내러티브 서술", topic["title"][:20]],
        ["chronology.xlsx", "연대기", topic["life_arc"][:16]],
        ["plot_codes.xlsx", "플롯·구조·유형", ""],
        ["documents_log.xlsx", "기록물", topic["docs"][:14]],
        ["reflexivity.md", "청자·윤리 성찰", ""],
        ["figA–C.png", "연대·구조·유형", ""],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "narrative"]]), ["생애사·내러티브 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」의 생애 서사({topic['life_arc']})는 어떻게 구성되며, "
        f"{method_name}({method_en})으로 {topic['hypothesis']}을(를) "
        f"내러티브로 재구성·해석할 수 있는가?"
    )
    doc.add_paragraph(f"중심 주장: {topic['hypothesis']}")
    doc.add_paragraph(
        f"참여자·기록물·초점: {topic['participants']} / {topic['docs']} / {topic['focus']}"
    )

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    doc.add_heading("【서사·생애사 설계 명세】", level=3)
    doc.add_paragraph("생애호·참여자·기록물·초점 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "ID", "회기", "파일", "연도/시기", "코드", "구조", "유형", "항목", "장", "층위",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str) -> str:
    base, en_c = name_ko, en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"서사·생애사 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 12"
        )
        doc.add_paragraph(
            "본 문서는 서사·생애사 아키타입의 6단계·12태스크(④통계=해당없음)에 따라 "
            "참여자선정·개방면담·기록물·전사·연대기·플롯·서사구조·유형화·내러티브서술·맥락해석 패키지를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"])
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
