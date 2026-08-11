# -*- coding: utf-8 -*-
"""Generate Symbolic Interactionism DOCX: 3 × 5 × 10 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from si_coding_examples import (
    coding_proc_negotiate,
    coding_proc_ritual,
    coding_proc_symbol,
    negotiate_rows,
    ritual_rows,
    symbol_rows,
)
from si_docx_profiles import METHOD_PROFILES
from si_figures import figures_for_analysis
from si_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-18.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\상징적상호작용 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "상징적상호작용 아키타입(Symbolic Interactionism Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①상호작용 장면에 대한 참여관찰 →②관련 행위자 심층면담으로 의미체계 파악",
        [
            "1) 상호작용 장면에 대한 참여관찰하기",
            "2) 관련 행위자 심층면담으로 의미체계 파악하기",
        ],
    ),
    (
        "② 전처리",
        "①현장기록 및 상호작용 장면의 전사 →②상황·맥락 정보 정리",
        [
            "1) 현장기록 및 상호작용 장면의 전사하기",
            "2) 상황·맥락 정보 정리하기",
        ],
    ),
    (
        "③ 코딩",
        "①개념적 코딩으로 상징·의미 단위 식별 →②상호작용 규칙 및 의례적 요소 범주화 →③행위자 간 의미협상 과정 분석",
        [
            "1) 개념적 코딩으로 상징·의미 단위 식별하기",
            "2) 상호작용 규칙 및 의례적 요소 범주화하기",
            "3) 행위자 간 의미협상 과정 분석하기",
        ],
    ),
    (
        "④ 통계기법",
        "해당없음 — 통계적 검정을 사용하지 않음",
        ["1) 해당없음 — 통계적 검정을 사용하지 않음"],
    ),
    (
        "⑤ 결과도출",
        "①상호작용 패턴을 개념화하여 이론적 서술로 통합",
        ["1) 상호작용 패턴을 개념화하여 이론적 서술로 통합하기"],
    ),
    (
        "⑥ 최종 산출물",
        "상호작용 이론적 서술",
        ["▶ 산출물: 상호작용 이론적 서술"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 0 and ti == 0:
        return [figs[0]] if figs else []
    if si == 2 and ti == 1:
        return [figs[1]] if len(figs) > 1 else []
    if si == 2 and ti == 2:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["행위자·현장", topic["participants"]],
        ["분석 모드", topic["mode"]],
        ["초점", topic["focus"]],
        ["분석 키", topic["analysis_key"]],
        ["중심 주장", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    n_obs = 12 + s % 8
    n_int = 8 + s % 5
    is_ir = topic["analysis_key"].startswith("ir_")
    is_ii = topic["analysis_key"].startswith("ii_")

    if si == 0 and ti == 0:
        proc = [
            f"【참여관찰】「{topic['title']}」 상호작용 장면을 관찰·기록한다.",
            f"【현장】행위자={topic['participants']}. 목표 장면 약 {n_obs}회.",
            f"【포인트】말·몸짓·사물·침묵. ({method_name}/{method_en})",
        ]
        del_r = [
            ["장면ID", "일시", "참석", "핵심상징", "메모"],
            ["SC01", "오전", "다수", "호칭", "개시"],
            ["SC02", "오후", "소그룹", "시선", ""],
            ["SC03", "저녁", "2인", topic["focus"][:8], "핵심"],
            ["SC04", "온라인", "팀", "이모지/침묵", ""],
        ]
        return STAGES[0][2][0], proc, (["장면n", str(n_obs)], [["윤리", "동의"]]), ["참여관찰 장면 로그"], del_r

    if si == 0 and ti == 1:
        proc = [
            f"【심층면담】관련 행위자 약 {n_int}명과 면담해 의미체계를 파악한다.",
            "【질문】그 말/몸짓/사물이 당신에게 무엇을 뜻하나요?",
            "【대비】관찰과 면담이 어긋나는 지점을 후속 단서로 표시한다.",
        ]
        del_r = [
            ["ID", "역할", "핵심상징해석", "관찰과의차", "동의"],
            ["P01", "내부자", "친밀표시", "부분일치", "Y"],
            ["P02", "경계자", "거리두기", "불일치", "Y"],
            ["P03", "관찰자", "체면관리", "보완", "Y"],
        ]
        return STAGES[0][2][1], proc, (["면담n", str(n_int)], [["의미체계", "초안"]]), ["면담·의미체계 로그"], del_r

    if si == 1 and ti == 0:
        proc = [
            "【전사】현장기록·장면녹음·면담을 텍스트로 옮긴다.",
            "【표기】화자·순서·비언어(시선·침묵)를 보존한다.",
            f"【QC】난해 구간 재생 대조. 「{topic['title']}」.",
        ]
        del_r = [
            ["파일", "유형", "분량", "화자라벨", "상태"],
            ["SC01.docx", "장면", "6p", "Y", "확정"],
            ["P01.docx", "면담", "10p", "N/A", "확정"],
            ["SC03.docx", "장면", "5p", "Y", "검수중"],
        ]
        return STAGES[1][2][0], proc, (["전사", "진행"], [["QC", "Y"]]), ["전사 체크"], del_r

    if si == 1 and ti == 1:
        proc = [
            "【맥락정리】장소·규범·참석자·직전사건·상황정의를 장면별로 붙인다.",
            f"【초점】{topic['focus']}가 드러나는 맥락 조건을 표시한다.",
            "【링크】장면ID↔면담ID를 연결한다.",
        ]
        del_r = [
            ["장면", "장소", "규범", "상황정의", "연결면담"],
            ["SC01", "현장A", "예의", "공식시작", "P01"],
            ["SC03", "현장B", "또래규범", topic["focus"][:10], "P02"],
            ["SC04", "온라인", "카메라규범", "회의", "P03"],
        ]
        return STAGES[1][2][1], proc, (["맥락표", "Y"], [["링크", "Y"]]), ["상황·맥락 정리표"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc_symbol(topic, method_name, s)
        del_r = symbol_rows(topic, s)
        return STAGES[2][2][0], proc, (["상징단위", str(35 + s % 20)], [["인비보", "Y"]]), ["상징·의미 단위표"], del_r

    if si == 2 and ti == 1:
        proc = coding_proc_ritual(topic, method_name)
        del_r = ritual_rows(topic)
        return STAGES[2][2][1], proc, (["의례범주", "4–5"], [["성패", "표시"]]), ["상호작용 규칙·의례 범주표"], del_r

    if si == 2 and ti == 2:
        proc = coding_proc_negotiate(topic, method_name)
        del_r = negotiate_rows(topic)
        return STAGES[2][2][2], proc, (["협상에피소드", "3–5"], [["재정의", "추적"]]), ["의미협상 과정표"], del_r

    if si == 3 and ti == 0:
        proc = [
            "【원칙】통계적 가설검정을 사용하지 않는다.",
            "【이유】핵심은 상징·의례·의미협상의 이론적 서술이다.",
            f"【주제】「{topic['title']}」.",
        ]
        del_r = [
            ["항목", "결정"],
            ["가설검정", "미사용"],
            ["주근거", "장면·상징·의례·협상 분석"],
            ["보조", "장면 출현빈도(선택, 증명용 아님)"],
        ]
        return STAGES[3][2][0], proc, (["검정", "N/A"], [["원칙", "상호작용"]]), ["통계 미사용 선언"], del_r

    if si == 4 and ti == 0:
        if is_ir:
            proc = [
                f"【의례이론】「{topic['title']}」의 공동초점·정서에너지·연대 패턴을 개념화한다.",
                f"【통합】주장: {topic['hypothesis']}",
                f"【방법】{method_name}({method_en}).",
            ]
        elif is_ii:
            proc = [
                f"【해석이론】전환점에서 상황 정의·재해석 패턴을 개념화한다. 「{topic['title']}」.",
                f"【통합】주장: {topic['hypothesis']}",
                f"【방법】{method_name}({method_en}).",
            ]
        else:
            proc = [
                f"【상징이론】상징이 자아·관계를 매개하는 패턴을 개념화한다. 「{topic['title']}」.",
                f"【통합】상징–의례–협상을 이론적 서술로 연결. 주장: {topic['hypothesis']}",
                f"【방법】{method_name}({method_en}).",
            ]
        del_r = [
            ["개념", "근거패턴", "대표장면", "한줄"],
            ["상징매개", "호칭·사물·시선", "SC01/SC03", "관계표시"],
            ["의례충전", "공동초점·응답", "SC02", "연대"],
            ["의미협상", "재정의·타협", "E2/E3", "조율"],
            ["통합명제", topic["focus"][:12], "다수", topic["hypothesis"][:24]],
        ]
        return STAGES[4][2][0], proc, (["이론서술", "초안"], [["한계", "명시"]]), ["상호작용 패턴 개념화 요지"], del_r

    # si == 5
    proc = [
        "【산출물】상호작용 이론적 서술 패키지.",
        f"【핵심】{topic['hypothesis']}",
        "【구성】장면록·면담·맥락표·상징코드·의례범주·협상과정·이론서술·성찰.",
        f"【모드】{topic['mode']} · {method_name}({method_en}).",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["interaction_theory.docx", "상호작용 이론적 서술", topic["title"][:20]],
        ["scenes.xlsx", "장면·맥락", ""],
        ["codebook.xlsx", "상징·의례·협상", ""],
        ["figA–C.png", "장면·의례·협상", ""],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "interaction theory"]]), ["상호작용 이론 서술 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」의 상호작용 장면을 참여관찰·면담으로 볼 때, "
        f"{method_name}({method_en})의 상징·의례·의미협상 분석으로 "
        f"{topic['hypothesis']}을(를) 이론적으로 서술할 수 있는가?"
    )
    doc.add_paragraph(f"중심 주장: {topic['hypothesis']}")
    doc.add_paragraph(
        f"행위자·현장 / 모드 / 초점: {topic['participants']} / {topic['mode']} / {topic['focus']}"
    )

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    doc.add_heading("【상징적상호작용 분석 명세】", level=3)
    doc.add_paragraph("행위자·장면·초점 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "블록", "ID", "그룹", "파일", "항목", "발췌ID", "범주", "주제", "관계",
        "장면ID", "단위ID", "에피소드", "개념", "장면",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str, no: int) -> str:
    base = name_ko
    en_c = en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"상징적상호작용 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 10"
        )
        doc.add_paragraph(
            "본 문서는 상징적상호작용 아키타입의 6단계·10태스크에 따라 "
            "참여관찰·심층면담·전사·맥락정리·상징코딩·의례범주·의미협상·이론적 서술 패키지를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"], no)
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
