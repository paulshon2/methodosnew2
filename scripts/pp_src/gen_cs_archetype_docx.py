# -*- coding: utf-8 -*-
"""Generate Case Study DOCX: 2 × 5 × 14 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from cs_coding_examples import (
    coding_proc_cross,
    coding_proc_pattern,
    coding_proc_within,
    cross_rows,
    pattern_rows,
    within_rows,
)
from cs_docx_profiles import METHOD_PROFILES
from cs_figures import figures_for_analysis
from cs_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-19.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\QL-020 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "QL-020 아키타입(Case Study Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①사례 선정(단일/다중사례, 전형적·극단적 사례 표집) →②면담·문서·관찰·기록물 등 복수자료원 수집(자료원 삼각검증) →③필요시 설문 등 양적자료 병행 수집",
        [
            "1) 사례 선정(단일/다중사례, 전형적·극단적 사례 표집)하기",
            "2) 면담·문서·관찰·기록물 등 복수자료원 수집(자료원 삼각검증)하기",
            "3) 필요시 설문 등 양적자료 병행 수집하기",
        ],
    ),
    (
        "② 전처리",
        "①자료원별로 정리 및 사례 데이터베이스 구축 →②사례별 연대기 및 사례기술 작성",
        [
            "1) 자료원별로 정리 및 사례 데이터베이스 구축하기",
            "2) 사례별 연대기(chronology) 및 사례기술(case description) 작성하기",
        ],
    ),
    (
        "③ 코딩",
        "①사례 내 분석 →②사례 간 비교분석 →③패턴매칭·설명구축",
        [
            "1) 사례 내 분석(within-case analysis)으로 각 사례 코딩하기",
            "2) 사례 간 비교분석(cross-case analysis)으로 공통 패턴 도출하기",
            "3) 패턴매칭(pattern matching)·설명구축(explanation building) 기법 적용하기",
        ],
    ),
    (
        "④ 통계기법",
        "①원칙적으로 통계 미사용 →②다중사례·혼합설계 시 사례특성 기술통계 보조 제시",
        [
            "1) 원칙적으로 통계 미사용하기",
            "2) 다중사례·혼합설계 시 사례특성 기술통계(빈도, 평균) 보조 제시하기",
        ],
    ),
    (
        "⑤ 결과도출",
        "①사례 내/간 패턴 종합 →②이론적 명제 도출 →③분석적 일반화로 결과 제시",
        [
            "1) 사례 내/사례 간 패턴을 종합하기",
            "2) 이론적 명제(propositions) 도출하기",
            "3) 분석적 일반화(analytic generalization)로 결과 제시하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "사례기술 및 이론적 명제",
        ["▶ 산출물: 사례기술 및 이론적 명제"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 0 and ti == 0:
        return [figs[0]] if figs else []
    if si == 2 and ti == 1:
        return [figs[1]] if len(figs) > 1 else []
    if si == 4 and ti == 2:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["사례·참여자", topic["participants"]],
        ["설계 모드", topic["mode"]],
        ["초점", topic["focus"]],
        ["분석 키", topic["analysis_key"]],
        ["중심 주장", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    is_multi = "다중" in topic["mode"] or "대비" in topic["mode"] or topic["analysis_key"].startswith("cs2_")
    n_cases = 3 + (s % 2) if is_multi else 1
    n_int = 10 + s % 6

    if si == 0 and ti == 0:
        sel = "다중·전형/극단/대비" if is_multi else "단일·목적표집(정보풍부)"
        proc = [
            f"【사례선정】「{topic['title']}」 — {sel}. 모드={topic['mode']}.",
            f"【경계】참여자·현장: {topic['participants']}. 사례 수 가안={n_cases}.",
            f"【논리】선정 이유·제외 이유를 문서화. ({method_name}/{method_en})",
        ]
        del_r = [
            ["사례ID", "유형", "선정이유", "경계", "접근"],
            ["C1", "전형/핵심", "정보풍부", "조직/기간", "동의"],
            ["C2", "대비", "조건차이", "동등경계", "동의"],
            ["C3", "극단/경계", "이론확장", "확장경계", "협의중"],
        ]
        if not is_multi:
            del_r = [
                ["항목", "내용"],
                ["설계", "단일사례"],
                ["선정이유", "과정 깊이·다원자료 가능"],
                ["경계", topic["participants"]],
                ["내포단위", "팀/사건/시기(선택)"],
            ]
        return STAGES[0][2][0], proc, (["사례n", str(n_cases)], [["선정논리", "Y"]]), ["사례 선정 명세"], del_r

    if si == 0 and ti == 1:
        proc = [
            f"【복수자료】면담·문서·관찰·기록물 수집. 면담 약 {n_int}건(+).",
            "【삼각검증】동일 주장에 대해 자료원 교차표를 만든다.",
            f"【초점】{topic['focus']}.",
        ]
        del_r = [
            ["자료원", "건수/분량", "역할", "교차주장예"],
            ["면담", str(n_int), "의미·과정", "조율역량"],
            ["문서", "20+", "공식규칙", "권한구조"],
            ["관찰", "8회+", "실행괴리", "보고문화"],
            ["기록물", "로그/회의록", "시계열", "전환점"],
        ]
        return STAGES[0][2][1], proc, (["삼각검증", "계획"], [["로그", "Y"]]), ["복수자료원 수집·교차표"], del_r

    if si == 0 and ti == 2:
        proc = [
            "【양적병행】필요 시 설문·지표를 보조 수집한다(증명용 아님).",
            "【역할】사례 선별·맥락 요약·이상치 단서.",
            f"【주제】「{topic['title']}」 — 질적 주장과 연결 지점 명시.",
        ]
        del_r = [
            ["양적자료", "내용", "연결지점", "비고"],
            ["미니설문", "인식·만족", "면담 표집", "선택"],
            ["행정지표", "성과/준수율", "문서교차", ""],
            ["미사용시", "—", "질적만으로 충분", "선언"],
        ]
        return STAGES[0][2][2], proc, (["양적", "선택"], [["통합계획", "Y"]]), ["양적자료 병행(선택) 명세"], del_r

    if si == 1 and ti == 0:
        proc = [
            "【사례DB】자료원별 폴더·메타(날짜·출처·동의·태그)를 통일한다.",
            "【인덱싱】사례ID–자료ID–주제태그를 연결한다.",
            "【QC】누락·중복·버전 충돌을 점검한다.",
        ]
        del_r = [
            ["자료ID", "사례", "유형", "태그", "상태"],
            ["D01", "C1", "면담", "촉진", "확정"],
            ["D02", "C1", "문서", "장벽", "확정"],
            ["D03", "C2", "관찰", "전환점", "검수"],
            ["D04", "C3", "기록", "거버넌스", "확정"],
        ]
        return STAGES[1][2][0], proc, (["DB", "구축"], [["QC", "Y"]]), ["사례 데이터베이스 인덱스"], del_r

    if si == 1 and ti == 1:
        proc = [
            f"【연대기·사례기술】「{topic['title']}」 사례별 시간순 사건과 서술 초안을 작성한다.",
            "【공유】팀이 같은 사례상을 보도록 확정본을 둔다.",
            f"【초점】{topic['focus']}가 드러나는 국면을 표시한다.",
        ]
        del_r = [
            ["시점", "사건", "행위자", "자료", "비고"],
            ["T0", "도입/문제인식", "관리층", "D02", ""],
            ["T1", "실행개시", "현장", "D01/D03", ""],
            ["T2", "갈등/전환", "중간관리", "D04", "핵심"],
            ["T3", "안정/잔여", "다수", "종합", ""],
        ]
        return STAGES[1][2][1], proc, (["연대기", "v1"], [["사례기술", "초안"]]), ["연대기·사례기술 요지"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc_within(topic, method_name, s)
        del_r = within_rows(topic, s)
        return STAGES[2][2][0], proc, (["within", "완료"], [["코드북", "v0.9"]]), ["사례 내 코딩표"], del_r

    if si == 2 and ti == 1:
        proc = coding_proc_cross(topic, method_name)
        del_r = cross_rows(topic)
        if not is_multi:
            del_r = [
                ["비교축", "내포단위A", "내포단위B", "패턴"],
                ["촉진", "강", "중", "공통"],
                ["장벽", "중", "강", "대비"],
                ["전환점", "유", "유", "반복"],
                ["초점", "유", "부분", topic["focus"][:12]],
            ]
        return STAGES[2][2][1], proc, (["cross", "Y"], [["예외", "보존"]]), ["사례 간(또는 내포) 비교표"], del_r

    if si == 2 and ti == 2:
        proc = coding_proc_pattern(topic, method_name)
        del_r = pattern_rows(topic)
        return STAGES[2][2][2], proc, (["매칭", "수행"], [["설명", "구축"]]), ["패턴매칭·설명구축 로그"], del_r

    if si == 3 and ti == 0:
        proc = [
            "【원칙】통계적 가설검정을 사용하지 않는다.",
            "【이유】핵심은 사례기술·패턴·이론적 명제·분석적 일반화이다.",
            f"【주제】「{topic['title']}」.",
        ]
        del_r = [
            ["항목", "결정"],
            ["가설검정", "미사용"],
            ["주근거", "within/cross·명제"],
            ["보조", "기술통계(선택)"],
        ]
        return STAGES[3][2][0], proc, (["검정", "N/A"], [["원칙", "QL-020"]]), ["통계 미사용 선언"], del_r

    if si == 3 and ti == 1:
        proc = [
            "【보조】다중사례·혼합설계 시 사례특성 빈도·평균을 보조 제시한다.",
            "【금지】기술통계로 인과·일반화를 ‘증명’한다고 서술하지 않는다.",
            "【제시】해석의 참고용 요약표.",
        ]
        del_r = [
            ["사례", "규모지표", "핵심점수/빈도", "비고"],
            ["C1", str(80 + s % 40), f"{3.2 + (s % 5) * 0.1:.1f}", ""],
            ["C2", str(60 + s % 30), f"{2.8 + (s % 4) * 0.1:.1f}", ""],
            ["C3", str(50 + s % 25), f"{3.0 + (s % 3) * 0.1:.1f}", "보조"],
        ]
        return STAGES[3][2][1], proc, (["기술통계", "보조"], [["증명용", "N"]]), ["사례특성 기술통계(보조)"], del_r

    if si == 4 and ti == 0:
        proc = [
            f"【종합】사례 내/간 패턴을 하나의 결과 틀로 모은다. 「{topic['title']}」.",
            f"【증거】패턴마다 자료원을 연결. 주장 점검: {topic['hypothesis']}",
            f"【방법】{method_name}({method_en}).",
        ]
        del_r = [
            ["패턴", "근거사례", "자료원", "한줄"],
            ["조율역량", "C1/C2", "면담+문서", "촉진"],
            ["안전문화", "C2", "관찰+기록", "준수"],
            ["의제실질", "C3", "회의록", "갈등완화"],
            ["통합", "다수", "삼각검증", topic["hypothesis"][:22]],
        ]
        return STAGES[4][2][0], proc, (["종합", "Y"], [["증거표", "Y"]]), ["패턴 종합표"], del_r

    if si == 4 and ti == 1:
        proc = [
            "【명제】검증·수정 가능한 이론적 명제를 문장화한다.",
            f"【핵심명제】{topic['hypothesis']}",
            "【범위】적용 조건·경계·잔여 예외를 명시한다.",
        ]
        del_r = [
            ["ID", "명제", "지지패턴", "조건/한계"],
            ["P1", topic["hypothesis"][:40], "조율·문화", "유사조직"],
            ["P2", "형식참여≠실질권한", "의제패턴", "토큰참여 시"],
            ["P3", "전환점에서 학습분기", "연대기", "비난문화 시 약화"],
        ]
        return STAGES[4][2][1], proc, (["명제", "3"], [["범위", "명시"]]), ["이론적 명제표"], del_r

    if si == 4 and ti == 2:
        proc = [
            "【분석적 일반화】표본 확장이 아니라 이론·유사사례로 일반화한다.",
            f"【이전】독자가 「{topic['focus']}」 맥락에 전이할 수 있도록 두꺼운 기술을 남긴다.",
            "【한계】사례 특수성·접근제약·관찰자 효과를 명시한다.",
        ]
        del_r = [
            ["구분", "내용"],
            ["일반화유형", "분석적(이론/명제)"],
            ["통계적일반화", "해당없음"],
            ["전이안내", "유사 조건의 조직·정책"],
            ["한계", "n소규모·맥락특수"],
            ["핵심", topic["hypothesis"][:36]],
        ]
        return STAGES[4][2][2], proc, (["analytic gen.", "Y"], [["한계", "Y"]]), ["분석적 일반화 진술"], del_r

    # si == 5
    proc = [
        "【산출물】사례기술 및 이론적 명제 패키지.",
        f"【핵심】{topic['hypothesis']}",
        "【구성】선정논리·복수자료·DB·연대기·within/cross·패턴매칭·명제·분석적 일반화.",
        f"【모드】{topic['mode']} · {method_name}({method_en}).",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["case_report.docx", "사례기술·명제", topic["title"][:20]],
        ["case_db.xlsx", "자료인덱스", ""],
        ["chronology.xlsx", "연대기", ""],
        ["cross_matrix.xlsx", "비교·매칭", ""],
        ["propositions.pdf", "명제·일반화", ""],
        ["figA–C.png", "설계·행렬·흐름", ""],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "case+propositions"]]), ["사례기술·명제 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」을(를) {topic['mode']}로 연구할 때, "
        f"{method_name}({method_en})의 복수자료·within/cross·패턴매칭 절차로 "
        f"{topic['hypothesis']}을(를) 사례기술과 이론적 명제로 제시할 수 있는가?"
    )
    doc.add_paragraph(f"중심 주장: {topic['hypothesis']}")
    doc.add_paragraph(
        f"사례·참여자 / 모드 / 초점: {topic['participants']} / {topic['mode']} / {topic['focus']}"
    )

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    doc.add_heading("【QL-020 설계 명세】", level=3)
    doc.add_paragraph("사례 경계·모드·초점 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "블록", "ID", "그룹", "파일", "항목", "발췌ID", "범주", "주제", "관계",
        "사례ID", "자료원", "양적자료", "자료ID", "시점", "사례", "패턴", "기법",
        "구분",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str, no: int) -> str:
    base = name_ko
    en_c = en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"QL-020 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 14"
        )
        doc.add_paragraph(
            "본 문서는 QL-020 아키타입의 6단계·14태스크에 따라 "
            "사례선정·복수자료·양적병행·사례DB·연대기·within/cross·패턴매칭·"
            "기술통계보조·패턴종합·명제·분석적 일반화·사례기술 패키지를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"], no)
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
