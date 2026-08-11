# -*- coding: utf-8 -*-
"""Generate Mixed Methods Design DOCX: 1 × 5 × 11 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from mm_coding_examples import (
    coding_proc_theme,
    coding_proc_transform,
    quantitize_rows,
    theme_rows,
)
from mm_docx_profiles import METHOD_PROFILES
from mm_figures import figures_for_analysis
from mm_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-28.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\MM-001 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "MM-001 아키타입(Mixed Methods Design Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①혼합설계 유형 결정(수렴적/설명적 순차/탐색적 순차 등) →②질적자료(면담 등)와 양적자료(설문 등)를 설계에 따라 동시 또는 순차적으로 수집",
        [
            "1) 혼합설계 유형 결정(수렴적/설명적 순차/탐색적 순차 등)하기",
            "2) 질적자료(면담 등)와 양적자료(설문 등)를 설계에 따라 동시 또는 순차적으로 수집하기",
        ],
    ),
    (
        "② 전처리",
        "①질적자료 전사,양적자료 코딩(변수화) 각각 수행 →②두 자료를 통합분석을 위한 공통 분석틀로 정리",
        [
            "1) 질적자료 전사,양적자료 코딩(변수화) 각각 수행하기",
            "2) 두 자료를 통합분석을 위한 공통 분석틀로 정리하기",
        ],
    ),
    (
        "③ 코딩",
        "①질적자료에 주제코딩 적용 →②코딩결과를 양적변수로 변환(quantitizing)하거나 그 반대(qualitizing)로 통합 준비",
        [
            "1) 질적자료에 주제코딩 적용하기",
            "2) 코딩결과를 양적변수로 변환(quantitizing)하거나 그 반대(qualitizing)로 통합 준비하기",
        ],
    ),
    (
        "④ 통계기법",
        "①양적자료에 기술통계~추론통계(설계목적에 따라 회귀, ANOVA 등) 적용 →②질적 결과와의 합치도(convergence) 비교",
        [
            "1) 양적자료에 기술통계~추론통계(설계목적에 따라 회귀, ANOVA 등) 적용하기",
            "2) 질적 결과와의 합치도(convergence) 비교하기",
        ],
    ),
    (
        "⑤ 결과도출",
        "①질적·양적 결과를 나란히 배치(side-by-side) 또는 통합표로 비교 →②통합해석(meta-inference) 도출",
        [
            "1) 질적·양적 결과를 나란히 배치(side-by-side) 또는 통합표로 비교하기",
            "2) 통합해석(meta-inference) 도출하기",
        ],
    ),
    (
        "⑥ 최종 산출물",
        "통합적 연구결과 보고서",
        ["▶ 산출물: 통합적 연구결과 보고서"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 0 and ti == 0:
        return [figs[0]] if figs else []
    if si == 3 and ti == 0:
        return [figs[1]] if len(figs) > 1 else []
    if si == 3 and ti == 1:
        return [figs[2]] if len(figs) > 2 else []
    if si == 4 and ti == 0:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["표집·가닥", topic["participants"]],
        ["혼합설계 유형", topic["mode"]],
        ["양적 가닥", topic["quant_strand"]],
        ["질적 가닥", topic["qual_strand"]],
        ["초점", topic["focus"]],
        ["분석 키", topic["analysis_key"]],
        ["통합 명제(초안)", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    n_q = 150 + s % 100
    n_i = 10 + s % 8
    mode = topic["mode"]

    if si == 0 and ti == 0:
        if "설명" in mode:
            rationale = "양적 패턴을 먼저 보고, 차이를 질적으로 설명한다."
            timing = "양→질 순차"
        elif "탐색" in mode:
            rationale = "질적으로 범주를 도출한 뒤 양적으로 일반화·검증한다."
            timing = "질→양 순차"
        else:
            rationale = "동일 개념을 질·양으로 동시 측정해 합치·불일치를 본다."
            timing = "동시·병렬"
        proc = [
            f"【설계결정】「{topic['title']}」에 {mode}를 채택한다.",
            f"【근거】{rationale}",
            f"【통합시점】{timing}. 초점={topic['focus']}. ({method_name}/{method_en})",
        ]
        del_r = [
            ["항목", "결정"],
            ["설계유형", mode],
            ["주도/보조", "양주도" if "설명" in mode else ("질주도" if "탐색" in mode else "동등")],
            ["통합시점", timing],
            ["양가닥", topic["quant_strand"]],
            ["질가닥", topic["qual_strand"]],
            ["연구질문정렬", "동일 개념·다른 증거"],
        ]
        return STAGES[0][2][0], proc, (["설계", mode[:8]], [["문서화", "Y"]]), ["혼합설계 유형 결정서"], del_r

    if si == 0 and ti == 1:
        if "설명" in mode:
            seq = f"1차 설문(n≈{n_q}) → 2차 면담(n≈{n_i})"
        elif "탐색" in mode:
            seq = f"1차 면담(n≈{n_i}) → 2차 설문(n≈{n_q})"
        else:
            seq = f"설문(n≈{n_q})·면담(n≈{n_i}) 동시"
        proc = [
            f"【이중수집】{seq}로 자료를 수집한다.",
            f"【연계】참여자 ID로 질·양 매칭(필요 시). 「{topic['title']}」.",
            "【윤리】동의·익명·철회권. 민감정보 최소화.",
        ]
        del_r = [
            ["가닥", "도구", "n(목표)", "상태"],
            ["양", topic["quant_strand"][:14], str(n_q), "수집"],
            ["질", topic["qual_strand"][:14], str(n_i), "수집"],
            ["순서", seq[:28], "—", mode[:10]],
            ["파일럿", "양·질 각 1회", "—", "양식개정"],
        ]
        return STAGES[0][2][1], proc, (["수집", "Y"], [["연계ID", "Y"]]), ["질·양 자료수집 로그"], del_r

    if si == 1 and ti == 0:
        proc = [
            "【전사】면담 녹음을 표준 규칙(화자·시간·비언어)으로 전사한다.",
            f"【변수화】{topic['quant_strand']} 항목을 분석 변수로 코딩한다.",
            "【QC】결측·역채점·전사 교차점검.",
        ]
        del_r = [
            ["자료", "처리", "완료", "비고"],
            ["면담녹음", "전사", "Y", f"{n_i}건"],
            ["설문응답", "변수화", "Y", f"유효n≈{n_q - s % 12}"],
            ["척도", "신뢰도α", "Y", f"α≈{0.78 + (s % 8) * 0.01:.2f}"],
            ["백업", "원본보관", "Y", ""],
        ]
        return STAGES[1][2][0], proc, (["전사", "Y"], [["변수화", "Y"]]), ["전사·변수화 점검표"], del_r

    if si == 1 and ti == 1:
        proc = [
            "【공통틀】연구질문·개념을 공유하는 통합 분석틀을 구성한다.",
            f"【매핑】질 코드 ↔ 양 변수. 초점={topic['focus']}.",
            "【키】후속 quantitizing·합치비교용 공통키를 부여한다.",
        ]
        del_r = [
            ["개념", "질적코드", "양적변수", "통합역할"],
            ["핵심1", "T01", "척도A", "합치검증"],
            ["핵심2", "T02", "척도B", "보완설명"],
            ["초점", "T03", topic["quant_strand"][:8], "핵심"],
            ["맥락", "메모", "통제변수", "조건부"],
        ]
        return STAGES[1][2][1], proc, (["분석틀", "v1"], [["매핑", "Y"]]), ["공통 분석틀·매핑표"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc_theme(topic, method_name, s)
        del_r = theme_rows(topic, s)
        return STAGES[2][2][0], proc, (["코드북", "v0.9"], [["이중코딩", "표본"]]), ["주제코딩 결과"], del_r

    if si == 2 and ti == 1:
        proc = coding_proc_transform(topic, method_name)
        del_r = quantitize_rows(topic, s)
        return STAGES[2][2][1], proc, (["변환", "규칙화"], [["재현", "Y"]]), ["양적화·질적화 변환표"], del_r

    if si == 3 and ti == 0:
        p = 0.01 + (s % 9) * 0.005
        beta = 0.18 + (s % 10) * 0.02
        proc = [
            f"【양적분석】「{topic['title']}」 양 가닥에 기술·추론통계를 적용한다.",
            f"【모형】회귀/ANOVA 등(설계목적). β≈{beta:.2f}, p≈{p:.3f}(도식).",
            f"【가닥】{topic['quant_strand']}. ({method_name})",
        ]
        del_r = [
            ["분석", "결과요약", "비고"],
            ["기술통계", f"M·SD·n≈{n_q}", ""],
            ["집단비교/회귀", f"β={beta:.2f}, p={p:.3f}", "도식"],
            ["효과크기", f"d/η²≈{0.25 + (s % 5) * 0.05:.2f}", ""],
            ["가정점검", "잔차·등분산", "통과/주의"],
        ]
        return STAGES[3][2][0], proc, (["통계", "완료"], [["가정", "점검"]]), ["양적 분석 요약"], del_r

    if si == 3 and ti == 1:
        conv = 0.58 + (s % 12) * 0.02
        proc = [
            "【합치도】질적 주제와 양적 결과의 convergence를 비교한다.",
            f"【판정(도식)】전체 합치≈{conv:.2f}. 불일치 주제는 별도 해석.",
            f"【초점】{topic['focus']}.",
        ]
        del_r = [
            ["주제", "양적신호", "질적신호", "판정"],
            ["T01", "일치방향", "서사일치", "합치"],
            ["T02", "약함", "강함", "질보완"],
            ["T03", "강함", "부분", "부분합치"],
            ["T04", "불일치", "반례서사", "불일치해석"],
            ["종합", f"≈{conv:.2f}", "—", "meta준비"],
        ]
        return STAGES[3][2][1], proc, (["합치", f"{conv:.2f}"], [["불일치", "기록"]]), ["합치도 비교표"], del_r

    if si == 4 and ti == 0:
        proc = [
            "【나란히】질·양 결과를 side-by-side 또는 통합표로 배치한다.",
            f"【증거】각 칸에 수치·인용을 함께 둔다. 「{topic['title']}」.",
            "【가독】합치/불일치/보완을 색·기호로 표시한다.",
        ]
        del_r = [
            ["연구질문/개념", "양적결과", "질적결과", "통합메모"],
            ["개념A", "평균↓", "단절서사", "합치"],
            ["개념B", "효과↑", "지원서사", "합치"],
            ["초점", topic["focus"][:10], "관련발화", "핵심"],
            ["불일치", "수치↑", "경험↓", "조건부해석"],
        ]
        return STAGES[4][2][0], proc, (["joint", "Y"], [["표", "완성"]]), ["Side-by-side 통합표"], del_r

    if si == 4 and ti == 1:
        proc = [
            f"【통합해석】meta-inference를 도출한다. 명제: {topic['hypothesis']}",
            "【조건】합치·불일치·보완을 한 문장으로 종합한다.",
            "【한계】표집·설계·통합 한계와 일반화 범위를 명시한다.",
        ]
        del_r = [
            ["구분", "내용"],
            ["meta-inference", topic["hypothesis"]],
            ["합치지점", "핵심 개념에서 질·양 일치"],
            ["불일치함의", "하위집단·맥락 조건"],
            ["이론·실무", "개입·설계 시사점"],
            ["한계", "혼합설계·표집·변환규칙"],
        ]
        return STAGES[4][2][1], proc, (["meta", "Y"], [["한계", "명시"]]), ["통합해석(meta-inference)"], del_r

    # si == 5
    proc = [
        "【산출물】통합적 연구결과 보고서 패키지.",
        f"【핵심】{topic['hypothesis']}",
        "【구성】설계결정·수집로그·전사/변수화·분석틀·코딩·변환·통계·합치·통합표·meta·한계.",
        f"【설계】{mode} · {method_name}({method_en}).",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["mm_report.docx", "통합적 연구결과", topic["title"][:20]],
        ["design_memo.docx", "설계유형·통합시점", mode[:12]],
        ["codebook_map.xlsx", "주제·변수 매핑", ""],
        ["stats_joint.xlsx", "통계·합치표", ""],
        ["figA–C.png", "설계·양·합치", ""],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "integrated report"]]), ["통합보고서 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」을(를) {topic['mode']}로 수행할 때, "
        f"{method_name}({method_en})의 이중수집·통합코딩·합치·meta-inference 절차로 "
        f"{topic['hypothesis']}을(를) 근거 있게 제시할 수 있는가?"
    )
    doc.add_paragraph(f"통합 명제: {topic['hypothesis']}")
    doc.add_paragraph(
        f"설계 / 양·질 가닥: {topic['mode']} / {topic['quant_strand']} · {topic['qual_strand']}"
    )
    doc.add_paragraph(f"표집: {topic['participants']}")

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(
        topic["analysis_key"], topic["title"], topic["mode"], fig_dir, fig_seed
    )

    doc.add_heading("【MM-001 명세】", level=3)
    doc.add_paragraph("설계유형·가닥·통합 시점 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "블록", "ID", "그룹", "파일", "항목", "발췌ID", "범주", "주제", "관계",
        "회차", "규칙", "장면ID", "장면", "범주코드", "구간/사건", "지표", "패턴",
        "경향", "구분", "가닥", "자료", "개념", "코드", "참여자", "분석", "연구질문/개념",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str, no: int) -> str:
    base = name_ko
    en_c = en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"MM-001 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 11"
        )
        doc.add_paragraph(
            "본 문서는 MM-001 아키타입의 6단계·11태스크에 따라 "
            "설계유형결정·이중수집·전사/변수화·공통분석틀·주제코딩·양적화/질적화·"
            "양적통계·합치도·side-by-side·meta-inference·통합보고서를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"], no)
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
