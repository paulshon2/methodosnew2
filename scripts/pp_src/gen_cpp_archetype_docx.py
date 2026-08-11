# -*- coding: utf-8 -*-
"""Generate Comparative Politics–Pattern Analysis DOCX: 1 × 5 × 6 (publication-grade)."""
from __future__ import annotations

import json
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Inches, Pt

from cpp_coding_examples import category_rows, coding_proc_cat
from cpp_docx_profiles import METHOD_PROFILES
from cpp_figures import figures_for_analysis
from cpp_task_notes import task_step_annotations

ROOT = Path(__file__).resolve().parents[1]
BUNDLE = ROOT / "public" / "archetypes" / "bundles" / "bundle-27.json"
OUT_DIR = Path(r"C:\Users\saran\OneDrive\Desktop\making-M\비교정치·패턴분석 아키타입-연구방법별 샘플_정교화")
ARCHETYPE = "비교정치·패턴분석 아키타입(Comparative Politics–Pattern Analysis Archetype)"

STAGES = [
    (
        "① 자료수집",
        "①사료·문헌·구술자료 수집",
        ["1) 사료·문헌·구술자료 수집하기"],
    ),
    (
        "② 전처리",
        "①사료비평 및 연대기 정리",
        ["1) 사료비평 및 연대기 정리하기"],
    ),
    (
        "③ 코딩",
        "①주제적·개념적 코딩으로 범주화",
        ["1) 주제적·개념적 코딩으로 범주화하기"],
    ),
    (
        "④ 통계기법",
        "①기초빈도(보조적으로만 사용)",
        ["1) 기초빈도(보조적으로만 사용)하기"],
    ),
    (
        "⑤ 결과도출",
        "①역사적 해석·비교명제 도출",
        ["1) 역사적 해석·비교명제 도출하기"],
    ),
    (
        "⑥ 최종 산출물",
        "역사·비교분석 서술",
        ["▶ 산출물: 역사·비교분석 서술"],
    ),
]


def _seed(*parts) -> int:
    return abs(hash(tuple(parts))) % 997


def add_table(doc, headers, rows):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"
    for j, h in enumerate(headers):
        tbl.rows[0].cells[j].text = str(h)
    for i, row in enumerate(rows, 1):
        for j, val in enumerate(row):
            tbl.rows[i].cells[j].text = str(val)
    doc.add_paragraph()


def add_figure(doc, path: Path, caption: str):
    p = doc.add_paragraph()
    run = p.add_run(caption)
    run.bold = True
    run.font.size = Pt(10)
    if path.exists():
        doc.add_picture(str(path), width=Inches(5.5))
    doc.add_paragraph()


def figures_for_task(si, ti, figs):
    if si == 1 and ti == 0:
        return [figs[0]] if figs else []
    if si == 2 and ti == 0:
        return [figs[1]] if len(figs) > 1 else []
    if si == 4 and ti == 0:
        return [figs[2]] if len(figs) > 2 else []
    if si == 5:
        return figs
    return []


def design_spec_table(topic: dict) -> list[list[str]]:
    return [
        ["항목", "내용"],
        ["연구 주제", topic["title"]],
        ["비교사례", topic["participants"]],
        ["분석 모드", topic["mode"]],
        ["초점", topic["focus"]],
        ["분석 키", topic["analysis_key"]],
        ["비교 명제(초안)", topic["hypothesis"]],
    ]


def task_content(si, ti, topic, method_name, method_en):
    s = _seed(method_name, topic["title"], si, ti)
    n_cases = 4 + s % 3

    if si == 0 and ti == 0:
        proc = [
            f"【수집】「{topic['title']}」에 대해 사료·문헌·구술자료를 수집한다.",
            f"【사례】{topic['participants']} (약 {n_cases}개 단위).",
            f"【로그】출처·판본·접근일. ({method_name}/{method_en})",
        ]
        del_r = [
            ["유형", "예", "역할", "비고"],
            ["사료/법령", "헌법·선거법", "제도배치", ""],
            ["문헌", "이차연구·보고서", "해석대조", ""],
            ["구술/증언", "정치인·활동가", "행위자시점", "선택"],
            ["언론·의회", "토론·표결", "국면재구성", ""],
        ]
        return STAGES[0][2][0], proc, (["사례n", str(n_cases)], [["로그", "Y"]]), ["사료·문헌·구술 수집목록"], del_r

    if si == 1 and ti == 0:
        proc = [
            "【사료비평】출처 신뢰·의도·편향·교차검증을 수행한다.",
            "【연대기】제도변화·위기·개혁을 시간순으로 정리하고 사례 창을 맞춘다.",
            f"【초점】{topic['focus']}.",
        ]
        del_r = [
            ["시점", "사건/제도", "사례A", "사례B", "비평메모"],
            ["T1", "이행/개혁개시", "유", "유", "교차확인"],
            ["T2", "제도선택", "내각제", "대통령제", "타협강도"],
            ["T3", "위기국면", "유", "유", "비상권"],
            ["T4", "재편/현재", "부분", "부분", topic["focus"][:8]],
        ]
        return STAGES[1][2][0], proc, (["연대기", "v1"], [["비평", "Y"]]), ["사료비평·연대기"], del_r

    if si == 2 and ti == 0:
        proc = coding_proc_cat(topic, method_name, s)
        del_r = category_rows(topic, s)
        return STAGES[2][2][0], proc, (["범주", "5"], [["패턴후보", "Y"]]), ["주제·개념 범주표"], del_r

    if si == 3 and ti == 0:
        proc = [
            "【기초빈도】제도유형·국면유형 출현을 보조적으로만 집계한다.",
            "【금지】빈도로 인과·일반화를 증명한다고 서술하지 않는다.",
            f"【주제】「{topic['title']}」 — 핵심은 비교 해석.",
        ]
        del_r = [
            ["항목", "사례수/횟수", "비고"],
            ["합의제 배치", str(2 + s % 2), "보조"],
            ["다수제 배치", str(2 + s % 2), "보조"],
            ["위기국면", str(3 + s % 3), "보조"],
            ["주근거", "연대기·코딩·해석", "빈도≠증명"],
        ]
        return STAGES[3][2][0], proc, (["빈도", "보조"], [["증명용", "N"]]), ["기초빈도(보조)"], del_r

    if si == 4 and ti == 0:
        proc = [
            f"【역사적 해석】맥락을 고려해 패턴을 해석한다. 「{topic['title']}」.",
            f"【비교명제】{topic['hypothesis']}",
            f"【한계】사례 특수·자료편향. ({method_name}/{method_en})",
        ]
        del_r = [
            ["ID", "비교명제", "지지사례", "조건/한계"],
            ["P1", topic["hypothesis"][:40], "다수", "유사국면"],
            ["P2", "타협강→제도선택", "이행기", "외압시변동"],
            ["P3", "견제장치 약화→침식위험", "위기사례", "한시성 조건"],
        ]
        return STAGES[4][2][0], proc, (["명제", "3"], [["한계", "명시"]]), ["역사적 해석·비교명제"], del_r

    # si == 5
    proc = [
        "【산출물】역사·비교분석 서술 패키지.",
        f"【핵심】{topic['hypothesis']}",
        "【구성】자료목록·사료비평·연대기·범주코딩·보조빈도·비교명제·서술·한계.",
        f"【모드】{topic['mode']} · {method_name}({method_en}).",
    ]
    del_r = [
        ["파일", "내용", "비고"],
        ["comparative_democracy.docx", "역사·비교분석 서술", topic["title"][:20]],
        ["sources.xlsx", "사료·문헌목록", ""],
        ["chronology.xlsx", "연대기", ""],
        ["coding.xlsx", "범주·패턴", ""],
        ["figA–C.png", "연대기·행렬·명제", ""],
        ["date", date.today().isoformat(), "v1.0"],
    ]
    return STAGES[5][2][0], proc, (["package", "v1.0"], [["deliverable", "hist-comp narrative"]]), ["역사·비교분석 서술 패키지"], del_r


def write_topic_section(doc, topic_idx, topic, method_name, method_en, method_no):
    doc.add_heading(f"주제{topic_idx}: {topic['title']}", level=2)
    p = doc.add_paragraph()
    p.add_run("연구질문: ").bold = True
    p.add_run(
        f"「{topic['title']}」을(를) 사료·문헌 기반 {topic['mode']}로 볼 때, "
        f"{method_name}({method_en})의 연대기·범주·비교명제 절차로 "
        f"{topic['hypothesis']}을(를) 역사·비교적으로 제시할 수 있는가?"
    )
    doc.add_paragraph(f"비교 명제(초안): {topic['hypothesis']}")
    doc.add_paragraph(
        f"사례 / 모드 / 초점: {topic['participants']} / {topic['mode']} / {topic['focus']}"
    )

    fig_seed = _seed(method_name, topic["title"], method_no)
    fig_dir = OUT_DIR / "_fig_cache" / f"m{method_no}_t{topic_idx}"
    fig_bundle = figures_for_analysis(topic["analysis_key"], topic["title"], fig_dir, fig_seed)

    doc.add_heading("【비교정치·패턴분석 설계 명세】", level=3)
    doc.add_paragraph("비교사례·초점 명세이다. ①자료수집 태스크와 동일 체계로 확장한다.")
    spec = design_spec_table(topic)
    add_table(doc, spec[0], spec[1:])

    header_first = {
        "블록", "ID", "그룹", "파일", "항목", "발췌ID", "범주", "주제", "관계",
        "유형", "시점", "구분",
    }

    for si, (slab, ssum, tasks) in enumerate(STAGES):
        doc.add_heading(slab, level=3)
        doc.add_paragraph(ssum)
        for ti, tlab in enumerate(tasks):
            _, proc, ptbl, dparas, drows = task_content(si, ti, topic, method_name, method_en)
            doc.add_heading(tlab, level=4)

            doc.add_paragraph().add_run("■ 세부 태스크(Procedural Task)").bold = True
            doc.add_paragraph(tlab)

            notes = task_step_annotations(si, ti, proc, topic["analysis_key"], topic, method_name)
            doc.add_paragraph().add_run("■ 각테스크별 주석").bold = True
            for i, (para, note) in enumerate(zip(proc, notes), start=1):
                head = doc.add_paragraph()
                head.add_run(f"({i}) 실행 절차 요지: ").bold = True
                head.add_run(para)
                for block in note.split("\n"):
                    if block.strip():
                        doc.add_paragraph(block.strip())

            doc.add_paragraph().add_run("■ 상세 실행 절차").bold = True
            for para in proc:
                doc.add_paragraph(para)
            if ptbl[1]:
                add_table(doc, ptbl[0], ptbl[1])

            doc.add_paragraph().add_run("■ 결과물 샘플(Sample Deliverable)").bold = True
            for dp in dparas:
                doc.add_paragraph(dp)
            if drows:
                if drows[0][0] in header_first or (len(drows) > 1 and len(drows[0]) == len(drows[1])):
                    add_table(doc, drows[0], drows[1:])
                else:
                    ncols = max(len(r) for r in drows)
                    hdr = ["항목"] + [f"열{i}" for i in range(1, ncols)]
                    add_table(doc, hdr[:ncols], [list(r) + [""] * (ncols - len(r)) for r in drows])

            for cap, fp in figures_for_task(si, ti, fig_bundle):
                add_figure(doc, fp, cap)
            doc.add_paragraph()


def safe_filename(name_ko: str, en: str, no: int) -> str:
    base = name_ko
    en_c = en
    for ch in '<>:"/\\|?*':
        base = base.replace(ch, "")
        en_c = en_c.replace(ch, "-")
    return f"비교정치·패턴분석 아키타입-{base}({en_c}).docx"


def main():
    bundle = json.loads(BUNDLE.read_text(encoding="utf-8"))
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for m in bundle["methods"]:
        no, name = m["no"], m["name"]
        prof = METHOD_PROFILES.get(no)
        if not prof:
            print("skip", no)
            continue
        doc = Document()
        sect = doc.sections[0]
        sect.top_margin = Cm(2)
        sect.bottom_margin = Cm(2)
        t = doc.add_heading(ARCHETYPE, level=0)
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_heading(f"연구방법: {name}({prof['en']})", level=1)
        doc.add_paragraph(
            f"생성일: {date.today().isoformat()} · 번호 {no} · 주제 5개 · 절차 6단계 · 태스크 6"
        )
        doc.add_paragraph(
            "본 문서는 비교정치·패턴분석 아키타입의 6단계·6태스크에 따라 "
            "사료·문헌·구술 수집·사료비평·연대기·주제·개념 범주화·보조빈도·"
            "역사적 해석·비교명제·역사·비교분석 서술 패키지를 "
            "설문·양적통계 및 민족지 샘플과 동일한 출판 형식으로 제시한다."
        )
        doc.add_page_break()
        for i, topic in enumerate(prof["topics"], 1):
            write_topic_section(doc, i, topic, name, prof["en"], no)
            if i < 5:
                doc.add_page_break()
        path = OUT_DIR / safe_filename(name, prof["en"], no)
        tmp = path.with_suffix(".docx.generating")
        doc.save(tmp)
        try:
            if path.exists():
                path.unlink()
            tmp.replace(path)
        except OSError:
            alt = path.with_name(path.stem + "_갱신본.docx")
            if alt.exists():
                alt.unlink()
            tmp.replace(alt)
            print("Wrote (locked)", alt.name)
            continue
        print("Wrote", path.name)
    print("Done.", OUT_DIR)


if __name__ == "__main__":
    main()
